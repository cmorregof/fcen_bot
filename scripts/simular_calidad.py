"""Simula el servicio de calidad con la plantilla FICTICIA y deja lo que vería
cada programa en docs/calidad/simulacion-AAAA-MM-DD.md.

Usa una base temporal: no toca ninguna base real ni envía nada a Telegram.

    uv run python scripts/simular_calidad.py [--hoy AAAA-MM-DD]
"""

from __future__ import annotations

import argparse
import asyncio
import io
import re
import tempfile
from datetime import date, datetime
from html import unescape
from pathlib import Path

import aiosqlite
from openpyxl import load_workbook

from app.servicios.calidad import alertas, plan, reporte, textos
from app.servicios.calidad.db import init_calidad
from app.servicios.calidad.reloj import ahora

RAIZ = Path(__file__).resolve().parents[1]
PLANTILLA = RAIZ / "docs/calidad/plantilla_plan_mejoramiento.xlsx"
EVIDENCIAS = (
    ("ALF-01", "enlace", "https://ejemplo.invalid/autoevaluacion-alfa"),
    ("ALF-03", "archivo", "informe-encuesta.pdf"),
    ("BET-02", "enlace", "https://ejemplo.invalid/guias-laboratorio-beta"),
)


def a_markdown(html: str) -> str:
    texto = re.sub(r"</?b>", "**", html)
    texto = re.sub(r"</?i>", "_", texto)
    return unescape(texto)


async def simular(hoy: date) -> str:
    lectura = plan.leer_plan(PLANTILLA.read_bytes())
    if lectura.errores:
        raise SystemExit("La plantilla no es válida: " + "; ".join(lectura.errores))
    with tempfile.TemporaryDirectory() as tmp:
        db = await aiosqlite.connect(Path(tmp) / "simulacion.db")
        db.row_factory = aiosqlite.Row
        await init_calidad(db)
        marca = datetime.combine(hoy, ahora().timetz()).isoformat()
        resumen = await plan.resumir(db, lectura.filas)
        await plan.aplicar_importacion(
            db, lectura.filas, archivo=PLANTILLA.name, usuario_id=1, creado=marca
        )
        for codigo, tipo, valor in EVIDENCIAS:
            tarea = (await plan.buscar_tareas_por_codigo(db, codigo))[0]
            await plan.registrar_evidencia(
                db, tarea.id, tipo=tipo, valor=valor, referencia=None,
                usuario_id=1, usuario_nombre="Persona Ficticia", creado=marca,
            )  # fmt: skip
        mapa = await alertas.construir_alertas(db, hoy)

        partes = [
            f"# Simulación del servicio de calidad — {hoy.isoformat()}",
            "",
            "Datos 100 % ficticios (`docs/calidad/plantilla_plan_mejoramiento.xlsx`). "
            "Base temporal; no se envió nada a Telegram.",
            "",
            f"**Importación:** {len(resumen.nuevas)} tareas nuevas, "
            f"{len(resumen.actualizadas)} actualizadas, {resumen.sin_cambios} sin cambios.",
            f"**Evidencias registradas:** {', '.join(c for c, _, _ in EVIDENCIAS)}.",
            "",
            f"> {textos.ALCANCE}",
        ]
        for programa in await plan.listar_programas(db):
            tareas = await plan.listar_tareas(db, programa)
            partes += [
                "",
                f"## {programa}",
                "",
                "### Respuesta a `/plan`",
                "",
                a_markdown(textos.texto_plan(programa, tareas, hoy)),
                "",
                "### Alerta quincenal que recibiría",
                "",
                a_markdown(mapa[programa]),
                "",
                "### Reporte `/reporte` (Excel)",
                "",
            ]
            libro = load_workbook(
                io.BytesIO(reporte.generar_reporte(programa, tareas, hoy))
            )
            partes.append(f"Pestañas: {', '.join(libro.sheetnames)}.")
            partes.append("")
            partes.append("| Cifra | Valor | Registros en «fuentes» |")
            partes.append("|---|---|---|")
            fuentes = list(libro["fuentes"].iter_rows(min_row=2, values_only=True))
            for cifra, valor, *_ in libro["estado"].iter_rows(
                min_row=2, values_only=True
            ):
                n = sum(1 for f in fuentes if f[0] == cifra)
                partes.append(f"| {cifra} | {valor} | {n} |")
        await db.close()
    return "\n".join(partes) + "\n"


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--hoy", type=date.fromisoformat, default=ahora().date())
    hoy: date = ap.parse_args().hoy
    destino = RAIZ / f"docs/calidad/simulacion-{hoy.isoformat()}.md"
    destino.write_text(asyncio.run(simular(hoy)), encoding="utf-8")
    print(f"Simulación guardada en {destino}")


if __name__ == "__main__":
    main()
