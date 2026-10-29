"""Crea los planes de mejoramiento FICTICIOS con la plantilla institucional.

Misma estructura de cinco hojas que los planes reales de la oficina de calidad
(«1. Formulación», «2. Plan de Mejora», «3. Cronograma 20XX», «4. Evaluación» y
«Presupuesto») y las mismas variaciones de formato que aparecen entre programas.
Todo es inventado: programas, textos, cargos y enlaces (dominio .invalid).

- docs/calidad/plantilla_plan_mejoramiento.xlsx → «Ejemplo Alfa (ficticio)»:
  encabezado en la fila 5, códigos de tres niveles (1.1.1), fechas como datetime,
  % mezclados (formato 0 % con 0–1, General con 0–100 y el texto «100 %»), una
  fecha de fin reprogramada escrita como dos fechas en la celda, cronogramas
  « 3. Cronograma 2025» (con espacio al inicio) y «3. Cronograma 2026» con
  «Tercera reprogramación» escrito a mano, filas de promedio y una leyenda al pie.
- docs/calidad/plantilla_plan_mejoramiento_beta.xlsx → «Ejemplo Beta (ficticio)»:
  hoja «2. Plan de Mejoramiento 2025» con encabezado en la fila 6, columna de
  factores, códigos de dos niveles (1.1), fechas como texto («1 de enero de
  2023»), % de 0 a 100, cortes «DESCRIPCIÓN (23-02-2026)», una acción marcada
  «NO FUNCIONÓ» en el cronograma, hoja «Cronograma 2026» sin numeración y con las
  fechas de las tareas solo en el Gantt semanal, y «5. Presupuesto».

    uv run python scripts/crear_plantilla_calidad.py
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

DOCS = Path(__file__).resolve().parents[1] / "docs" / "calidad"
ALFA = DOCS / "plantilla_plan_mejoramiento.xlsx"
BETA = DOCS / "plantilla_plan_mejoramiento_beta.xlsx"
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO",
         "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]  # fmt: skip
ENLACE = "https://ejemplo.invalid/drive/"


@dataclass
class Tarea:
    texto: str
    soporte: str = ""
    responsable: str = ""
    drive: str = ""  # enlace o estado escrito a mano
    avance: float | None = None
    ok: bool | None = None
    inicio: date | None = None
    fin: date | None = None
    semanas: list[tuple[int, int]] = field(default_factory=list)  # (mes, semana)
    nota_gantt: str = ""


@dataclass
class Accion:
    codigo: str
    texto: str
    indicador: str
    responsables: str
    inicio: Any
    fin: Any
    avance: Any
    formato: str = "General"
    evidencia: str = ""
    cortes: tuple[str, str] = ("", "")
    marca_crono: str = ""
    tareas: dict[int, list[Tarea]] = field(default_factory=dict)


@dataclass
class Objetivo:
    codigo: str
    texto: str
    meta: str
    oportunidad: str
    linea_base: str
    factor: str
    acciones: list[Accion]


def _negrita(hoja: Worksheet, fila: int) -> None:
    for celda in hoja[fila]:
        celda.font = Font(bold=True)
        celda.alignment = Alignment(wrap_text=True, vertical="center")


def _anchos(hoja: Worksheet) -> None:
    for c in range(1, hoja.max_column + 1):
        hoja.column_dimensions[get_column_letter(c)].width = 22


def _cabecera(hoja: Worksheet, col: int, titulo: str, programa: str) -> None:
    hoja.cell(2, col, titulo)
    hoja.cell(3, col, programa)
    hoja.cell(2, col).font = Font(bold=True, size=13)


# ---------------------------------------------------------------------------
# Hojas
# ---------------------------------------------------------------------------


def formulacion(libro: Workbook, programa: str, objetivos: list[Objetivo]) -> None:
    h = libro.create_sheet("1. Formulación")
    _cabecera(
        h, 3, "OPORTUNIDADES DE MEJORA Y LÍNEA BASE - PLAN DE MEJORAMIENTO ", programa
    )
    h["B4"] = "FECHA DE INICIO DE PROCESO DE P.M: 2025"
    h.append([])
    h.append(["ID", "OPORTUNIDAD DE MEJORA", "LÍNEA BASE", "FACTOR ASOCIADO",
              "CARACTERÍSTICA ASOCIADA", "RELACIÓN CON EL PLAN DE ACCIÓN DE LA FACULTAD"])  # fmt: skip
    _negrita(h, 6)
    for n, o in enumerate(objetivos, start=1):
        h.append([n, o.oportunidad, o.linea_base, o.factor,
                  f"Característica {n} (ejemplo)", f"Eje {n % 3 + 1} del plan de acción (ejemplo)"])  # fmt: skip
    _anchos(h)


def plan_alfa(libro: Workbook, programa: str, objetivos: list[Objetivo]) -> None:
    h = libro.create_sheet("2. Plan de Mejora")
    _cabecera(h, 3, "PLAN DE MEJORAMIENTO ", programa)
    h.merge_cells("C2:R2")
    enc = {
        "B": "OBSERVACIONES CNA  (N° RESOLUCIÓN DE ACREDITACIÓN )",
        "C": "OPORTUNIDAD DE MEJORA/FORTALEZA", "D": "LÍNEA BASE",
        "E": "OBJETIVO(S)", "F": "META", "G": "ACCION(ES)/ACTIVIDAD",
        "H": "INDICADOR DE CUMPLIMIENTO", "I": "RESPONSABLES", "J": "RECURSOS",
        "M": "FECHAS EJECUCIÓN", "O": "% AVANCE", "P": "EVIDENCIAS",
        "Q": "DESCRIPCIÓN DE AVANCE ACTUAL (03-12-2025)",
        "R": "DESCRIPCIÓN DE AVANCE ACTUAL (21-05-2026)",
    }  # fmt: skip
    for col, texto in enc.items():
        h[f"{col}5"] = texto
    for col, texto in {
        "J": "HUMANOS",
        "K": "FÍSICOS",
        "L": "TECNOL.",
        "M": "INICIO",
        "N": "FIN",
    }.items():
        h[f"{col}6"] = texto
    for rango in (
        "J5:L5",
        "M5:N5",
        "O5:O6",
        "B5:B6",
        "G5:G6",
        "P5:P6",
        "Q5:Q6",
        "R5:R6",
    ):
        h.merge_cells(rango)
    _negrita(h, 5)
    _negrita(h, 6)
    h["G7"] = (
        "Descripción de actividades que se deben ejecutar para alcanzar el objetivo."
    )
    h["M7"] = (
        "Pueden ser a corto, mediano o largo plazo, pero tener en cuenta que el plan permanece durante el tiempo de acreditación otorgada"
    )
    h["O7"] = "Tener en cuenta las fechas establecidas para la ejecución de la acción."
    h["B8"] = "COMPONENTE 1. PERTINENCIA (EJEMPLO)"
    h.merge_cells("B8:R8")
    fila = 9
    for o in objetivos:
        primera = fila
        for a in o.acciones:
            h[f"G{fila}"] = f"{a.codigo}. {a.texto}"
            h[f"H{fila}"], h[f"I{fila}"] = a.indicador, a.responsables
            h[f"J{fila}"], h[f"L{fila}"] = "x", "x"
            h[f"M{fila}"], h[f"N{fila}"] = a.inicio, a.fin
            h[f"O{fila}"] = a.avance
            h[f"O{fila}"].number_format = a.formato
            h[f"P{fila}"] = a.evidencia
            h[f"Q{fila}"], h[f"R{fila}"] = a.cortes
            fila += 1
        h[f"B{primera}"] = "Observación del CNA (ejemplo)"
        h[f"C{primera}"], h[f"D{primera}"] = o.oportunidad, o.linea_base
        h[f"E{primera}"], h[f"F{primera}"] = f"{o.codigo}. {o.texto}", o.meta
        if fila - primera > 1:
            for col in "BCDEF":
                h.merge_cells(f"{col}{primera}:{col}{fila - 1}")
    # Filas de promedio al pie, sin acción: el lector las ignora.
    fila += 1
    for valor, formato in ((0.35, "0.00"), (41.5, "0.00"), (20.9, "0.0")):
        h[f"O{fila}"] = valor
        h[f"O{fila}"].number_format = formato
        fila += 1
    _anchos(h)


def plan_beta(libro: Workbook, programa: str, objetivos: list[Objetivo]) -> None:
    h = libro.create_sheet("2. Plan de Mejoramiento 2025")
    h["B2"], h["B3"] = "PLAN DE MEJORAMIENTO", programa
    enc = {
        "A": "FACTORES", "B": "OPORTUNIDAD DE MEJORA/FORTALEZA", "C": "LÍNEA BASE",
        "D": "OBJETIVO(S)", "E": "LÍNEA BASE", "F": "META", "G": "ACCION(ES)/ACTIVIDAD",
        "H": "INDICADOR DE CUMPLIMIENTO", "I": "RESPONSABLES", "J": "RECURSOS",
        "N": "FECHAS EJECUCIÓN", "P": "% AVANCE", "Q": "DESCRIPCIÓN (23-02-2026)",
        "R": "DESCRIPCIÓN (25-06-2026)", "S": "EVIDENCIAS (PENDIENTE)",
    }  # fmt: skip
    for col, texto in enc.items():
        h[f"{col}6"] = texto
    for col, texto in {"J": "HUMANOS", "K": "FÍSICOS", "L": "TECNOL.", "M": "FINANCIEROS",
                       "N": "INICIO", "O": "FIN"}.items():  # fmt: skip
        h[f"{col}7"] = texto
    for rango in ("J6:M6", "N6:O6", "P6:P7"):
        h.merge_cells(rango)
    _negrita(h, 6)
    _negrita(h, 7)
    h["P8"] = "Tener en cuenta las fechas establecidas para la ejecución de la acción."
    fila = 9
    for o in objetivos:
        primera = fila
        for a in o.acciones:
            h[f"G{fila}"] = f"{a.codigo}. {a.texto}"
            h[f"H{fila}"], h[f"I{fila}"] = a.indicador, a.responsables
            h[f"J{fila}"], h[f"M{fila}"] = "x", "N/A"
            h[f"N{fila}"], h[f"O{fila}"] = a.inicio, a.fin
            h[f"P{fila}"] = a.avance
            h[f"Q{fila}"], h[f"R{fila}"] = a.cortes
            h[f"S{fila}"] = a.evidencia
            fila += 1
        h[f"A{primera}"], h[f"B{primera}"] = o.factor, o.oportunidad
        h[f"C{primera}"], h[f"D{primera}"] = o.linea_base, f"{o.codigo}. {o.texto}"
        h[f"E{primera}"], h[f"F{primera}"] = o.linea_base, o.meta
        if fila - primera > 1:
            for col in "ABCDEF":
                h.merge_cells(f"{col}{primera}:{col}{fila - 1}")
    _anchos(h)


def cronograma(
    libro: Workbook,
    titulo: str,
    programa: str,
    anio: int,
    objetivos: list[Objetivo],
    *,
    fila_enc: int,
    con_fechas_tarea: bool,
) -> None:
    """Dos variantes reales: con columnas de fecha por tarea (2025) o con
    soporte y responsable y las fechas solo en el Gantt (2026)."""
    h = libro.create_sheet(titulo)
    h["B1"] = " CRONOGRAMA DE SEGUIMIENTO PARA EL PLAN DE MEJORAMIENTO"
    h["B3"] = f"{programa} {anio}"
    f0, f1, f2 = fila_enc, fila_enc + 1, fila_enc + 2
    for col, texto in {"B": "OBJETIVO", "C": "ACCIÓN", "D": "FECHA ACCIÓN",
                       "F": "INDICADOR DE CUMPLIMIENTO", "G": "TAREAS ESPECÍFICAS"}.items():  # fmt: skip
        h[f"{col}{f0}"] = texto
    h[f"D{f1}"], h[f"E{f1}"] = "Inicio", "Fin"
    h.merge_cells(f"D{f0}:E{f0}")
    if con_fechas_tarea:
        h[f"H{f0}"] = "FECHA TAREA"
        h[f"H{f1}"], h[f"I{f1}"], h[f"J{f1}"], h[f"K{f1}"] = (
            "Inicio",
            "Fin",
            "DRIVE",
            "Avance",
        )
        h[f"J{f2}"], h[f"K{f2}"], h[f"L{f2}"] = f"Año {anio}", "%", "OK"
        h.merge_cells(f"H{f0}:I{f0}")
    else:
        h[f"H{f0}"] = "Cumplimiento de tareas específicas"
        h[f"H{f1}"], h[f"I{f1}"], h[f"J{f1}"], h[f"K{f1}"] = (
            "Soporte",
            "Responsable",
            "DRIVE",
            "Avance",
        )
        h[f"H{f2}"], h[f"I{f2}"], h[f"J{f2}"] = (
            "Tarea Específica",
            "Tarea Específica",
            f"Año {anio}",
        )
        h[f"K{f2}"], h[f"L{f2}"] = "%", "OK"
        h.merge_cells(f"H{f0}:L{f0}")
    gantt = 13  # columna M: cuatro semanas por mes
    for m, mes in enumerate(MESES):
        c = gantt + 4 * m
        h.cell(f0, c, mes)
        h.merge_cells(start_row=f0, start_column=c, end_row=f0, end_column=c + 3)
        for s in range(4):
            h.cell(f1, c + s, "SEMANA")
            h.cell(f2, c + s, float(s + 1))
    _negrita(h, f0)
    fila = f0 + 3
    for o in objetivos:
        h[f"B{fila}"] = f"{o.codigo}. {o.texto}"
        for a in o.acciones:
            tareas = a.tareas.get(anio, [])
            h[f"C{fila}"] = f"{a.marca_crono} {a.codigo}. {a.texto}".strip()
            h[f"D{fila}"], h[f"E{fila}"] = a.inicio, a.fin
            h[f"F{fila}"] = a.indicador
            if not tareas:
                fila += 1
                continue
            primera = fila
            for t in tareas:
                h[f"G{fila}"] = t.texto
                if con_fechas_tarea:
                    h[f"H{fila}"], h[f"I{fila}"] = t.inicio, t.fin
                    h[f"J{fila}"] = t.drive
                    h[f"K{fila}"], h[f"L{fila}"] = t.avance, t.ok
                else:
                    h[f"H{fila}"], h[f"I{fila}"] = t.soporte, t.responsable
                    h[f"J{fila}"] = t.drive
                    h[f"K{fila}"], h[f"L{fila}"] = t.avance, t.ok
                h[f"K{fila}"].number_format = "0%"
                for mes, semana in t.semanas:
                    h.cell(fila, gantt + 4 * (mes - 1) + semana - 1, 1.0)
                if t.nota_gantt:
                    mes, semana = t.semanas[-1] if t.semanas else (6, 1)
                    h.cell(fila, gantt + 4 * (mes - 1) + semana, t.nota_gantt)
                fila += 1
            if fila - primera > 1:
                for col in "CDEF":
                    h.merge_cells(f"{col}{primera}:{col}{fila - 1}")
    # Leyenda de colores al pie, como en los planes reales.
    fila += 2
    for valor, texto in (
        (0.1, "Cumplimiento"),
        (0.2, "Incumplimiento - reprogramar"),
        (0.3, "En curso"),
    ):
        h[f"C{fila}"], h[f"D{fila}"] = valor, texto
        fila += 1
    _anchos(h)


def evaluacion(
    libro: Workbook, titulo: str, programa: str, objetivos: list[Objetivo], fecha: Any
) -> None:
    h = libro.create_sheet(titulo)
    _cabecera(h, 3, "EVALUACIÓN DEL SEGUIMIENTO AL PLAN DE MEJORAMIENTO", programa)
    for col, texto in {"B": "FECHA DE EVALUACIÓN", "C": "OBJETIVO PROPUESTO EN EL PLAN",
                       "D": "LÍNEA BASE ESTABLECIDA EN EL PLAN", "E": "META PROPUESTA EN EL PLAN DE MEJORAMIENTO",
                       "F": "ESTADO ACTUAL DE AVANCE % EN EL PLAN", "G": "DESCRIPCIÓN DEL AVANCE EN A LA FECHA",
                       "H": "ESTADO REAL DE AVANCE %"}.items():  # fmt: skip
        h[f"{col}5"] = texto
    _negrita(h, 5)
    h["H7"] = "Estado real de avance con relación a la meta (ejemplo de instrucción)"
    fila = 9
    for o in objetivos:
        avances = [a.avance for a in o.acciones if isinstance(a.avance, int | float)]
        declarado = sum(avances) / len(avances) if avances else 0
        h[f"B{fila}"], h[f"C{fila}"] = fecha, f"{o.codigo}. {o.texto}"
        h[f"D{fila}"], h[f"E{fila}"] = o.linea_base, o.meta
        h[f"F{fila}"], h[f"H{fila}"] = round(declarado, 2), round(declarado * 0.8, 2)
        h[f"G{fila}"] = "Descripción ficticia del avance del objetivo."
        fila += 1
    _anchos(h)


def presupuesto(libro: Workbook, titulo: str) -> None:
    h = libro.create_sheet(titulo)
    h.append(["PRESUPUESTO (EJEMPLO)"])
    h.append(["Rubro", "Valor", "Fuente"])
    h.append(["Eventos académicos (ejemplo)", 1_000_000, "Facultad (ejemplo)"])
    h.append(["Material de difusión (ejemplo)", 500_000, "Programa (ejemplo)"])
    _anchos(h)


# ---------------------------------------------------------------------------
# Datos ficticios
# ---------------------------------------------------------------------------

DIR = "Dirección del programa (ejemplo)"
COM = "Comité Asesor del programa (ejemplo)"
CAL = "Oficina de Calidad (ejemplo)"


def d(a: int, m: int, dia: int) -> datetime:
    return datetime(a, m, dia)


def objetivos_alfa() -> list[Objetivo]:
    return [
        Objetivo("1.1", "Aumentar el conocimiento del proyecto educativo del programa", "Proyecto educativo difundido", "Difusión del proyecto educativo", "Documento de 2021 (ejemplo)", "Factor 1", [
            Accion("1.1.1", "Publicar piezas de difusión del proyecto educativo", "Dos piezas publicadas por semestre", DIR, d(2025, 2, 1), d(2027, 6, 30), 0.6, "0%", ENLACE + "alfa-111", ("Se publicó una pieza (ficticio).", "Se publicaron dos piezas (ficticio)."), tareas={
                2025: [Tarea("Diseñar la pieza de difusión", inicio=d(2025, 3, 1), fin=d(2025, 4, 30), drive=ENLACE + "t1", avance=1.0, ok=True)],
                2026: [Tarea("Publicar dos piezas en 2026", "Enlaces a las publicaciones", DIR, "Tercera reprogramación", 0.5, False, semanas=[(4, 1), (9, 4)])],
            }),
            Accion("1.1.2", "Realizar jornadas con profesores y estudiantes", "Una jornada por semestre", COM, d(2025, 2, 1), d(2026, 9, 30), 50.0, tareas={
                2026: [Tarea("Realizar la jornada con profesores", "Lista de asistencia", COM, ENLACE + "t2", 0.5, False, semanas=[(5, 2)])],
            }),
            Accion("1.1.3", "Aplicar la encuesta de conocimiento del proyecto educativo", "Resultados de la encuesta", CAL, d(2025, 7, 1), d(2026, 11, 15), 40.0, evidencia=ENLACE + "alfa-113"),
        ]),
        Objetivo("1.2", "Actualizar el proyecto educativo del programa", "Documento actualizado y avalado", "Proyecto educativo desactualizado", "Versión de 2021 (ejemplo)", "Factor 1", [
            Accion("1.2.1", "Ajustar el documento según el plan de estudios", "Documento actualizado", COM, d(2024, 6, 1), d(2025, 8, 30), 100.0, evidencia=ENLACE + "alfa-121", cortes=("Documento entregado (ficticio).", "Sin cambios (ficticio).")),
            Accion("1.2.2", "Obtener el aval del Consejo de Facultad", "Acto de aval", DIR, d(2024, 6, 1), "01/02/2025\n\n30/07/2025", "100 %", cortes=("Aval recibido (ficticio).", "")),
        ]),
        Objetivo("2.1", "Fortalecer el contacto con los egresados", "Dos canales activos", "Comunicación con egresados", "Sin canal propio (ejemplo)", "Factor 6", [
            Accion("2.1.1", "Crear canales de comunicación con egresados", "Dos canales creados", DIR, d(2025, 2, 1), d(2026, 9, 30), 0.3, "0%", ENLACE + "alfa-211", tareas={
                2026: [
                    Tarea("Abrir el grupo de egresados", "Enlace del grupo", DIR, "Tercera reprogramación", 0.3, False, semanas=[(3, 1), (6, 4)], nota_gantt="Fechas pospuestas (ejemplo)"),
                    Tarea("Publicar el boletín semestral", "Boletín en PDF", CAL, ENLACE + "t3", 0.0, False, semanas=[(8, 2)]),
                ],
            }),
            Accion("2.1.2", "Actualizar la base de datos de egresados", "Base actualizada cada semestre", CAL, d(2027, 2, 1), d(2028, 3, 30), 0.0),
        ]),
    ]  # fmt: skip


def objetivos_beta() -> list[Objetivo]:
    return [
        Objetivo("1", "Brindar estrategias para reducir el tiempo de permanencia", "Informe de seguimiento semestral", "Permanencia estudiantil", "Promedio de 3 años (ejemplo)", "Factor 4. Estudiantes", [
            Accion("1.1", "Revisar el tiempo de permanencia por cohorte", "Informe semestral de seguimiento", COM, "1 de enero de 2023", "30 de agosto 2028", 20.0, marca_crono="NO FUNCIONÓ", cortes=("Se revisó una cohorte (ficticio).", "Sin avance (ficticio)."), tareas={
                2026: [Tarea("Estudio de tiempo de permanencia", "Informe", COM, "Tercera reprogramación", 0.2, False, semanas=[(5, 1), (5, 2)])],
            }),
            Accion("1.2", "Entregar la propuesta de asignatura de acompañamiento", "Acta del Comité Asesor", COM, "1 de enero de 2023", "30 de junio de 2026", 40.0, cortes=("Propuesta en borrador (ficticio).", "Propuesta en revisión (ficticio)."), tareas={
                2026: [Tarea("Presentar la propuesta al comité", "Acta", COM, ENLACE + "b1", 0.4, False, semanas=[(3, 2), (6, 4)])],
            }),
        ]),
        Objetivo("2", "Promover actividades con distintos sectores", "Un seminario permanente", "Relación con el sector externo", "Ningún seminario (ejemplo)", "Factor 7. Impacto", [
            Accion("2.1", "Implementar el seminario permanente del programa", "N° de sesiones del seminario", DIR, "1 de enero de 2023", "1 de diciembre de 2026", 75.0, cortes=("Tres sesiones (ficticio).", "Seis sesiones (ficticio)."), tareas={
                2026: [
                    Tarea("Programar las sesiones del semestre", "Cronograma del seminario", DIR, ENLACE + "b2", 1.0, True, semanas=[(2, 1)]),
                    Tarea("Realizar las sesiones", "Listas de asistencia", DIR, ENLACE + "b3", 0.5, False, semanas=[(9, 1), (11, 4)]),
                ],
            }),
            Accion("2.2", "Elaborar el diagnóstico de la situación laboral de egresados", "Informe diagnóstico", CAL, "1 de febrero de 2027", "30 de agosto 2028", 0.0),
            Accion("2.3", "Divulgar los resultados del seminario", "Memorias publicadas", DIR, "1 de enero de 2023", "15 de marzo de 2026", 100.0, cortes=("En preparación (ficticio).", "Memorias publicadas (ficticio).")),
        ]),
    ]  # fmt: skip


def crear_alfa(destino: Path = ALFA) -> None:
    programa = "Ejemplo Alfa (ficticio)"
    objetivos = objetivos_alfa()
    libro = Workbook()
    assert libro.active is not None
    libro.remove(libro.active)
    formulacion(libro, programa, objetivos)
    plan_alfa(libro, programa, objetivos)
    cronograma(
        libro,
        " 3. Cronograma 2025",
        programa,
        2025,
        objetivos,
        fila_enc=5,
        con_fechas_tarea=True,
    )
    cronograma(
        libro,
        "3. Cronograma 2026",
        programa,
        2026,
        objetivos,
        fila_enc=5,
        con_fechas_tarea=False,
    )
    evaluacion(libro, "4. Evaluación 2025", programa, objetivos, d(2025, 12, 3))
    presupuesto(libro, "Presupuesto")
    libro.save(destino)


def crear_beta(destino: Path = BETA) -> None:
    programa = "Ejemplo Beta (ficticio)"
    objetivos = objetivos_beta()
    libro = Workbook()
    assert libro.active is not None
    libro.remove(libro.active)
    formulacion(libro, programa, objetivos)
    plan_beta(libro, programa, objetivos)
    cronograma(
        libro,
        "Cronograma 2026",
        programa,
        2026,
        objetivos,
        fila_enc=7,
        con_fechas_tarea=False,
    )
    evaluacion(libro, "4. Evaluación", programa, objetivos, "23 de febrero de 2026")
    presupuesto(libro, "5. Presupuesto")
    libro.save(destino)


def main() -> None:
    DOCS.mkdir(parents=True, exist_ok=True)
    crear_alfa()
    crear_beta()
    print(f"Planes ficticios en {ALFA} y {BETA}")


if __name__ == "__main__":
    main()
