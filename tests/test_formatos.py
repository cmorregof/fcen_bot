"""Formatos de solicitud: se leen; los demás soportes solo se registran."""

from __future__ import annotations

import io
import shutil
from typing import Any

import aiosqlite
import pytest
from docx import Document as DocxDocument
from PIL import Image

from app.core.config import settings
from app.db import repo
from app.integrations import telegram
from app.services import bot, formatos
from tests.conftest import TelegramFalso
from tests.test_bot import _callback, _registrar, _update


def _formato_docx() -> bytes:
    doc = DocxDocument()
    doc.add_paragraph("FORMATO DE AUTORIZACIÓN PARA CURSAR MENOS DE LA CARGA MÍNIMA")
    doc.add_paragraph("Nombre completo: Persona Ficticia Uno")
    doc.add_paragraph("Documento de identidad: 1000000001")
    tabla = doc.add_table(rows=2, cols=3)
    for celda, valor in zip(
        tabla.rows[0].cells, ["Código", "Asignatura", "Créditos"], strict=True
    ):
        celda.text = valor
    for celda, valor in zip(
        tabla.rows[1].cells, ["4100849", "Termodinámica", "4"], strict=True
    ):
        celda.text = valor
    doc.add_paragraph("III. JUSTIFICACIÓN: trabajo de medio tiempo.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_escaneado() -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (400, 500), "white").save(buf, format="PDF")
    return buf.getvalue()


def _documento(nombre: str, unico: str) -> dict[str, Any]:
    return {"file_id": unico, "file_unique_id": unico, "file_name": nombre}


def test_texto_docx_con_tablas() -> None:
    texto = formatos.texto_docx(_formato_docx())
    assert formatos.es_formato(texto)
    assert "4100849 | Termodinámica | 4" in texto
    assert not formatos.es_formato("Certificado médico\nLa paciente…")


async def test_formato_docx_entra_al_caso(
    db: aiosqlite.Connection,
    tg: TelegramFalso,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await _registrar(db, monkeypatch)

    async def descargar(file_id: str, **_: Any) -> bytes:
        return _formato_docx()

    monkeypatch.setattr(telegram, "descargar_archivo", descargar)
    await bot.manejar_update(db, _update(document=_documento("solicitud.docx", "f1")))
    caso = await repo.get_caso(db, 1)
    assert caso and "--- Formato adjunto: solicitud.docx ---" in caso.texto
    assert "Persona Ficticia Uno" in caso.texto
    assert caso.anexos_archivos == []  # el formato no cuenta como soporte oculto
    assert any("formato de solicitud" in t for t in tg.textos())


@pytest.mark.skipif(shutil.which("pdftoppm") is None, reason="sin poppler")
async def test_pdf_escaneado_pregunta_y_se_convierte(
    db: aiosqlite.Connection,
    tg: TelegramFalso,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await _registrar(db, monkeypatch)

    async def descargar(file_id: str, **_: Any) -> bytes:
        return _pdf_escaneado()

    monkeypatch.setattr(telegram, "descargar_archivo", descargar)
    await bot.manejar_update(db, _update(document=_documento("scan.pdf", "u1")))
    pregunta = next(m for m in tg.mensajes if "escaneado" in m["texto"])
    assert pregunta["botones"][0][0]["callback_data"] == "f:1:u1:si"

    await bot.manejar_update(db, _callback("f:1:u1:si"))
    caso = await repo.get_caso(db, 1)
    assert caso and len(caso.imagenes) == 1 and caso.anexos_archivos == []

    # Un segundo PDF escaneado que es un soporte: no se guarda.
    await bot.manejar_update(db, _update(document=_documento("cert.pdf", "u2")))
    await bot.manejar_update(db, _callback("f:1:u2:no"))
    caso = await repo.get_caso(db, 1)
    assert caso and caso.anexos_archivos == ["cert.pdf"] and len(caso.imagenes) == 1
    assert not list((settings.media_dir / "1").glob("_pendiente_*"))
