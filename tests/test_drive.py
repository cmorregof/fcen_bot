"""Respaldo nocturno y archivo oficial en Drive, con un Drive falso (sin red).
Datos ficticios."""

from __future__ import annotations

import gzip
import io
import itertools
import json
import sqlite3
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.routes import router as router_interno
from app.api.web import crear_router
from app.core.config import settings
from app.db.session import connect
from app.services.drive import archivo, respaldo
from app.services.drive import cliente as drive
from tests.web.test_web import _token, web  # noqa: F401

RAIZ = "raiz"


@dataclass
class _Nodo:
    nombre: str
    padre: str
    carpeta: bool
    contenido: bytes = b""
    papelera: bool = False
    mime: str = ""


@dataclass
class DriveFalso:
    nodos: dict[str, _Nodo] = field(default_factory=dict)
    _ids: itertools.count[int] = field(default_factory=lambda: itertools.count(1))
    llamadas: list[str] = field(default_factory=list)

    def _nuevo(self, nodo: _Nodo) -> str:
        i = f"id{next(self._ids)}"
        self.nodos[i] = nodo
        return i

    async def listar(self, padre: str) -> list[drive.ArchivoDrive]:
        return [
            drive.ArchivoDrive(i, n.nombre, n.carpeta)
            for i, n in self.nodos.items()
            if n.padre == padre and not n.papelera
        ]

    async def crear_carpeta(self, nombre: str, padre: str) -> str:
        return self._nuevo(_Nodo(nombre, padre, True))

    async def subir(self, nombre: str, padre: str, contenido: bytes, mime: str) -> str:
        self.llamadas.append(f"subir {nombre}")
        return self._nuevo(_Nodo(nombre, padre, False, contenido, mime=mime))

    async def reemplazar(self, archivo_id: str, contenido: bytes, mime: str) -> None:
        self.llamadas.append(f"reemplazar {self.nodos[archivo_id].nombre}")
        self.nodos[archivo_id].contenido = contenido

    async def descargar(self, archivo_id: str) -> bytes:
        return self.nodos[archivo_id].contenido

    async def a_papelera(self, archivo_id: str) -> None:
        self.nodos[archivo_id].papelera = True

    # Ayudas de las pruebas
    def ruta(self, i: str) -> str:
        partes = []
        while i in self.nodos:
            partes.append(self.nodos[i].nombre)
            i = self.nodos[i].padre
        return "/".join(reversed(partes))

    def archivos(self) -> dict[str, bytes]:
        return {
            self.ruta(i): n.contenido
            for i, n in self.nodos.items()
            if not n.carpeta and not n.papelera
        }


@pytest.fixture
def configurado(monkeypatch: pytest.MonkeyPatch) -> DriveFalso:
    falso = DriveFalso()
    for campo, valor in (
        ("drive_client_id", "cliente-ficticio"),
        ("drive_client_secret", "secreto-ficticio"),
        ("drive_refresh_token", "token-ficticio"),
        ("drive_carpeta_id", RAIZ),
        ("memoria_practica", False),
    ):
        monkeypatch.setattr(settings, campo, valor)
    monkeypatch.setattr(drive, "obtener", lambda: falso)
    return falso


@pytest.fixture
def apagado(monkeypatch: pytest.MonkeyPatch) -> None:
    for campo in (
        "drive_client_id",
        "drive_client_secret",
        "drive_refresh_token",
        "drive_carpeta_id",
    ):
        monkeypatch.setattr(settings, campo, "")


def _pdf(texto: str) -> bytes:
    """PDF mínimo de una página con una línea de texto."""
    flujo = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode()
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792]"
        b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length "
        + str(len(flujo)).encode()
        + b" >>\nstream\n"
        + flujo
        + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    salida = bytearray(b"%PDF-1.4\n")
    posiciones = []
    for n, obj in enumerate(objetos, start=1):
        posiciones.append(len(salida))
        salida += f"{n} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(salida)
    salida += f"xref\n0 {len(objetos) + 1}\n0000000000 65535 f \n".encode()
    for p in posiciones:
        salida += f"{p:010d} 00000 n \n".encode()
    salida += (
        f"trailer\n<< /Size {len(objetos) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref}\n%%EOF\n"
    ).encode()
    return bytes(salida)


def _base(ruta: Path) -> None:
    con = sqlite3.connect(ruta)
    con.execute("CREATE TABLE casos (id INTEGER PRIMARY KEY, nota TEXT)")
    con.execute("INSERT INTO casos (nota) VALUES ('caso ficticio')")
    con.commit()
    con.close()


# ---------------------------------------------------------------------------
# Respaldo
# ---------------------------------------------------------------------------


def test_nombre_del_respaldo() -> None:
    assert (
        respaldo.nombre("calidad", date(2026, 10, 31), "db")
        == "memoria_calidad_20261031.db.gz"
    )


async def test_sin_configuracion_no_sube_nada(apagado: None) -> None:
    assert not drive.configurado()
    assert drive.obtener() is None
    assert await respaldo.respaldar(None, date(2026, 10, 31)) == {"activo": False}


async def test_practica_nunca_sube(
    configurado: DriveFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "memoria_practica", True)
    assert not drive.configurado()


async def test_respaldo_consistente_y_comprimido(
    configurado: DriveFalso, tmp_path: Path
) -> None:
    _base(settings.db_path)
    r = await respaldo.respaldar(configurado, date(2026, 10, 31), servicio="actas")
    assert r["activo"] is True and r["archivo"] == "memoria_actas_20261031.db.gz"
    archivos = configurado.archivos()
    assert list(archivos) == ["Respaldos/actas/memoria_actas_20261031.db.gz"]
    restaurada = tmp_path / "restaurada.db"
    restaurada.write_bytes(gzip.decompress(next(iter(archivos.values()))))
    con = sqlite3.connect(restaurada)
    assert con.execute("SELECT nota FROM casos").fetchall() == [("caso ficticio",)]
    assert con.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
    con.close()


async def test_respaldo_del_mismo_dia_se_reemplaza(configurado: DriveFalso) -> None:
    _base(settings.db_path)
    hoy = date(2026, 10, 31)
    await respaldo.respaldar(configurado, hoy, servicio="calidad")
    await respaldo.respaldar(configurado, hoy, servicio="calidad")
    assert list(configurado.archivos()) == [
        "Respaldos/calidad/memoria_calidad_20261031.db.gz"
    ]
    assert configurado.llamadas[-1] == "reemplazar memoria_calidad_20261031.db.gz"


async def test_retencion_solo_en_su_carpeta(configurado: DriveFalso) -> None:
    _base(settings.db_path)
    carpeta = await drive.carpeta(configurado, RAIZ, ["Respaldos", "actas"])
    otra = await drive.carpeta(configurado, RAIZ, ["Respaldos", "calidad"])
    for dia in range(1, 31):  # 1 al 30 de octubre
        await configurado.subir(
            f"memoria_actas_202610{dia:02d}.db.gz", carpeta, b"x", ""
        )
    await configurado.subir("memoria_calidad_20260101.db.gz", otra, b"x", "")
    await configurado.subir("nota de Carlos.txt", carpeta, b"x", "")

    r = await respaldo.respaldar(
        configurado, date(2026, 10, 31), servicio="actas", conservar=5
    )
    assert r["eliminadas"] == 26
    quedan = sorted(
        n.split("/")[-1]
        for n in configurado.archivos()
        if n.startswith("Respaldos/actas/")
    )
    assert quedan == [
        "memoria_actas_20261027.db.gz",
        "memoria_actas_20261028.db.gz",
        "memoria_actas_20261029.db.gz",
        "memoria_actas_20261030.db.gz",
        "memoria_actas_20261031.db.gz",
        "nota de Carlos.txt",
    ]
    assert "Respaldos/calidad/memoria_calidad_20260101.db.gz" in configurado.archivos()


def test_ruta_interna_apagada_responde_sin_error(
    apagado: None, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "internal_cron_secret", "cron-prueba")
    app = FastAPI()
    app.include_router(router_interno)
    c = TestClient(app)
    assert c.post("/internal/respaldo-drive").status_code == 403
    r = c.post("/internal/respaldo-drive", headers={"X-Cron-Secret": "cron-prueba"})
    assert r.status_code == 200 and r.json() == {"ok": True, "activo": False}


def test_ruta_interna_respalda(
    configurado: DriveFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    _base(settings.db_path)
    monkeypatch.setattr(settings, "internal_cron_secret", "cron-prueba")
    monkeypatch.setattr(settings, "bot_servicio", "calidad")
    app = FastAPI()
    app.include_router(router_interno)
    r = TestClient(app).post(
        "/internal/respaldo-drive", headers={"X-Cron-Secret": "cron-prueba"}
    )
    assert r.status_code == 200, r.text
    assert r.json()["archivo"].startswith("memoria_calidad_")


# ---------------------------------------------------------------------------
# Archivo oficial
# ---------------------------------------------------------------------------


def _doc(**cambios: object) -> archivo.Documento:
    base: dict[str, object] = {
        "tipo": "acta",
        "numero": "019",
        "fecha": date(2026, 10, 30),
        "programa": "ing_fisica",
        "contenido": _pdf("Acta 019 firmada"),
        "archivado_por": "area@unal.edu.co",
    }
    base.update(cambios)
    return archivo.Documento(**base)  # type: ignore[arg-type]


def _indice(falso: DriveFalso, ruta: str) -> list[tuple[object, ...]]:
    libro = load_workbook(io.BytesIO(falso.archivos()[ruta]))
    ws = libro.worksheets[0]
    return [f[1:12] for f in ws.iter_rows(min_row=11, values_only=True) if f[1]]


async def test_estructura_e_indice(configurado: DriveFalso) -> None:
    hoy = date(2026, 10, 31)
    ruta = await archivo.archivar(configurado, RAIZ, _doc(), hoy)
    assert (
        ruta
        == "Actas/2026/Ciencias Naturales/Ingeniería Física/Acta_019_ing_fisica_20261030.pdf"
    )
    await archivo.archivar(
        configurado,
        RAIZ,
        _doc(numero="Acta 7/2026", programa="matematicas", fecha=date(2026, 9, 1)),
        hoy,
    )
    await archivo.archivar(
        configurado,
        RAIZ,
        _doc(tipo="oficio", numero="Mz.FCEN.1.004-000-2026"),
        hoy,
    )
    archivos = configurado.archivos()
    assert set(archivos) == {
        "Actas/2026/Ciencias Naturales/Ingeniería Física/Acta_019_ing_fisica_20261030.pdf",
        "Actas/2026/Matemáticas y Estadística/Matemáticas/Acta_Acta-7-2026_matematicas_20260901.pdf",
        "Actas/2026/Índice Actas 2026.xlsx",
        "Oficios/2026/Ciencias Naturales/Ingeniería Física/Oficio_Mz.FCEN.1.004-000-2026_ing_fisica_20261030.pdf",
        "Oficios/2026/Índice Oficios 2026.xlsx",
    }
    filas = _indice(configurado, "Actas/2026/Índice Actas 2026.xlsx")
    assert [f[0] for f in filas] == ["1", "2"]
    assert filas[0][:7] == (
        "1",
        "2026-10-30",
        "Acta",
        "019",
        "Ciencias Naturales",
        "Ingeniería Física",
        "Acta_019_ing_fisica_20261030.pdf",
    )
    assert filas[0][7:9] == ("2026-10-31", "PDF")
    assert filas[1][3] == "Acta 7/2026"
    assert len(_indice(configurado, "Oficios/2026/Índice Oficios 2026.xlsx")) == 1


async def test_no_se_sobrescribe(configurado: DriveFalso) -> None:
    await archivo.archivar(configurado, RAIZ, _doc())
    with pytest.raises(archivo.YaArchivado):
        await archivo.archivar(configurado, RAIZ, _doc())
    filas = _indice(configurado, "Actas/2026/Índice Actas 2026.xlsx")
    assert len(filas) == 1


@pytest.mark.parametrize(
    "contenido",
    [b"PK\x03\x04 un docx", _pdf("Oficio [la firma se agrega al firmar]")],
    ids=["word", "pdf-borrador"],
)
async def test_no_se_suben_borradores(
    configurado: DriveFalso, contenido: bytes
) -> None:
    with pytest.raises(archivo.DocumentoNoValido):
        await archivo.archivar(configurado, RAIZ, _doc(contenido=contenido))
    assert configurado.archivos() == {}


def test_las_marcas_son_las_de_los_borradores() -> None:
    """Si cambia el texto del borrador, el filtro tiene que cambiar con él."""
    from app.services import documento, oficio_consejo

    fuentes = [
        Path(m.__file__ or "").read_text(encoding="utf-8")
        for m in (documento, oficio_consejo)
    ]
    assert archivo.MARCAS_BORRADOR[0] in fuentes[0]
    assert archivo.MARCAS_BORRADOR[1] in fuentes[1]


# ---------------------------------------------------------------------------
# Web
# ---------------------------------------------------------------------------


def _web(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setattr(settings, "bot_servicio", "actas")
    app = FastAPI()
    app.include_router(crear_router("actas"))
    return TestClient(app)


async def _caso_expedido(numero: str | None) -> int:
    db = await connect()
    try:
        cur = await db.execute(
            "INSERT INTO casos (chat_id, user_id, direccion, programa, estado,"
            " sentido, extraccion_json) VALUES (1, 1, 'ciencias_naturales',"
            " 'ing_fisica', 'cerrado', 'aprobar', ?)",
            (json.dumps({"tipos": ["cancelacion"]}),),
        )
        caso_id = int(cur.lastrowid or 0)
        if numero is not None:
            await db.execute(
                "INSERT INTO notificaciones (caso_id, acto, numero, fecha_expedicion,"
                " registrado_por) VALUES (?, 'resolucion', ?, '2026-10-29', 's@unal.edu.co')",
                (caso_id, numero),
            )
        await db.commit()
        return caso_id
    finally:
        await db.close()


async def test_web_apagada(apagado: None, monkeypatch: pytest.MonkeyPatch) -> None:
    c = _web(monkeypatch)
    cab = await _token("area@unal.edu.co", "administrativa", "ciencias_naturales")
    r = c.get("/api/web/actas/archivo", headers=cab)
    assert r.json()["activo"] is False
    assert {p["clave"] for p in r.json()["programas"]} >= {"ing_fisica"}
    r = c.post(
        "/api/web/actas/archivo/acta",
        data={"programa": "ing_fisica", "numero": "019", "fecha": "2026-10-30"},
        files={"pdf": ("acta.pdf", _pdf("Acta firmada"), "application/pdf")},
        headers=cab,
    )
    assert r.status_code == 503


async def test_web_acta_del_area(
    configurado: DriveFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    c = _web(monkeypatch)
    cab = await _token("area@unal.edu.co", "administrativa", "ciencias_naturales")
    datos = {"programa": "ing_fisica", "numero": "019", "fecha": "2026-10-30"}
    pdf = {"pdf": ("acta.pdf", _pdf("Acta firmada"), "application/pdf")}
    r = c.post("/api/web/actas/archivo/acta", data=datos, files=pdf, headers=cab)
    assert r.status_code == 200, r.text
    assert r.json()["ruta"].startswith(
        "Actas/2026/Ciencias Naturales/Ingeniería Física/"
    )
    # Repetida: no se sobrescribe.
    r = c.post("/api/web/actas/archivo/acta", data=datos, files=pdf, headers=cab)
    assert r.status_code == 409
    # Programa de otra área.
    otro = {**datos, "programa": "matematicas"}
    r = c.post("/api/web/actas/archivo/acta", data=otro, files=pdf, headers=cab)
    assert r.status_code == 404
    # Un Word es un borrador.
    word = {"pdf": ("acta.docx", b"PK\x03\x04", "application/octet-stream")}
    r = c.post(
        "/api/web/actas/archivo/acta",
        data={**datos, "numero": "020"},
        files=word,
        headers=cab,
    )
    assert r.status_code == 422
    # El comité no archiva.
    comite = await _token("comite@unal.edu.co", "comite", "ciencias_naturales")
    r = c.post("/api/web/actas/archivo/acta", data=datos, files=pdf, headers=comite)
    assert r.status_code == 403


async def test_web_secretaria_exige_acto_expedido(
    configurado: DriveFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    c = _web(monkeypatch)
    cab = await _token("secretaria@unal.edu.co", "secretaria")
    pdf = {"pdf": ("res.pdf", _pdf("Resolucion firmada"), "application/pdf")}
    sin_registro = await _caso_expedido(None)
    r = c.post(
        f"/api/web/actas/secretaria/casos/{sin_registro}/archivo",
        files=pdf,
        headers=cab,
    )
    assert r.status_code == 409
    assert configurado.archivos() == {}

    expedido = await _caso_expedido("Mz.FCEN.1.010-000-2026")
    r = c.post(
        f"/api/web/actas/secretaria/casos/{expedido}/archivo", files=pdf, headers=cab
    )
    assert r.status_code == 200, r.text
    assert set(configurado.archivos()) == {
        "Resoluciones/2026/Ciencias Naturales/Ingeniería Física/"
        "Resolución_Mz.FCEN.1.010-000-2026_ing_fisica_20261029.pdf",
        "Resoluciones/2026/Índice Resoluciones 2026.xlsx",
    }
    adm = await _token("adm@unal.edu.co", "administrativa", "ciencias_naturales")
    r = c.post(
        f"/api/web/actas/secretaria/casos/{expedido}/archivo", files=pdf, headers=adm
    )
    assert r.status_code == 403


# ---------------------------------------------------------------------------
# Cliente HTTP, contra un transporte simulado (sin red)
# ---------------------------------------------------------------------------


async def test_cliente_http(monkeypatch: pytest.MonkeyPatch) -> None:
    import httpx

    pedidos: list[tuple[str, str]] = []

    def responder(req: httpx.Request) -> httpx.Response:
        pedidos.append((req.method, req.url.path))
        if req.url.host == "oauth2.googleapis.com":
            assert b"refresh_token=token-ficticio" in req.content
            return httpx.Response(200, json={"access_token": "a", "expires_in": 3600})
        assert (
            req.headers.get("Authorization") == "Bearer a" or req.url.host == "subida"
        )
        if req.method == "GET" and req.url.params.get("alt") == "media":
            return httpx.Response(200, content=b"datos")
        if req.method == "GET":
            if req.url.params.get("pageToken"):
                return httpx.Response(
                    200, json={"files": [{"id": "2", "name": "b", "mimeType": "x"}]}
                )
            return httpx.Response(
                200,
                json={
                    "files": [{"id": "1", "name": "a", "mimeType": drive.MIME_CARPETA}],
                    "nextPageToken": "p2",
                },
            )
        if req.url.params.get("uploadType") == "resumable":
            return httpx.Response(200, headers={"Location": "https://subida/sesion"})
        if req.url.host == "subida":
            return httpx.Response(200, json={"id": "nuevo"})
        return httpx.Response(200, json={"id": "carpeta"})

    transporte = httpx.MockTransport(responder)
    original = httpx.AsyncClient
    monkeypatch.setattr(
        drive.httpx,
        "AsyncClient",
        lambda **kw: original(transport=transporte, **kw),
    )
    c = drive.ClienteDriveHttp("id", "secreto", "token-ficticio")
    assert await c.listar("raiz") == [
        drive.ArchivoDrive("1", "a", True),
        drive.ArchivoDrive("2", "b", False),
    ]
    assert await c.crear_carpeta("Actas", "raiz") == "carpeta"
    assert await c.subir("x.pdf", "carpeta", b"%PDF-", "application/pdf") == "nuevo"
    assert await c.descargar("nuevo") == b"datos"
    await c.a_papelera("nuevo")
    # El token se pide una sola vez y se reutiliza.
    assert sum(1 for _, ruta in pedidos if ruta == "/token") == 1
    assert ("PUT", "/sesion") in pedidos


async def test_cliente_http_error(monkeypatch: pytest.MonkeyPatch) -> None:
    import httpx

    transporte = httpx.MockTransport(
        lambda req: httpx.Response(400, json={"error": "invalid_grant"})
    )
    original = httpx.AsyncClient
    monkeypatch.setattr(
        drive.httpx, "AsyncClient", lambda **kw: original(transport=transporte, **kw)
    )
    with pytest.raises(drive.ErrorDrive, match="invalid_grant"):
        await drive.ClienteDriveHttp("id", "s", "t").listar("raiz")
