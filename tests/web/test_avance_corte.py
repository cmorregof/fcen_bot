"""Borrador de avance de un corte desde la web, con el plan FICTICIO y el modelo simulado."""

from __future__ import annotations

import re
import time
from datetime import date
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import BaseModel

from app.servicios.calidad.avance import borrador as bor
from app.servicios.calidad.avance import corrida as avance
from app.servicios.calidad.avance import fuentes
from app.servicios.calidad.config import calidad_settings
from app.servicios.calidad.soportes.cliente import Respuesta
from app.servicios.calidad.soportes.describir import Cache
from tests.calidad.test_avance import ModeloFalso as Modelo
from tests.calidad.test_avance import con_enlace, guardado
from tests.web.test_revisor_soportes import _plan_cargado
from tests.web.test_web import Cliente, _token, cliente, web  # noqa: F401

RUTA = "/api/web/calidad/avance"
CORTE = "2026-11-30"


class ModeloFalso:
    def __init__(self) -> None:
        self.llamadas = 0
        self.clientes = 0  # veces que se pidió el cliente de la API

    async def llamar(
        self,
        *,
        modelo: str,
        sistema: str,
        bloques: list[dict[str, Any]],
        texto: str,
        salida: type[BaseModel],
        nombre_herramienta: str,
        esfuerzo: str,
        max_tokens: int,
    ) -> Respuesta:
        self.llamadas += 1
        nuevas = re.findall(r"^\[(\w+)\] \(nueva en este corte\)", texto, re.M)
        frases = [{"texto": f"Hay un enlace [{i}].", "citas": [i]} for i in nuevas]
        return Respuesta({"frases": frases}, 1000, 200, modelo=modelo)


@pytest.fixture
def falso(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> ModeloFalso:
    modelo = ModeloFalso()
    monkeypatch.setattr(calidad_settings, "avance_activo", True)
    monkeypatch.setattr(calidad_settings, "avance_tope_usd", 2.0)
    monkeypatch.setattr(calidad_settings, "revisor_dir", tmp_path / "volumen")
    monkeypatch.setattr(avance, "_tarea", None)

    def crear() -> ModeloFalso:
        modelo.clientes += 1
        return modelo

    monkeypatch.setattr(avance, "crear_cliente", crear)
    return modelo


def _esperar(c: TestClient, cab: dict[str, str]) -> dict[str, Any]:
    for _ in range(400):
        estado: dict[str, Any] = c.get(RUTA, headers=cab).json()
        if estado["estado"] != "corriendo":
            return estado
        time.sleep(0.025)
    raise AssertionError("el borrador no terminó")


def _registrar_enlace(c: TestClient, cab: dict[str, str]) -> None:
    r = c.post(
        "/api/web/calidad/evidencia",
        json={
            "programa": "alfa",
            "codigo": "1.1.1",
            "enlace": "https://ejemplo.invalid/drive/nuevo",
        },
        headers=cab,
    )
    assert r.status_code == 200, r.text


async def test_apagado_responde_403_y_no_llama_al_modelo(
    cliente: Cliente,  # noqa: F811
    falso: ModeloFalso,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(calidad_settings, "avance_activo", False)
    c = cliente("calidad")
    cab = await _token("lili@unal.edu.co", "calidad")
    pedido = {"programa": "alfa", "corte": CORTE, "confirmo_costo": True}
    for metodo, ruta, cuerpo in [
        ("GET", RUTA, None),
        ("GET", f"{RUTA}/cortes?programa=alfa", None),
        ("POST", f"{RUTA}/estimar", pedido),
        ("POST", f"{RUTA}/generar", pedido),
        ("GET", f"{RUTA}/descargar/docx", None),
        ("DELETE", RUTA, None),
    ]:
        r = c.request(metodo, ruta, json=cuerpo, headers=cab)
        assert r.status_code == 403, ruta
        assert "apagado" in r.json()["detail"]
    assert falso.llamadas == 0
    assert not avance.base().exists()


@pytest.mark.parametrize(("rol", "esperado"), [("calidad", 200), ("decanatura", 403)])
async def test_solo_calidad_y_admin(
    rol: str,
    esperado: int,
    cliente: Cliente,  # noqa: F811
    falso: ModeloFalso,
) -> None:
    c = cliente("calidad")
    cab = await _token(f"{rol}@unal.edu.co", rol)
    assert c.get(RUTA, headers=cab).status_code == esperado
    assert c.get(RUTA).status_code == 401


async def test_flujo_completo_estimar_confirmar_generar_descargar(
    cliente: Cliente,  # noqa: F811
    falso: ModeloFalso,
) -> None:
    with cliente("calidad") as c:
        cab = await _token("lili@unal.edu.co", "calidad")
        await _plan_cargado(c, cab)
        _registrar_enlace(c, cab)

        r = c.get(f"{RUTA}/cortes?programa=alfa", headers=cab)
        assert r.json()["cortes"] == ["2026-05-21", "2025-12-03"]
        assert r.json()["proximo"] is not None

        r = c.post(
            f"{RUTA}/estimar", json={"programa": "alfa", "corte": CORTE}, headers=cab
        )
        est = r.json()
        assert r.status_code == 200, r.text
        assert est["acciones"] == 7 and est["al_modelo"] == 1
        assert est["sin_evidencia_nueva"] == 6
        assert 0 < est["costo_usd"] < est["tope_usd"] and not est["excede_tope"]
        assert falso.llamadas == 0

        malo = c.post(
            f"{RUTA}/estimar", json={"programa": "alfa", "corte": "mañana"}, headers=cab
        )
        assert malo.status_code == 400

        # Sin confirmar el costo, o sin la cifra mostrada, no corre.
        r = c.post(
            f"{RUTA}/generar", json={"programa": "alfa", "corte": CORTE}, headers=cab
        )
        assert r.status_code == 400 and falso.llamadas == 0
        pedido = {"programa": "alfa", "corte": CORTE, "confirmo_costo": True}
        assert c.post(f"{RUTA}/generar", json=pedido, headers=cab).status_code == 400
        # La confirmación va atada a la cifra: otra cifra, no corre.
        r = c.post(f"{RUTA}/generar", json={**pedido, "costo_usd": 0.0}, headers=cab)
        assert r.status_code == 409 and "La estimación cambió" in r.json()["detail"]
        assert falso.llamadas == 0 and falso.clientes == 0

        r = c.post(
            f"{RUTA}/generar",
            json={
                "programa": "alfa",
                "corte": CORTE,
                "confirmo_costo": True,
                "costo_usd": est["costo_usd"],
            },
            headers=cab,
        )
        assert r.status_code == 200, r.text
        estado = _esperar(c, cab)
        assert estado["estado"] == "terminado", estado
        assert falso.llamadas == 1
        assert estado["resumen"]["sin_evidencia_nueva"] == 6
        assert estado["resumen"]["con_propuesta"] == 1
        assert sorted(estado["descargas"]) == ["docx", "xlsx"]

        docx = c.get(f"{RUTA}/descargar/docx", headers=cab)
        xlsx = c.get(f"{RUTA}/descargar/xlsx", headers=cab)
        assert docx.status_code == xlsx.status_code == 200
        assert docx.content[:2] == xlsx.content[:2] == b"PK"
        assert ".docx" in docx.headers["content-disposition"]
        assert c.get(f"{RUTA}/descargar/otro", headers=cab).status_code == 404

        carpeta = avance.base()
        doc = Document(str(carpeta / "borrador_avance.docx"))
        assert doc.paragraphs[0].text.startswith("BORRADOR — propuesta")
        assert "Sin evidencia nueva en el corte." in "\n".join(
            p.text for p in doc.paragraphs
        )
        ws = load_workbook(carpeta / "revision_avance.xlsx")["Revisión del corte"]
        assert ws.max_row == 8

        # Repetir el mismo borrador no paga otra vez.
        r = c.post(
            f"{RUTA}/estimar", json={"programa": "alfa", "corte": CORTE}, headers=cab
        )
        assert r.json()["al_modelo"] == 0 and r.json()["en_cache"] == 1

        assert c.delete(RUTA, headers=cab).json()["estado"] == "vacio"
        assert not carpeta.exists()


async def test_tope_de_costo_no_deja_generar(
    cliente: Cliente,  # noqa: F811
    falso: ModeloFalso,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(calidad_settings, "avance_tope_usd", 0.0001)
    with cliente("calidad") as c:
        cab = await _token("lili@unal.edu.co", "calidad")
        await _plan_cargado(c, cab)
        _registrar_enlace(c, cab)
        r = c.post(
            f"{RUTA}/estimar", json={"programa": "alfa", "corte": CORTE}, headers=cab
        )
        assert "pasa del tope" in r.json()["excede_tope"]
        r = c.post(
            f"{RUTA}/generar",
            json={
                "programa": "alfa",
                "corte": CORTE,
                "confirmo_costo": True,
                "costo_usd": 0.0,
            },
            headers=cab,
        )
        assert r.status_code == 409 and "pasa del tope" in r.json()["detail"]
        assert falso.llamadas == 0 and falso.clientes == 0
        assert c.get(RUTA, headers=cab).json()["estado"] == "vacio"


async def test_el_cliente_con_tope_detiene_la_corrida(tmp_path: Path) -> None:
    """Aunque la estimación quepa, el gasto real tampoco pasa del tope."""
    pg = guardado()
    for codigo in ("1.1.1", "1.1.2", "1.2.1"):
        con_enlace(pg, codigo, f"https://ejemplo.invalid/{codigo}", "2026-10-01T10:00")
    modelo = Modelo()
    acciones = fuentes.armar(pg, date.fromisoformat(CORTE), {}, set())
    b = await bor.generar(
        "Ejemplo", CORTE, acciones, set(), modelo, Cache(tmp_path), tope_usd=0.0001
    )
    assert modelo.textos == [] and "tope" in b.detenido
    assert all(a.error for a in b.acciones if not a.sin_evidencia_nueva)
