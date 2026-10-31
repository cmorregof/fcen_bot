"""Notificaciones de la Secretaría de Facultad. Datos ficticios."""

from __future__ import annotations

import io
import json
from datetime import date

import pytest
from openpyxl import load_workbook

from app.core.config import settings
from app.db.session import connect
from app.services import notificaciones
from tests.web.test_web import Cliente, _token, cliente, web  # noqa: F401

RUTA = "/api/web/actas/secretaria/notificaciones"


async def _caso_decidido(
    acta_fecha: str = "2026-09-17", sentido: str = "aprobo", instancia: str = "facultad"
) -> int:
    db = await connect()
    try:
        await db.execute(
            "INSERT OR IGNORE INTO actas_consejo (id, instancia, numero, fecha, texto,"
            " subido_por) VALUES (1, 'facultad', '35', ?, 'acta ficticia', 'x@unal.edu.co')",
            (acta_fecha,),
        )
        cur = await db.execute(
            "INSERT INTO casos (chat_id, user_id, direccion, programa, estado,"
            " sentido, extraccion_json, recomendado) VALUES (1, 1, 'fisica',"
            " 'ing_fisica', 'cerrado', 'aprobar', ?, '2026-09-10 15:00:00')",
            (
                json.dumps(
                    {
                        "tipos": ["cancelacion"],
                        "nombre_estudiante": "Persona Ficticia Uno",
                    }
                ),
            ),
        )
        caso_id = int(cur.lastrowid or 0)
        await db.execute(
            "INSERT INTO decisiones_consejo (caso_id, acta_id, decision_id, instancia,"
            " sentido, texto_decision, acta_numero, acta_fecha, fragmento,"
            " confirmado_por) VALUES (?, 1, 1, ?, ?, 'texto', '35', ?, 'frag',"
            " 'x@unal.edu.co')",
            (caso_id, instancia, sentido, acta_fecha),
        )
        await db.commit()
        return caso_id
    finally:
        await db.close()


def test_semestre() -> None:
    hoy = date(2026, 10, 30)
    s = notificaciones.semestre_de("", hoy)
    assert (s.texto, s.desde, s.hasta) == (
        "2026-2S",
        date(2026, 7, 1),
        date(2026, 12, 31),
    )
    assert notificaciones.semestre_de("2026-1s", hoy).hasta == date(2026, 6, 30)
    with pytest.raises(notificaciones.PeriodoInvalido):
        notificaciones.semestre_de("2026-3S", hoy)


def test_orden_de_fechas() -> None:
    hoy = date(2026, 10, 30)
    r = notificaciones.Registro(
        fecha_ingreso=date(2026, 9, 20), fecha_expedicion=date(2026, 9, 18)
    )
    assert "anterior" in (r.error(hoy) or "")
    futura = notificaciones.Registro(fecha_notificacion=date(2026, 11, 2))
    assert "posterior a hoy" in (futura.error(hoy) or "")
    assert notificaciones.Registro(fecha_ingreso=date(2026, 9, 1)).error(hoy) is None


async def test_flujo_secretaria(
    cliente: Cliente,  # noqa: F811
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app.services import indicadores

    monkeypatch.setattr(indicadores, "hoy_local", lambda: date(2026, 10, 30))
    c = cliente("actas")
    caso_id = await _caso_decidido()
    aplazado = await _caso_decidido(sentido="aplazo")
    de_sede = await _caso_decidido(instancia="sede")
    cab = await _token("secretaria@unal.edu.co", "secretaria")

    datos = c.get(f"{RUTA}?semestre=2026-2S", headers=cab).json()
    ids = [x["caso_id"] for x in datos["casos"]]
    assert ids == [caso_id] and aplazado not in ids and de_sede not in ids
    fila = datos["casos"][0]
    assert fila["estudiante"] == "Persona Ficticia Uno"
    assert fila["sugerida_ingreso"] == "2026-09-10"
    assert fila["registro"] is None

    cuerpo = {
        "acto": "oficio",
        "numero": "Mz.FCEN.1.004-000-2026",
        "fecha_ingreso": "2026-09-18",
        "fecha_expedicion": "2026-09-21",
        "fecha_notificacion": "2026-09-23",
    }
    r = c.put(f"{RUTA}/{caso_id}", json=cuerpo, headers=cab)
    assert r.status_code == 200, r.text
    assert c.put(f"{RUTA}/{aplazado}", json=cuerpo, headers=cab).status_code == 404
    malo = {**cuerpo, "fecha_notificacion": "2026-09-01"}
    assert c.put(f"{RUTA}/{caso_id}", json=malo, headers=cab).status_code == 422

    datos = c.get(f"{RUTA}?semestre=2026-2S", headers=cab).json()
    assert datos["resumen"]["completos"] == 1
    assert datos["resumen"]["ingreso_expedicion"]["mediana"] == 3
    assert datos["resumen"]["total"]["p90"] == 5

    r = c.get(f"{RUTA}/xlsx?semestre=2026-2S", headers=cab)
    assert r.status_code == 200
    libro = load_workbook(io.BytesIO(r.content))
    ws = libro["Casos"]
    valores = [ws.cell(2, j).value for j in range(1, 15)]
    assert valores[0] == caso_id and valores[7] == "Mz.FCEN.1.004-000-2026"
    assert valores[11:] == [3, 2, 5]
    # Sin nombres en el Excel.
    for hoja in libro.worksheets:
        for fila_x in hoja.iter_rows(values_only=True):
            assert "Persona Ficticia Uno" not in [str(v) for v in fila_x]


async def test_permisos(cliente: Cliente) -> None:  # noqa: F811
    c = cliente("actas")
    caso_id = await _caso_decidido()
    adm = await _token("adm@unal.edu.co", "administrativa", "fisica")
    dec = await _token("dec@unal.edu.co", "decanatura")
    assert c.get(RUTA, headers=adm).status_code == 403
    assert c.put(f"{RUTA}/{caso_id}", json={}, headers=dec).status_code == 403
    assert c.get(f"{RUTA}/xlsx", headers=dec).status_code == 200
    assert c.get(RUTA).status_code == 401
    assert settings.db_path.exists()


async def test_secretaria_descarga_el_borrador(cliente: Cliente) -> None:  # noqa: F811
    c = cliente("actas")
    caso_id = await _caso_decidido()
    aplazado = await _caso_decidido(sentido="aplazo")
    cab = await _token("secretaria@unal.edu.co", "secretaria")
    r = c.get(f"/api/web/actas/secretaria/casos/{caso_id}/oficio-consejo", headers=cab)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"
    ruta = f"/api/web/actas/secretaria/casos/{aplazado}/oficio-consejo"
    assert c.get(ruta, headers=cab).status_code == 409
    adm = await _token("adm2@unal.edu.co", "administrativa", "fisica")
    ruta = f"/api/web/actas/secretaria/casos/{caso_id}/oficio-consejo"
    assert c.get(ruta, headers=adm).status_code == 403
