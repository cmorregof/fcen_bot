"""Revisor de soportes con una carpeta FICTICIA y el modelo simulado."""

from __future__ import annotations

import io
import json
import re
import zipfile
from pathlib import Path
from typing import Any

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image
from pydantic import BaseModel
from pypdf import PdfWriter

from app.servicios.calidad.soportes import contenido
from app.servicios.calidad.soportes.cliente import Respuesta
from app.servicios.calidad.soportes.describir import Cache
from app.servicios.calidad.soportes.inventario import accion_de_ruta, inventariar
from app.servicios.calidad.soportes.modelos import (
    Accion,
    Archivo,
    Cruce,
    CruceAccion,
    Veredicto,
)
from app.servicios.calidad.soportes.plan_acciones import leer_acciones
from app.servicios.calidad.soportes.revisor import estimar, revisar
from app.servicios.calidad.soportes.salidas import ENCABEZADO_BORRADOR, citas_a_nombres
from app.servicios.calidad.soportes.verificador import verificar_cruce

SOPORTES = "Soporte de las acciones"
PERSONA_FICTICIA = "Zoila Ficticia Pérez"
CORREO_FICTICIO = "zoila.ficticia@ejemplo.test"


def pdf_ficticio(paginas: int) -> bytes:
    escritor = PdfWriter()
    for _ in range(paginas):
        escritor.add_blank_page(width=200, height=200)
    salida = io.BytesIO()
    escritor.write(salida)
    return salida.getvalue()


def imagen_ficticia() -> bytes:
    salida = io.BytesIO()
    Image.new("RGB", (3000, 1000), "white").save(salida, format="PNG")
    return salida.getvalue()


def xlsx_con_personas() -> bytes:
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.title = "Asistentes"
    ws.append(["Nombre", "Correo", "Semestre"])
    ws.append([PERSONA_FICTICIA, CORREO_FICTICIO, 3])
    ws.append(["Otro Ficticio", "otro@ejemplo.test", 2])
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def plan_ficticio(ruta: Path) -> None:
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.append([])
    ws.append(
        [None, "Aspectos a mejorar", "Objetivo", "Acciones", "Responsable Nombre y correo",
         "Origen de los recursos", "Fechas Ejecución", None, "Indicadores de cumplimiento"]
    )  # fmt: skip
    ws.append([None, None, None, None, None, None, "Inicio", "Fin", None])
    ws.append(
        [None, "Aspecto uno", "Objetivo uno", "Hacer seminarios", CORREO_FICTICIO, "Humanos",
         "2020-01-01", "2021-12-31", "Número de seminarios por semestre"]
    )  # fmt: skip
    ws.append(
        [
            None,
            None,
            None,
            "Difundir convocatorias",
            None,
            None,
            None,
            None,
            "Correos enviados",
        ]
    )
    ws.append(
        [None, "Aspecto dos", "Objetivo dos", "Escribir documento", None, None,
         "2020-01-01", "2020-06-30", "Documento aprobado"]
    )  # fmt: skip
    libro.save(ruta)


@pytest.fixture
def carpeta(tmp_path: Path) -> Path:
    raiz = tmp_path / "plan_ficticio"
    a11 = raiz / SOPORTES / "ACCION 1.1"
    a12 = raiz / SOPORTES / "ACCIÓN 1.2" / "subcarpeta"
    a11.mkdir(parents=True)
    a12.mkdir(parents=True)
    (raiz / SOPORTES / "ACCION 2.1").mkdir(parents=True)
    (a11 / "seminario_ficticio.pdf").write_bytes(pdf_ficticio(7))
    (a11 / "foto_ficticia.png").write_bytes(imagen_ficticia())
    (a11 / "asistentes_ficticios.xlsx").write_bytes(xlsx_con_personas())
    (a11 / "desktop.ini").write_text("[.ShellClassInfo]")
    (a12 / "convocatoria_ficticia.pdf").write_bytes(pdf_ficticio(1))
    with zipfile.ZipFile(a12 / "correos_ficticios.zip", "w") as zf:
        zf.writestr("correos/uno.pdf", pdf_ficticio(2))
        zf.writestr("correos/Thumbs.db", b"x")
    (raiz / "PLAN DE ACCION 2016.xlsx").write_bytes(xlsx_con_personas())
    plan_ficticio(raiz / "plan.xlsx")
    return raiz


class ModeloFalso:
    """Responde según la herramienta; registra lo que se le envió."""

    def __init__(self) -> None:
        self.llamadas: list[dict[str, Any]] = []

    async def llamar(
        self,
        *,
        modelo: str,
        sistema: str,
        bloques: list[dict[str, Any]],
        texto: str,
        salida: type[BaseModel],
        nombre_herramienta: str,
        esfuerzo: str,
        max_tokens: int,
    ) -> Respuesta:
        self.llamadas.append(
            {
                "modelo": modelo,
                "bloques": bloques,
                "texto": texto,
                "herramienta": nombre_herramienta,
            }
        )
        if nombre_herramienta == "registrar_descripcion":
            return Respuesta(
                {
                    "que_es": "Documento ficticio",
                    "fecha_documento": "2020-05",
                    "que_prueba": "Un seminario",
                    "ilegible_o_incompleto": False,
                    "observacion": "",
                },
                1000,
                100,
            )
        if nombre_herramienta == "registrar_cruce":
            propios = texto.split("Documentos generales")[0]
            ids = [
                linea.split("]")[0][1:]
                for linea in propios.splitlines()
                if linea.startswith("[S")
            ]
            if "Acción 1.1" in texto:
                return Respuesta(
                    {
                        "veredicto": "parcial",
                        "justificacion": f"Hay un seminario [{ids[0]}] y uno inventado [S999].",
                        "falta": "Actas de los demás semestres.",
                        "avance_propuesto": 50,
                        "citas": [ids[0], "S999"],
                        "parrafo_cierre": f"Se realizaron seminarios [{ids[0]}].",
                    },
                    3000,
                    800,
                )
            veredicto = "evidencia suficiente" if ids else "sin evidencia"
            return Respuesta(
                {
                    "veredicto": veredicto,
                    "justificacion": "Según los soportes.",
                    "falta": "",
                    "avance_propuesto": 100 if ids else 0,
                    "citas": ids[:1],
                    "parrafo_cierre": f"Cierre [{ids[0]}]." if ids else "Sin soportes.",
                },
                3000,
                800,
            )
        if nombre_herramienta == "registrar_sugerencias":
            codigos = re.findall(r"^## Acción (\d+\.\d+)\.", texto, re.MULTILINE)
            return Respuesta(
                {
                    "sugerencias": [
                        {
                            "accion": c,
                            "donde_buscar": "Actas del comité de esos semestres.",
                            "accion_proximo_plan": "Registrar cada seminario en acta.",
                        }
                        for c in codigos
                    ]
                },
                2000,
                400,
            )
        return Respuesta({"resumen": "Resumen ficticio [S001]."}, 2000, 400)


# --- Inventario --------------------------------------------------------------


def test_asignacion_por_ruta() -> None:
    assert accion_de_ruta("Soportes/ACCION 1.3/x.pdf") == "1.3"
    assert accion_de_ruta("Soportes/ACCIÓN 4.5/sub/x.pdf") == "4.5"
    assert accion_de_ruta("Soportes/accion 2.1/x.pdf") == "2.1"
    # Un archivo cuyo nombre dice «ACCION» no es una carpeta de acción.
    assert accion_de_ruta("Plan/CRONOGRAMA PLAN DE ACCION 2016 MMA.xlsx") == ""
    assert accion_de_ruta("Plan/Soporte de las acciones- Plan 2016/x.pdf") == ""


def test_inventario_ficticio(carpeta: Path) -> None:
    archivos = inventariar(carpeta)
    nombres = {a.nombre: a for a in archivos}
    assert "desktop.ini" not in nombres and "Thumbs.db" not in nombres
    assert nombres["seminario_ficticio.pdf"].accion == "1.1"
    assert nombres["seminario_ficticio.pdf"].paginas == 7
    assert nombres["foto_ficticia.png"].tipo == "imagen"
    assert nombres["convocatoria_ficticia.pdf"].accion == "1.2"
    assert nombres["PLAN DE ACCION 2016.xlsx"].accion == ""
    zip_ = nombres["correos_ficticios.zip"]
    assert zip_.miembros == ["correos/uno.pdf"]
    miembro = nombres["uno.pdf"]
    assert miembro.dentro_de_zip == zip_.id and miembro.accion == "1.2"
    assert [a.id for a in archivos] == [
        f"S{i:03d}" for i in range(1, len(archivos) + 1)
    ]
    assert all(a.sha256 for a in archivos)


def test_lector_del_plan(carpeta: Path) -> None:
    acciones = leer_acciones(carpeta / "plan.xlsx")
    assert [a.codigo for a in acciones] == ["1.1", "1.2", "2.1"]
    assert acciones[1].aspecto == "Aspecto uno"
    assert acciones[1].fin == "2021-12-31"
    assert acciones[2].indicador == "Documento aprobado"
    # La columna de responsables no se lee.
    assert all(CORREO_FICTICIO not in str(a) for a in acciones)


# --- Minimización -------------------------------------------------------------


def test_xlsx_solo_hojas_encabezados_y_conteo() -> None:
    preparado = contenido.xlsx(xlsx_con_personas())
    assert preparado.bloques == []
    assert "Asistentes" in preparado.texto
    assert "Nombre" in preparado.texto and "Semestre" in preparado.texto
    assert "2 filas" in preparado.texto
    assert PERSONA_FICTICIA not in preparado.texto
    assert "ejemplo.test" not in preparado.texto


def test_pdf_largo_se_recorta() -> None:
    preparado = contenido.pdf(pdf_ficticio(9))
    assert preparado.paginas_enviadas == contenido.PAGINAS_PDF
    assert "9 páginas" in preparado.nota
    assert preparado.bloques[0]["type"] == "document"


def test_correos_se_ocultan() -> None:
    assert (
        contenido.sin_correos(f"escribir a {CORREO_FICTICIO} ya")
        == "escribir a [correo] ya"
    )


# --- Verificador ---------------------------------------------------------------


def _archivo(id_: str, accion: str) -> Archivo:
    return Archivo(id_, f"r/{id_}.pdf", f"{id_}.pdf", accion, "pdf", ".pdf", 10)


def test_verificador_marca_y_no_borra() -> None:
    inventario = {"S001": _archivo("S001", "1.1"), "S002": _archivo("S002", "")}
    accion = Accion("1.1", "A", "O", "Hacer", "Indicador", "2020", "2021")
    propuesta = CruceAccion(
        veredicto=Veredicto.SUFICIENTE,
        justificacion="Ver [S001], [S002] y [S404].",
        falta="",
        avance_propuesto=80,
        citas=["S001", "S404"],
        parrafo_cierre="Cierre sin citas.",
    )
    cruce = Cruce(accion, propuesta)
    verificar_cruce(cruce, inventario)
    estados = {e.cita: e for e in cruce.verificacion}
    assert estados["S001"].valida and estados["S001"].nota == ""
    assert (
        estados["S002"].valida and estados["S002"].nota == "documento general del plan"
    )
    assert not estados["S404"].valida
    # La propuesta queda intacta: la cita mala se marca, no se borra.
    assert "S404" in cruce.propuesta.citas  # type: ignore[union-attr]
    alertas = " ".join(cruce.alertas)
    assert "S404" in alertas
    assert "menor que 100" in alertas
    assert "no cita ningún soporte" in alertas


def test_citas_a_nombres() -> None:
    inventario = {"S001": _archivo("S001", "1.1")}
    texto = citas_a_nombres("Ver [S001] y [S777].", inventario)
    assert texto == "Ver [S001 · S001.pdf] y [S777 ⚠ cita no verificada]."


def test_avance_se_acota() -> None:
    p = CruceAccion(
        veredicto=Veredicto.PARCIAL, justificacion="", falta="", avance_propuesto=140,
        citas=[], parrafo_cierre="",
    )  # fmt: skip
    assert p.avance_propuesto == 100


# --- Corrida completa, caché y salidas -------------------------------------------


async def test_corrida_completa_cache_y_salidas(carpeta: Path, tmp_path: Path) -> None:
    salidas = tmp_path / "salidas"
    modelo = ModeloFalso()
    informe = await revisar(
        carpeta,
        carpeta / "plan.xlsx",
        salidas,
        modelo,
        "Programa Ficticio",
        concurrencia=2,
    )
    # El zip no va al modelo; sus archivos internos sí.
    descritos = [
        c for c in modelo.llamadas if c["herramienta"] == "registrar_descripcion"
    ]
    no_zip = [a for a in informe.archivos if a.tipo != "zip"]
    assert len(descritos) == len(no_zip)
    # Ninguna fila con personas del xlsx llegó al modelo.
    todo = json.dumps(modelo.llamadas, ensure_ascii=False)
    assert PERSONA_FICTICIA not in todo and CORREO_FICTICIO not in todo
    # Los PDF largos se mandan recortados.
    largo = next(c for c in descritos if "seminario_ficticio.pdf" in c["texto"])
    assert "PDF de 7 páginas; se envían las primeras 4" in largo["texto"]
    assert largo["bloques"][0]["type"] == "document"

    assert informe.veredictos() == {
        "parcial": 1, "evidencia suficiente": 1, "sin evidencia": 1,
    }  # fmt: skip
    c11 = next(c for c in informe.cruces if c.accion.codigo == "1.1")
    assert any(not e.valida and e.cita == "S999" for e in c11.verificacion)
    assert informe.uso.costo() > 0

    # Las tres salidas.
    inv = load_workbook(informe.salidas["inventario"])["Inventario"]
    encabezado = [c.value for c in inv[1]]
    assert "¿Correcto? (Sí/No)" in encabezado and "Corrección" in encabezado
    assert inv.max_row == len(informe.archivos) + 1
    cru = load_workbook(informe.salidas["cruce"])["Cruce por acción"]
    encabezado = [c.value for c in cru[1]]
    assert "Aprobado" in encabezado and "Ajuste" in encabezado
    assert cru.max_row == 4
    doc = Document(str(informe.salidas["borrador"]))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert ENCABEZADO_BORRADOR in texto
    assert "Citas que no están en el inventario: S999" in texto
    assert "Acción 1.1" in texto and "Acción 2.1" in texto
    # Sugerencias: una sola llamada, solo para las acciones sin evidencia.
    sug = [c for c in modelo.llamadas if c["herramienta"] == "registrar_sugerencias"]
    assert len(sug) == 1 and "## Acción 2.1." in sug[0]["texto"]
    assert "## Acción 1.1." not in sug[0]["texto"]
    c21 = next(c for c in informe.cruces if c.accion.codigo == "2.1")
    assert c21.sugerencia is not None
    assert "Sugerencias para las acciones sin evidencia" in texto
    assert informe.salidas["reporte_sin_evidencia"].exists()
    assert doc.sections[0].header.paragraphs[0].text == ENCABEZADO_BORRADOR

    # Repetir la corrida no llama al modelo: todo sale de la caché.
    otra = ModeloFalso()
    informe2 = await revisar(
        carpeta, carpeta / "plan.xlsx", salidas, otra, "Programa Ficticio"
    )
    assert otra.llamadas == []
    assert informe2.uso.costo() == 0
    assert all(d.desde_cache for d in informe2.descripciones.values() if d.datos and
               informe2.archivos[int(d.archivo_id[1:]) - 1].tipo != "zip")  # fmt: skip


def test_estimacion_descuenta_cache(carpeta: Path, tmp_path: Path) -> None:
    archivos = inventariar(carpeta)
    acciones = leer_acciones(carpeta / "plan.xlsx")
    cache = Cache(tmp_path / "cache")
    antes = estimar(archivos, acciones, cache)
    assert antes.archivos_al_modelo == len([a for a in archivos if a.tipo != "zip"])
    assert antes.costo > 0


async def test_sin_creditos_detiene_la_corrida() -> None:
    from app.servicios.calidad.soportes.cliente import ClienteAnthropic
    from app.servicios.calidad.soportes.modelos import DescripcionArchivo

    intentos = [0]

    class Mensajes:
        async def create(self, **_: Any) -> Any:
            intentos[0] += 1
            raise RuntimeError("Your credit balance is too low to access the API.")

    class Beta:
        messages = Mensajes()

    class Falso:
        beta = Beta()

    cliente = ClienteAnthropic("clave-ficticia")
    cliente._cliente = Falso()  # type: ignore[assignment]
    kwargs: dict[str, Any] = {
        "modelo": "m", "sistema": "s", "bloques": [], "texto": "t",
        "salida": DescripcionArchivo, "nombre_herramienta": "h",
        "esfuerzo": "low", "max_tokens": 10,
    }  # fmt: skip
    primera = await cliente.llamar(**kwargs)
    segunda = await cliente.llamar(**kwargs)
    assert primera.error == "no quedan créditos en la cuenta de la API"
    assert segunda.error.startswith("corrida detenida")
    assert intentos[0] == 1


def test_parrafo_con_formato_quita_marcas_markdown() -> None:
    from docx import Document

    from app.servicios.calidad.soportes.salidas import parrafo_con_formato

    doc = Document()
    parrafo_con_formato(doc, "- **Becas (2.3):** listado [S001].")
    p = doc.paragraphs[-1]
    assert p.text == "Becas (2.3): listado [S001]."
    assert p.style.name == "List Bullet"
    assert p.runs[0].bold and p.runs[0].text == "Becas (2.3):"


def test_conteos_xlsx_por_periodo_sin_nombres() -> None:
    from openpyxl import Workbook

    from app.servicios.calidad.soportes.contenido import conteos_xlsx, xlsx

    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.append(["Periodo Académico", "Nombre", "Fecha terminación"])
    ws.append(["2016 - 2S", PERSONA_FICTICIA, "2030-01-01"])
    ws.append(["2016 - 2S", "Otra Persona Ficticia", "2030-01-01"])
    ws.append(["2017-1S", "Tercera Persona Ficticia", "2030-01-01"])
    salida = io.BytesIO()
    libro.save(salida)
    conteos = conteos_xlsx(salida.getvalue())
    assert "2016-2S: 2" in conteos and "2017-1S: 1" in conteos
    assert PERSONA_FICTICIA not in conteos
    enviado = xlsx(salida.getvalue()).texto
    assert "2016-2S: 2" in enviado and PERSONA_FICTICIA not in enviado
