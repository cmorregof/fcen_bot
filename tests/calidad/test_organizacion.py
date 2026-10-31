"""Organización de evidencias según el instructivo de la Dirección Académica.

Nombres y archivos ficticios.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.servicios.calidad.soportes.inventario import inventariar
from app.servicios.calidad.soportes.modelos import Archivo
from app.servicios.calidad.soportes.organizacion import (
    carpeta_de,
    nombre_sugerido,
    revisar,
    revisar_nombre,
)
from app.servicios.calidad.soportes.salidas import organizacion_xlsx

BASE = "Soportes de acciones/Acción 1.1"


def _archivo(ruta: str, *, sha: str = "", fecha: str | None = None) -> Archivo:
    nombre = ruta.rsplit("/", 1)[-1]
    return Archivo(
        id="S001",
        ruta=ruta,
        nombre=nombre,
        accion="1.1",
        tipo="pdf",
        extension="." + nombre.rsplit(".", 1)[-1],
        tamano=10,
        sha256=sha,
        fecha=fecha,
    )


def test_nombres_correctos_del_instructivo() -> None:
    for nombre in (
        "01 Correo estudiantes PEP 20251001.pdf",
        "02 Invitacion reunion PEP 20251005.pdf",
        "06 Acta reunion 20251020.pdf",
    ):
        assert revisar_nombre(nombre) == [], nombre


def test_nombres_que_se_deben_evitar() -> None:
    obs = revisar_nombre(
        "1. Correo para estudiantes y profesores ¿Sabes qué es el PEP?.pdf"
    )
    texto = " ".join(obs)
    assert "caracteres (máximo 50)" in texto
    assert "consecutivo" in texto
    assert "AAAAMMDD" in texto
    assert "tildes" in texto and "especiales" in texto
    assert any("especiales" in o for o in revisar_nombre("Foto reunión!!!.jpg"))
    assert any("especiales" in o for o in revisar_nombre("Soporte%avance#1.xlsx"))
    assert any("no es válida" in o for o in revisar_nombre("01 Acta 20251340.pdf"))


def test_nombre_sugerido_sin_tildes_ni_signos_y_con_fecha() -> None:
    a = _archivo(
        f"{BASE}/Soportes de producto final/Difusión en pantallas octubre 2025.jpeg",
        fecha="2025-10-10",
    )
    assert nombre_sugerido(a, 3) == "03 Difusion en pantallas octubre 20251010.jpeg"
    b = _archivo(f"{BASE}/Soportes de producto final/Documento_final_v3 (1).docx")
    sugerido = nombre_sugerido(b, 1)
    assert sugerido == "01 Documento final v3 AAAAMMDD.docx"
    largo = _archivo(
        f"{BASE}/Soportes de producto final/{'Pieza comunicativa difundida ' * 4}20251012.png"
    )
    s = nombre_sugerido(largo, 4)
    assert len(s) <= 50 and s.endswith(" 20251012.png") and revisar_nombre(s) == []


def test_carpetas_dentro_de_la_accion() -> None:
    assert (
        carpeta_de(_archivo(f"{BASE}/Producto final/Informe.pdf")) == "producto final"
    )
    assert (
        carpeta_de(_archivo(f"{BASE}/Soportes de producto final/x.pdf")) == "soportes"
    )
    assert carpeta_de(_archivo(f"{BASE}/x.pdf")) == "suelto"
    zipeado = _archivo("Acción 1.1.zip!/Soportes de producto final/x.pdf")
    assert carpeta_de(zipeado) == "soportes"


def test_estructura_orden_y_duplicados() -> None:
    archivos = [
        _archivo(f"{BASE}/Producto final/Informe.pdf"),
        _archivo(f"{BASE}/Producto final/Anexo.docx"),
        _archivo(f"{BASE}/Soportes de producto final/01 Acta 20251020.pdf", sha="a"),
        _archivo(f"{BASE}/Soportes de producto final/02 Correo 20251001.pdf", sha="b"),
        _archivo(f"{BASE}/Soportes de producto final/03 Copia 20251021.pdf", sha="a"),
        _archivo(f"{BASE}/suelto.pdf"),
        _archivo(
            "Soportes de acciones/Acción 1.2/Soportes de producto final/01 Acta 20251001.pdf"
        ).__class__(
            id="S009",
            ruta="Soportes de acciones/Acción 1.2/Producto final/Final.pdf",
            nombre="Final.pdf",
            accion="1.2",
            tipo="pdf",
            extension=".pdf",
            tamano=1,
        ),
    ]
    org = revisar(archivos)
    assert any("uno solo" in o for o in org.acciones["1.1"])
    assert any("Soportes de producto final" in o for o in org.acciones["1.2"])
    por_nombre = {r.archivo.nombre: r for r in org.archivos}
    assert "el producto final debe ser un PDF" in por_nombre["Anexo.docx"].observaciones
    assert any(
        "orden cronológico" in o
        for o in por_nombre["02 Correo 20251001.pdf"].observaciones
    )
    assert any(
        "es igual a 01 Acta" in o
        for o in por_nombre["03 Copia 20251021.pdf"].observaciones
    )
    assert any("suelto" in o for o in por_nombre["suelto.pdf"].observaciones)
    # Renumerado por fecha: el correo del 1 de octubre pasa a ser el 01.
    assert por_nombre["02 Correo 20251001.pdf"].sugerido == "01 Correo 20251001.pdf"
    assert org.resumen()["acciones_con_observaciones"] == ["1.1", "1.2"]


def test_carpeta_bien_organizada_no_tiene_observaciones(tmp_path: Path) -> None:
    accion = tmp_path / "Soportes de acciones" / "Acción 2.1"
    (accion / "Producto final").mkdir(parents=True)
    (accion / "Soportes de producto final").mkdir()
    (accion / "Producto final" / "Informe accion 2 1.pdf").write_bytes(
        b"%PDF-1.4 ficticio"
    )
    for nombre in ("01 Acta reunion 20251001.pdf", "02 Correo PEP 20251005.pdf"):
        (accion / "Soportes de producto final" / nombre).write_bytes(nombre.encode())
    org = revisar(inventariar(tmp_path))
    assert org.acciones == {"2.1": []}
    assert org.con_observaciones == 0

    salida = tmp_path / "salidas" / "organizacion.xlsx"
    organizacion_xlsx(salida, org)
    libro = load_workbook(salida)
    assert libro.sheetnames == ["Acciones", "Archivos", "Léeme"]
    assert libro["Acciones"]["B2"].value == "Sin observaciones."
    assert libro["Archivos"].max_row == 4


def test_indice_electronico_sigue_el_formato(tmp_path: Path) -> None:
    from app.servicios.calidad.soportes.salidas import indice_electronico_xlsx

    archivos = [
        _archivo(f"{BASE}/Soportes de producto final/01 Acta 20251001.pdf"),
        _archivo(f"{BASE}/Producto final/Informe accion.pdf", fecha="2025-11-02"),
        Archivo(
            id="S003",
            ruta="Plan general.xlsx",
            nombre="Plan general.xlsx",
            accion="",
            tipo="xlsx",
            extension=".xlsx",
            tamano=290_000,
        ),
    ]
    archivos[0].paginas, archivos[1].paginas = 2, 5
    salida = tmp_path / "indice.xlsx"
    indice_electronico_xlsx(salida, revisar(archivos))
    ws = load_workbook(salida)["Formato"]
    assert ws["C9"].value == "PLANES DE MEJORAMIENTO DE PROGRAMAS CURRICULARES"
    assert [ws.cell(12, j).value for j in range(2, 13)][:4] == [
        "Nombre Documento",
        "Fecha Creación Documento",
        "Fecha Incorporación Expediente",
        "Orden Documento",
    ]
    # Producto final primero, luego soportes, luego lo general.
    assert [ws[f"B{n}"].value for n in (13, 14, 15)] == [
        "Informe accion",
        "01 Acta 20251001",
        "Plan general",
    ]
    assert ws["C13"].value == "02/11/2025" and ws["E13"].value == 1
    assert ws["G13"].value == '=IF(F13=0,"0","1")'
    assert ws["G14"].value == '=IF(F14=0,"0",(1+H13))'
    assert ws["H14"].value == "=G14+(F14-1)"
    assert ws["I15"].value == "excel" and ws["J15"].value == "283 kb"
    assert ws["K13"].value == "Electrónico"
    assert ws["L13"].value == "Acción 1.1 — Producto final"
    assert ws["L15"].value == "Documento general del plan"
    # Lo que no se sabe queda en amarillo: páginas del Excel.
    assert ws["F15"].fill.fgColor.rgb.endswith("FFFF00")
