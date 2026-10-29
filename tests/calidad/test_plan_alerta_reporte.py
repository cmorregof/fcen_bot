from __future__ import annotations

import io
from collections.abc import Callable
from datetime import date

import aiosqlite
from openpyxl import load_workbook

from app.servicios.calidad import alertas, plan
from tests.calidad.conftest import HOY, Chat_, fila, xlsx
from tests.conftest import TelegramFalso


async def _cargar(chat: Chat_, servir: Callable[[bytes], None]) -> None:
    servir(
        xlsx(
            [
                fila(codigo="UNO-01", limite=date(2026, 9, 1)),  # vencida
                fila(codigo="UNO-02", limite=date(2026, 11, 10)),  # próxima
                fila(codigo="UNO-03", limite=date(2027, 3, 1)),  # pendiente
                fila(
                    codigo="UNO-04", limite=date(2026, 8, 1)
                ),  # vencida, luego cumplida
                fila(
                    programa="Programa Dos", codigo="DOS-01", limite=date(2026, 10, 1)
                ),
            ]
        )
    )
    await chat.documento("plan.xlsx", caption="/importar_plan")
    await chat.boton("cal:imp:ok")


async def test_plan_agrupa_por_estado_con_responsable(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    await chat.texto("/evidencia UNO-04 https://ejemplo.invalid/doc")
    texto = await chat.texto("/plan programa uno")
    assert "Vencidas (1)" in texto and "Próximas a vencer (4 semanas) (1)" in texto
    assert "Pendientes (1)" in texto and "Cumplidas (1)" in texto
    assert "Cargo de prueba — Dependencia de prueba" in texto
    assert "DOS-01" not in texto


async def test_evidencia_guarda_quien_y_cuando(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    respuesta = await chat.texto("/evidencia uno-01 https://ejemplo.invalid/informe")
    assert "marcada como cumplida" in respuesta
    tarea = (await plan.buscar_tareas_por_codigo(db, "UNO-01"))[0]
    ev = tarea.evidencias[0]
    assert tarea.cumplida and ev.tipo == "enlace"
    assert ev.registrado_por == 777 and ev.registrado_nombre == "Persona Ficticia"
    assert ev.creado.startswith("2026-10-28T09:00")


async def test_evidencia_en_dos_pasos_con_archivo(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    assert "Envíame el archivo" in await chat.texto("/evidencia UNO-02")
    assert "marcada como cumplida" in await chat.documento("informe.pdf")
    tarea = (await plan.buscar_tareas_por_codigo(db, "UNO-02"))[0]
    assert (
        tarea.evidencias[0].tipo == "archivo"
        and tarea.evidencias[0].valor == "informe.pdf"
    )


async def test_evidencia_rechaza_texto_que_no_es_enlace(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    assert "enlace" in await chat.texto("/evidencia UNO-01 listo")
    assert "No hay una tarea" in await chat.texto(
        "/evidencia NOEXISTE https://x.invalid"
    )
    assert not (await plan.buscar_tareas_por_codigo(db, "UNO-01"))[0].cumplida


async def test_alerta_incluye_vencidas_y_proximas_sin_evidencia(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    await chat.texto("/evidencia UNO-04 https://ejemplo.invalid/doc")
    mapa = await alertas.construir_alertas(db, HOY)
    assert set(mapa) == {"Programa Uno", "Programa Dos"}
    uno = mapa["Programa Uno"]
    assert "UNO-01" in uno and "UNO-02" in uno
    assert "UNO-03" not in uno and "UNO-04" not in uno  # pendiente y cumplida
    assert "Cargo de prueba — Dependencia de prueba" in uno
    assert "No reemplazo el seguimiento de la DAMA" in uno


async def test_vista_previa_avisa_que_no_se_envia(
    chat: Chat_, tg: TelegramFalso, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    texto = await chat.texto("/alerta_vista_previa")
    assert "DESACTIVADO" in texto and "DOS-01" in texto


async def test_envio_real_solo_con_alertas_activas(
    chat: Chat_,
    db: aiosqlite.Connection,
    tg: TelegramFalso,
    servir_excel: Callable[[bytes], None],
    monkeypatch: object,
) -> None:
    from app.servicios.calidad.config import calidad_settings

    await _cargar(chat, servir_excel)
    antes = len(tg.mensajes)
    assert await alertas.enviar_alertas(db, HOY) == {"activas": False, "enviados": 0}
    assert len(tg.mensajes) == antes

    monkeypatch.setattr(calidad_settings, "alertas_activas", True)  # type: ignore[attr-defined]
    assert await alertas.enviar_alertas(db, HOY) == {"activas": True, "enviados": 2}
    assert len(tg.mensajes) == antes + 2


async def test_reporte_tiene_estado_y_fuentes(
    chat: Chat_, tg: TelegramFalso, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    await chat.texto("/evidencia UNO-04 https://ejemplo.invalid/doc")
    await chat.texto("/reporte Programa Uno")
    libro = load_workbook(io.BytesIO(tg.documentos[-1]["contenido"]))
    assert libro.sheetnames == ["estado", "tareas", "fuentes"]
    cifras = {
        r[0]: r[1] for r in libro["estado"].iter_rows(min_row=2, values_only=True)
    }
    assert cifras["Total de tareas"] == 4
    assert cifras["Tareas vencidas"] == 1 and cifras["Tareas cumplidas"] == 1
    fuentes = list(libro["fuentes"].iter_rows(min_row=2, values_only=True))
    de_cumplidas = [f for f in fuentes if f[0] == "Tareas cumplidas"]
    assert len(de_cumplidas) == 1 and de_cumplidas[0][1] == "UNO-04"
    assert de_cumplidas[0][7] == "https://ejemplo.invalid/doc"
    # Cada cifra tiene tantos registros como su valor.
    for nombre, valor in cifras.items():
        assert len([f for f in fuentes if f[0] == nombre]) == valor
