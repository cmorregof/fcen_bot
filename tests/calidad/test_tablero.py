from __future__ import annotations

import io
from collections.abc import Callable

import aiosqlite
import pytest
from openpyxl import load_workbook

from app.servicios.calidad import alertas, plan
from app.servicios.calidad.config import calidad_settings
from tests.calidad.conftest import ALFA, BETA, HOY, Chat_, plantilla
from tests.conftest import TelegramFalso


async def _cargar(chat: Chat_, servir: Callable[[bytes], None]) -> None:
    for ruta in (ALFA, BETA):
        servir(plantilla(ruta))
        await chat.documento(ruta.name, caption="/importar_plan")
        await chat.boton("cal:imp:ok")


async def test_tablero_del_programa(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    texto = await chat.texto("/plan alfa")
    assert "<b>Plan de mejoramiento — Ejemplo Alfa (ficticio)</b>" in texto
    assert "7 acciones · avance promedio 54 % · 5 tareas en el cronograma (2025, 2026)" in texto
    assert "🔴 2 vencidas" in texto and "🟢 2 cumplidas" in texto
    assert "🔁 3 reprogramadas" in texto
    assert "Próximo corte: 2026-11-21 (estimado" in texto
    # Cada acción con su % de avance.
    assert "🔵 1.1.1 60 %" in texto and "🟢 1.2.2 100 %" in texto
    # Vencidas: fecha de fin pasada y avance < 100 %, con responsables.
    vencidas = texto.split("<b>🔴 Vencidas (2)</b>")[1].split("<b>🔁")[0]
    assert "<b>1.1.2</b> 50 %" in vencidas and "<b>2.1.1</b> 30 %" in vencidas
    assert "Responsables: Comité Asesor del programa (ejemplo)" in vencidas
    reprogramadas = texto.split("<b>🔁 Reprogramadas (3)</b>")[1]
    assert "<b>1.2.2</b> fin anterior 2025-02-01" in reprogramadas
    assert "«Tercera reprogramación»" in reprogramadas
    assert "<b>🔵 En curso (2)</b>: 1.1.1, 1.1.3" in texto
    assert "<b>⚪ Por iniciar (1)</b>: 2.1.2" in texto
    assert "<b>🟢 Cumplidas (2)</b>: 1.2.1, 1.2.2" in texto
    assert "<b>Responsables (acciones abiertas)</b>" in texto
    assert "• Dirección del programa (ejemplo): 1.1.1, 2.1.1" in texto
    assert "Beta" not in texto


async def test_tablero_muestra_las_marcadas_no_funciono(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    texto = await chat.texto("/plan beta")
    assert "<b>⛔ Marcadas «NO FUNCIONÓ» (1)</b>" in texto
    assert "<b>1.1</b> 20 %" in texto.split("NO FUNCIONÓ» (1)</b>")[1]


async def test_plan_sin_programa_lista_los_cargados(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    texto = await chat.texto("/plan")
    assert "Ejemplo Alfa (ficticio)" in texto and "Ejemplo Beta (ficticio)" in texto
    assert "No encontré el programa «gamma»" in await chat.texto("/plan gamma")


async def test_evidencia_queda_por_revisar_y_no_cambia_el_avance(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    respuesta = await chat.texto("/evidencia alfa 1.1.2 https://ejemplo.invalid/acta")
    assert "Soporte de la acción 1.1.2 (Ejemplo Alfa (ficticio)) registrado" in respuesta
    assert "Queda por revisar" in respuesta
    guardado = await plan.cargar_plan(db, "Ejemplo Alfa (ficticio)")
    assert guardado is not None
    soporte = guardado.soportes["1.1.2"][0]
    assert soporte.tipo == "enlace" and soporte.registrado_nombre == "Persona Ficticia"
    assert soporte.creado.startswith("2026-10-28T09:00")
    assert guardado.accion("1.1.2").avance == 50  # type: ignore[union-attr]
    assert "Soportes registrados en el bot, por revisar: 1." in await chat.texto("/plan alfa")


async def test_evidencia_en_dos_pasos_con_archivo(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    pedido = await chat.texto("/evidencia beta 2.1")
    assert "Envíame el archivo o el enlace" in pedido
    assert "N° de sesiones del seminario" in pedido  # el indicador de la acción
    assert "registrado" in await chat.documento("listas.pdf")
    guardado = await plan.cargar_plan(db, "Ejemplo Beta (ficticio)")
    assert guardado is not None and guardado.soportes["2.1"][0].valor == "listas.pdf"


async def test_evidencia_rechaza_lo_que_no_sirve(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    assert "Uso:" in await chat.texto("/evidencia alfa")
    assert "enlace" in await chat.texto("/evidencia alfa 1.1.1 listo")
    assert "No hay una acción 9.9" in await chat.texto("/evidencia alfa 9.9 https://x.invalid")
    # Con dos programas cargados hay que decir cuál.
    assert "Uso:" in await chat.texto("/evidencia 1.1 https://x.invalid")


async def test_alerta_por_programa(
    chat: Chat_, db: aiosqlite.Connection, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    mapa = await alertas.construir_alertas(db, HOY)
    assert set(mapa) == {"Ejemplo Alfa (ficticio)", "Ejemplo Beta (ficticio)"}
    alfa = mapa["Ejemplo Alfa (ficticio)"]
    assert "🔴 Vencidas (2)" in alfa
    assert "🟡 Vencen en las próximas 4 semanas (1)" in alfa and "<b>1.1.3</b>" in alfa
    assert "🔁 Reprogramadas sin cumplir (2)" in alfa  # 1.2.2 ya está cumplida
    assert "Próximo corte: 2026-11-21" in alfa
    assert "No reemplazo el seguimiento de la DAMA" in alfa
    assert "⛔ Marcadas «NO FUNCIONÓ» (1)" in mapa["Ejemplo Beta (ficticio)"]


async def test_vista_previa_avisa_que_no_se_envia(
    chat: Chat_, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    texto = await chat.texto("/alerta_vista_previa")
    assert "DESACTIVADO" in texto and "Ejemplo Beta" in texto


async def test_envio_real_solo_con_alertas_activas(
    chat: Chat_,
    db: aiosqlite.Connection,
    tg: TelegramFalso,
    servir_excel: Callable[[bytes], None],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    await _cargar(chat, servir_excel)
    antes = len(tg.mensajes)
    assert await alertas.enviar_alertas(db, HOY) == {"activas": False, "enviados": 0}
    assert len(tg.mensajes) == antes
    monkeypatch.setattr(calidad_settings, "alertas_activas", True)
    assert await alertas.enviar_alertas(db, HOY) == {"activas": True, "enviados": 2}


async def test_reporte_con_pestana_de_fuentes(
    chat: Chat_, tg: TelegramFalso, servir_excel: Callable[[bytes], None]
) -> None:
    await _cargar(chat, servir_excel)
    await chat.texto("/evidencia alfa 1.1.2 https://ejemplo.invalid/acta")
    await chat.texto("/reporte alfa")
    libro = load_workbook(io.BytesIO(tg.documentos[-1]["contenido"]))
    assert libro.sheetnames == ["estado", "acciones", "tareas", "cortes", "fuentes"]
    cifras = {r[0]: r[1] for r in libro["estado"].iter_rows(min_row=2, values_only=True)}
    assert cifras["Acciones"] == 7 and cifras["Acciones vencidas"] == 2
    assert cifras["Acciones reprogramadas"] == 3
    assert cifras["Acciones con soportes por revisar"] == 1
    assert cifras["Avance promedio (%)"] == 54.3
    assert cifras["Próximo corte (estimado)"] == "2026-11-21"
    assert libro["tareas"].max_row == 6 and libro["cortes"].max_row == 6
    fuentes = list(libro["fuentes"].iter_rows(min_row=2, values_only=True))
    # Cada cifra de conteo tiene tantas filas en «fuentes» como su valor.
    for nombre, valor in cifras.items():
        if nombre.startswith("Acciones"):
            assert len([f for f in fuentes if f[0] == nombre]) == valor
    vencida = next(f for f in fuentes if f[0] == "Acciones vencidas" and f[1] == "1.1.2")
    assert vencida[5] == "«2. Plan de Mejora» fila 10"  # origen en el plan
    assert vencida[4] == ALFA.name and vencida[9] == "https://ejemplo.invalid/acta"
