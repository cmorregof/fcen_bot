"""Borrador de avance de un corte: fuentes, citas, nombres y costo. Datos FICTICIOS."""

from __future__ import annotations

import re
from datetime import date
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import load_workbook
from pydantic import BaseModel

from app.servicios.calidad.avance import borrador as bor
from app.servicios.calidad.avance import fuentes, salidas
from app.servicios.calidad.plan import PlanGuardado, Soporte
from app.servicios.calidad.soportes.cliente import Respuesta
from app.servicios.calidad.soportes.describir import Cache
from app.servicios.calidad.soportes.indice import guardar_indice, leer_indice
from app.servicios.calidad.soportes.modelos import (
    Archivo,
    Descripcion,
    DescripcionArchivo,
)
from tests.calidad.conftest import leer

CORTE = date(2026, 11, 30)
PERSONA = "Zoila Ficticia Pérez"


class ModeloFalso:
    """Redacta una frase por fuente nueva y cuenta las llamadas."""

    def __init__(self, frases: list[dict[str, Any]] | None = None) -> None:
        self.textos: list[str] = []
        self.frases = frases

    async def llamar(
        self,
        *,
        modelo: str,
        sistema: str,
        bloques: list[dict[str, Any]],
        texto: str,
        salida: type[BaseModel],
        nombre_herramienta: str,
        esfuerzo: str,
        max_tokens: int,
    ) -> Respuesta:
        self.textos.append(texto)
        if self.frases is not None:
            return Respuesta({"frases": self.frases}, 1000, 200, modelo=modelo)
        nuevas = re.findall(r"^\[(\w+)\] \(nueva en este corte\)", texto, re.M)
        frases = [{"texto": f"Se registró avance [{i}].", "citas": [i]} for i in nuevas]
        return Respuesta({"frases": frases}, 1000, 200, modelo=modelo)


def guardado(responsable: str | None = None) -> PlanGuardado:
    plan = leer()
    if responsable:
        plan.acciones[0].responsables = responsable
        plan.acciones[0].tareas[0].responsable = responsable
    return PlanGuardado(plan=plan, importacion_id=1, archivo="plan.xlsx")


def con_enlace(pg: PlanGuardado, codigo: str, url: str, creado: str) -> None:
    pg.soportes[codigo] = [
        Soporte(1, 1, "enlace", url, None, 0, "x@ejemplo.test", creado)
    ]


def armar(pg: PlanGuardado, **kw: Any) -> list[fuentes.AccionCorte]:
    return fuentes.armar(pg, CORTE, kw.get("soportes", {}), fuentes.nombres_en_plan(pg))


def accion(acciones: list[fuentes.AccionCorte], codigo: str) -> fuentes.AccionCorte:
    return next(a for a in acciones if a.accion.codigo == codigo)


async def correr(
    pg: PlanGuardado, modelo: ModeloFalso, tmp_path: Path, **kw: Any
) -> bor.Borrador:
    acciones = armar(pg, **kw)
    return await bor.generar(
        "Ejemplo",
        CORTE.isoformat(),
        acciones,
        fuentes.nombres_en_plan(pg),
        modelo,
        Cache(tmp_path / "c"),
    )


# ---------------------------------------------------------------------------


def test_apagado_por_defecto() -> None:
    from app.servicios.calidad.config import CalidadSettings

    assert CalidadSettings.model_fields["avance_activo"].default is False
    assert CalidadSettings.model_fields["avance_tope_usd"].default == 2.0


def test_fuentes_con_hoja_y_fila_del_plan() -> None:
    a = accion(armar(guardado()), "1.1.1")
    plan = [f for f in a.fuentes if f.tipo == "plan"]
    assert plan[0].ref == "«2. Plan de Mejora» fila 9"
    assert any(f.ref == "«3. Cronograma 2026» fila 8" for f in plan)
    assert [f.id for f in a.fuentes if f.tipo == "corte"] == ["C1", "C2"]
    # Lo viejo no es evidencia nueva: nada cambió desde el corte de mayo.
    assert not a.hay_evidencia_nueva


def test_enlace_registrado_despues_del_corte_anterior_es_nuevo() -> None:
    pg = guardado()
    con_enlace(pg, "1.1.1", "https://ejemplo.invalid/drive/nuevo", "2026-10-01T10:00")
    con_enlace(pg, "1.2.1", "https://ejemplo.invalid/drive/viejo", "2026-01-01T10:00")
    acciones = armar(pg)
    assert accion(acciones, "1.1.1").hay_evidencia_nueva
    assert not accion(acciones, "1.2.1").hay_evidencia_nueva
    enlace = next(f for f in accion(acciones, "1.1.1").fuentes if f.nueva)
    assert enlace.ref.startswith("https://ejemplo.invalid/drive/nuevo")


def test_soportes_del_revisor_se_leen_de_la_cache(tmp_path: Path) -> None:
    cache = Cache(tmp_path / "cache")
    archivo = Archivo("S007", "x/a.pdf", "a.pdf", "1.1.2", "pdf", ".pdf", 10)
    desc = DescripcionArchivo(
        que_es="Acta ficticia",
        fecha_documento="2026-09",
        que_prueba="Una jornada con profesores",
        ilegible_o_incompleto=False,
        observacion="",
    )
    guardar_indice(
        cache, "Ejemplo Alfa", {"S007": Descripcion("S007", desc)}, [archivo]
    )
    soportes = leer_indice(cache, "ejemplo alfa")
    a = accion(armar(guardado(), soportes=soportes), "1.1.2")
    s = next(f for f in a.fuentes if f.id == "S007")
    assert s.tipo == "soporte" and a.hay_evidencia_nueva
    assert "a.pdf" not in s.texto  # el nombre del archivo no se manda
    assert leer_indice(cache, "Otro programa") == {}


async def test_citas_verificadas_y_lo_que_no_tiene_fuente_se_marca(
    tmp_path: Path,
) -> None:
    pg = guardado()
    con_enlace(pg, "1.1.1", "https://ejemplo.invalid/drive/nuevo", "2026-10-01T10:00")
    modelo = ModeloFalso(
        [
            {"texto": "Se publicó una pieza [P1, E1].", "citas": ["P1", "E1"]},
            {"texto": "Asistieron 40 personas.", "citas": []},
            {"texto": "Se aprobó el cierre [P99].", "citas": ["P99"]},
            {"texto": "Hay un acta [S555].", "citas": ["S555"]},
        ]
    )
    b = await correr(pg, modelo, tmp_path)
    a = next(x for x in b.acciones if x.base.accion.codigo == "1.1.1")
    assert [f.valida for f in a.frases] == [True, False, False, False]
    assert [f.motivo for f in a.rechazadas] == [
        bor.MOTIVO_SIN_CITA,
        bor.MOTIVO_CITA_INVALIDA,
        bor.MOTIVO_CITA_INVALIDA,
    ]
    assert a.parrafo == "Se publicó una pieza [P1, E1]."
    assert "40 personas" not in a.parrafo and a.citas == ["P1", "E1"]
    assert a.alertas


async def test_accion_sin_evidencia_no_llama_al_modelo_ni_inventa(
    tmp_path: Path,
) -> None:
    modelo = ModeloFalso()
    b = await correr(guardado(), modelo, tmp_path)
    assert modelo.textos == []
    for a in b.acciones:
        assert a.sin_evidencia_nueva
        assert a.parrafo.startswith(bor.SIN_EVIDENCIA_NUEVA)
        assert a.citas[0].startswith("P")
    assert b.resumen()["sin_evidencia_nueva"] == len(b.acciones)


async def test_sin_nombres_de_personas(tmp_path: Path) -> None:
    pg = guardado(PERSONA)
    con_enlace(pg, "1.1.1", "https://ejemplo.invalid/drive/nuevo", "2026-10-01T10:00")
    nombres = fuentes.nombres_en_plan(pg)
    assert PERSONA in nombres
    # Los cargos del resto del plan no se toman por nombres.
    assert not any("Dirección" in n or "Comité" in n for n in nombres)

    modelo = ModeloFalso(
        [
            {"texto": f"{PERSONA} publicó la pieza [P1].", "citas": ["P1"]},
            {"texto": "Escribir a ana@ejemplo.test [P1].", "citas": ["P1"]},
            {"texto": "El profesor Juan Gómez lo aprobó [P1].", "citas": ["P1"]},
            {
                "texto": "La dirección del programa publicó la pieza [E1].",
                "citas": ["E1"],
            },
        ]
    )
    b = await correr(pg, modelo, tmp_path)
    a = next(x for x in b.acciones if x.base.accion.codigo == "1.1.1")
    assert [f.valida for f in a.frases] == [False, False, False, True]
    assert all(f.motivo == bor.MOTIVO_NOMBRE for f in a.rechazadas)
    # Ni el prompt ni las fuentes llevan el nombre.
    assert PERSONA not in modelo.textos[0]
    assert all(PERSONA not in f.texto for f in a.base.fuentes)

    salidas.borrador_docx(tmp_path / "b.docx", b)
    salidas.revision_xlsx(tmp_path / "r.xlsx", b)
    # En las salidas el nombre aparece solo dentro de la frase rechazada.
    doc = Document(str(tmp_path / "b.docx"))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "no se afirma" in texto and "La dirección del programa publicó" in texto
    propuesta = load_workbook(tmp_path / "r.xlsx")["Revisión del corte"]["E2"].value
    assert PERSONA not in propuesta and "ana@" not in propuesta


def test_el_prompt_no_manda_responsables_ni_correos() -> None:
    pg = guardado("Zoila Ficticia Pérez, zoila@ejemplo.test")
    a = accion(armar(pg), "1.1.1")
    prompt = bor.texto_para_modelo(a)
    assert "Zoila" not in prompt and "zoila@" not in prompt
    assert "Fuentes (cita solo estas)" in prompt


async def test_salidas_docx_y_xlsx(tmp_path: Path) -> None:
    pg = guardado()
    con_enlace(pg, "1.1.1", "https://ejemplo.invalid/drive/nuevo", "2026-10-01T10:00")
    b = await correr(pg, ModeloFalso(), tmp_path)
    salidas.borrador_docx(tmp_path / "b.docx", b)
    salidas.revision_xlsx(tmp_path / "r.xlsx", b)

    doc = Document(str(tmp_path / "b.docx"))
    assert doc.sections[0].header.paragraphs[0].text.startswith("BORRADOR — propuesta")
    assert doc.paragraphs[0].text.startswith("BORRADOR — propuesta")
    titulos = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    assert any(t.startswith("Objetivo 1.") for t in titulos)
    assert any(t.startswith("Acción 1.1.1") for t in titulos)
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "[P1] «2. Plan de Mejora» fila" in texto  # acción sin evidencia nueva
    assert re.search(r"\[E\d\] https://ejemplo.invalid/drive/nuevo", texto)
    assert bor.SIN_EVIDENCIA_NUEVA in texto

    ws = load_workbook(tmp_path / "r.xlsx")["Revisión del corte"]
    columnas = [c.value for c in ws[1]]
    assert "Decisión (Aprobado/Editar/Descartar)" in columnas
    assert "Texto final" in columnas
    assert ws.max_row == 1 + len(b.acciones)
    assert (
        ws.data_validations.dataValidation[0].formula1 == '"Aprobado,Editar,Descartar"'
    )
