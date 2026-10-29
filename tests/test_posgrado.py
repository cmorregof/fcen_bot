"""Posgrado: programas, proyecto de tesis y «cursos» en el oficio. Datos ficticios."""

from __future__ import annotations

import io
from datetime import date

from docx import Document

from app.core.programas import PROGRAMAS, identificar_programa, programas_de
from app.schemas.caso import (
    AsignaturaSolicitada,
    Caso,
    EstadoCaso,
    Jurado,
    ProyectoTesis,
    Sentido,
    TipoSolicitud,
)
from app.services.documento import CasoDocumento, DatosOficio, construir_oficio
from app.services.ficha import ficha
from app.services.reglas import analizar
from tests.conftest import extraccion_ejemplo


def _extraccion(**cambios):
    base = dict(
        tipos=[TipoSolicitud.PROYECTO_TESIS],
        objeto_solicitud="inscripción del proyecto de tesis de maestría",
        asignaturas=[
            AsignaturaSolicitada(
                codigo="9100001",
                nombre="Tesis de Maestría Ficticia",
                grupo="1",
                creditos=12,
                accion="inscribir",
            )
        ],
        proyecto=ProyectoTesis(
            titulo="Modelo ficticio de transporte de calor",
            director="Profesor Ficticio A",
        ),
        jurados=[Jurado(nombre=f"Profesor Ficticio {c}") for c in "BCDEFG"],
        sia=None,
    )
    base.update(cambios)
    return extraccion_ejemplo(**base)


def _oficio(e, clave="mcf") -> Document:
    programa = PROGRAMAS[clave]
    caso = Caso(
        id=1,
        chat_id=1,
        direccion=programa.direccion,
        programa=clave,
        estado=EstadoCaso.LISTO,
        sentido=Sentido.APROBAR,
        texto="x",
        extraccion=e,
    )
    item = CasoDocumento(
        caso,
        e,
        analizar(e, programa, periodo_inicio=None, hay_imagenes=False),
        "Buenos días,\n\nSolicito inscribir mi proyecto.\n\nPersona Ficticia Uno",
    )
    datos = DatosOficio(
        programa=programa, fecha=date(2026, 10, 29), casos=[item], elaboro="AB"
    )
    return Document(io.BytesIO(construir_oficio(datos, "2026-2S")))


def test_posgrados_por_area() -> None:
    assert PROGRAMAS["mcf"].nivel == "posgrado"
    assert PROGRAMAS["mcf"] in programas_de("ciencias_naturales")
    assert PROGRAMAS["mcma"] in programas_de("matematicas")
    assert PROGRAMAS["mecen"] in programas_de("matematicas")
    assert PROGRAMAS["matematicas"].nivel == "pregrado"


def test_se_identifica_por_el_nombre_del_plan() -> None:
    assert (
        identificar_programa(None, "MAESTRIA EN CIENCIAS - FISICA") == PROGRAMAS["mcf"]
    )
    assert (
        identificar_programa(None, "Maestría en Ciencias - Matemática Aplicada")
        == PROGRAMAS["mcma"]
    )


def test_oficio_trae_proyecto_cursos_y_jurados() -> None:
    doc = _oficio(_extraccion())
    textos = [p.text for p in doc.paragraphs]
    assert "Proyecto de tesis o trabajo final:" in textos
    assert "Inscribir los siguientes cursos:" in textos
    assert "Inscribir las siguientes asignaturas:" not in textos
    assert "Se recomiendan los siguientes jurados:" in textos
    proyecto = doc.tables[0]
    assert proyecto.rows[0].cells[1].text == "Modelo ficticio de transporte de calor"
    assert proyecto.rows[1].cells[1].text == "Profesor Ficticio A"


def test_director_faltante_se_resalta_y_avisa() -> None:
    e = _extraccion(proyecto=ProyectoTesis(titulo="Proyecto ficticio"))
    doc = _oficio(e)
    assert 'w:fill="FFFF00"' in doc.tables[0].rows[1].cells[1]._tc.xml
    a = analizar(e, PROGRAMAS["mcf"], periodo_inicio=None, hay_imagenes=False)
    assert any("Falta el director" in x.mensaje for x in a.alertas)


def test_proyecto_en_pregrado_avisa() -> None:
    a = analizar(
        _extraccion(), PROGRAMAS["ing_fisica"], periodo_inicio=None, hay_imagenes=False
    )
    assert any("Revisa el programa del caso" in x.mensaje for x in a.alertas)


def test_proyecto_no_usa_normas_de_trabajo_de_grado_de_pregrado() -> None:
    texto = " ".join(p.text for p in _oficio(_extraccion()).paragraphs)
    assert "CFCEN-060" not in texto


def test_ficha_muestra_el_proyecto() -> None:
    e = _extraccion()
    caso = Caso(
        id=7,
        chat_id=1,
        direccion="ciencias_naturales",
        programa="mcf",
        estado=EstadoCaso.LISTO,
        extraccion=e,
    )
    a = analizar(e, PROGRAMAS["mcf"], periodo_inicio=None, hay_imagenes=False)
    texto = ficha(caso, a)
    assert "Proyecto: «Modelo ficticio de transporte de calor»" in texto
    assert "Director(a): Profesor Ficticio A" in texto
