"""Indicadores de tiempos y reporte mensual del servicio de actas.

Cálculo de medianas, «pendiente de acta» nunca contado como aprobado, filtros
por área y rol, reporte mensual apagado y Excel con la pestaña «fuentes».
Todos los datos son ficticios.
"""

from __future__ import annotations

import io
from datetime import UTC, date, datetime
from pathlib import Path

import aiosqlite
import pytest
import yaml
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api.web import crear_router
from app.api.web.config import web_settings
from app.core import solicitudes
from app.core.config import RAIZ, settings
from app.db import repo
from app.main import app
from app.schemas.caso import Sentido, TipoSolicitud
from app.schemas.consejo import SentidoConsejo
from app.schemas.telegram import Update
from app.services import bot, indicadores, reporte_mensual, resumen
from app.services.indicadores import FilaCaso
from tests.conftest import TelegramFalso, extraccion_ejemplo
from tests.web.test_web import _token

DECANA = 9001
ADMIN = 9100
VICEDECANA = 9200
ADM_NATURALES = 555
ADM_MATE = 556
HOY = date(2026, 11, 1)
OCTUBRE = indicadores.mes(2026, 10)
SIN_PLAZOS = solicitudes.Plazos()
NOMBRE = "Persona Ficticia Uno"
DOCUMENTO = "1000000001"


@pytest.fixture(autouse=True)
def entorno(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "decanatura_telegram_ids", str(DECANA))
    monkeypatch.setattr(settings, "admin_telegram_ids", str(ADMIN))
    monkeypatch.setattr(settings, "vicedecanatura_telegram_ids", str(VICEDECANA))
    monkeypatch.setattr(settings, "reporte_mensual_activo", True)
    monkeypatch.setattr(settings, "periodo_inicio", date(2026, 8, 3))
    monkeypatch.setattr(settings, "periodo_actual", "2026-2S")
    monkeypatch.setattr(settings, "timezone", "America/Bogota")
    monkeypatch.setattr(indicadores, "hoy_local", lambda: HOY)
    monkeypatch.setattr(
        resumen, "_ahora", lambda: datetime(2026, 11, 2, 12, 30, tzinfo=UTC)
    )


# ---------------------------------------------------------------------------
# Escenario ficticio
# ---------------------------------------------------------------------------


async def _caso(
    db: aiosqlite.Connection,
    *,
    programa: str = "ing_fisica",
    creado: str = "2026-10-05 15:00:00",
    fecha_solicitud: str | None = "2026-10-01",
    sentido: Sentido = Sentido.PENDIENTE,
    recomendado: str | None = None,
    tipos: list[TipoSolicitud] | None = None,
) -> int:
    direccion = (
        "matematicas"
        if programa in ("matematicas", "ciencias_computacion", "estadistica")
        else "ciencias_naturales"
    )
    caso_id = await repo.crear_caso(db, chat_id=1, user_id=1, direccion=direccion)
    ext = extraccion_ejemplo(
        fecha_solicitud=fecha_solicitud,
        tipos=tipos or [TipoSolicitud.CARGA_MINIMA],
    )
    await repo.guardar_extraccion(db, caso_id, ext, programa)
    await db.execute(
        "UPDATE casos SET creado = ?, sentido = ?, recomendado = ? WHERE id = ?",
        (creado, sentido.value, recomendado, caso_id),
    )
    await db.commit()
    return caso_id


async def _acta(
    db: aiosqlite.Connection,
    caso_id: int,
    sentido: SentidoConsejo,
    fecha: str | None = "2026-10-20",
    *,
    numero: str = "14",
    instancia: str = "facultad",
    confirmado: str = "2026-10-30 12:00:00",
) -> None:
    cur = await db.execute(
        "INSERT INTO actas_consejo (instancia, numero, fecha, texto, subido_por)"
        " VALUES (?, ?, ?, 'Acta ficticia', 'prueba')",
        (instancia, numero, fecha),
    )
    await db.execute(
        "INSERT INTO decisiones_consejo (caso_id, acta_id, decision_id, instancia,"
        " sentido, texto_decision, acta_numero, acta_fecha, fragmento,"
        " confirmado_por, confirmado)"
        " VALUES (?, ?, 1, ?, ?, 'texto ficticio', ?, ?, 'fragmento ficticio',"
        " 'prueba', ?)",
        (caso_id, cur.lastrowid, instancia, sentido.value, numero, fecha, confirmado),
    )
    await db.commit()


async def _escenario(db: aiosqlite.Connection) -> None:
    """Octubre: Física 4 casos, Matemáticas 2. Uno de septiembre fuera."""
    # Física: recomendado aprobar y aprobado por el Consejo (2 y 10 días).
    c1 = await _caso(db, sentido=Sentido.APROBAR, recomendado="2026-10-03 15:00:00")
    await _acta(db, c1, SentidoConsejo.APROBO, "2026-10-13")
    # Física: recomendado aprobar, SIN acta → pendiente, nunca aprobado (6 días).
    await _caso(db, sentido=Sentido.APROBAR, recomendado="2026-10-07 15:00:00")
    # Física: recomendado no aprobar y aplazado por el Consejo (4 días).
    c3 = await _caso(db, sentido=Sentido.NO_APROBAR, recomendado="2026-10-05 15:00:00")
    await _acta(db, c3, SentidoConsejo.APLAZO, "2026-10-20")
    # Física: todavía sin recomendación.
    await _caso(db)
    # Matemáticas: no aprobar, y el Consejo no aprobó (12 y 4 días).
    c5 = await _caso(
        db,
        programa="matematicas",
        sentido=Sentido.NO_APROBAR,
        recomendado="2026-10-13 15:00:00",
        tipos=[TipoSolicitud.CANCELACION],
    )
    await _acta(db, c5, SentidoConsejo.NO_APROBO, "2026-10-17", numero="15")
    # Matemáticas: recomendado antes de que MemorIA guardara la fecha.
    await _caso(
        db,
        programa="matematicas",
        sentido=Sentido.APROBAR,
        recomendado=None,
        tipos=[TipoSolicitud.CANCELACION],
    )
    # Septiembre: no entra en octubre.
    await _caso(db, creado="2026-09-20 15:00:00", fecha_solicitud="2026-09-18")


# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------


def test_mediana_y_percentil_90() -> None:
    assert indicadores.estadistica([]) == indicadores.Estadistica(0, None, None)
    assert indicadores.estadistica([5]) == indicadores.Estadistica(1, 5.0, 5)
    e = indicadores.estadistica([4, 1, 3, 2])
    assert (e.casos, e.mediana, e.p90) == (4, 2.5, 4)
    # Rango más cercano: posición ⌈0,9·10⌉ = 9.
    assert indicadores.percentil(list(range(1, 11)), 0.9) == 9
    assert indicadores.percentil(list(range(1, 21)), 0.9) == 18


async def test_cifras_del_periodo(db: aiosqlite.Connection) -> None:
    await _escenario(db)
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    t = r.total
    assert r.alcance == indicadores.FACULTAD
    assert t.recibidos == 6
    assert (t.recomendar_aprobar, t.recomendar_no_aprobar) == (3, 2)
    assert (t.consejo_aprobo, t.consejo_no_aprobo, t.consejo_aplazo) == (1, 1, 1)
    assert t.decididos_consejo == 2
    # El de Física sin acta, el aplazado y el de Matemáticas sin acta.
    assert t.pendientes_acta == 3
    # Solicitud → recomendación: 2, 6, 4 y 12 días (el sin fecha no entra).
    assert t.dias_comite == indicadores.Estadistica(4, 5.0, 12)
    # Recomendación → acta: 10 (aprobó) y 4 (no aprobó). El aplazado no entra.
    assert t.dias_consejo == indicadores.Estadistica(2, 7.0, 10)
    # Sin plazos en config/solicitudes.yaml: no se cuenta nada.
    assert t.fuera_plazo_comite is None and t.fuera_plazo_consejo is None

    programas = {g.nombre: g.indicadores for g in r.de("Programa")}
    assert programas["Ingeniería Física"].recibidos == 4
    assert programas["Matemáticas"].recibidos == 2
    tipos = {g.nombre: g.indicadores.recibidos for g in r.de("Tipo de solicitud")}
    assert sum(tipos.values()) == 6  # un caso cuenta una sola vez
    instancias = {g.nombre: g.indicadores.recibidos for g in r.de("Instancia")}
    # Con acta, la del acta; sin acta, la ruta (cancelación: por confirmar).
    assert instancias == {"Consejo de Facultad": 5, "Instancia por confirmar": 1}
    assert {g.dimension for g in r.grupos} >= {"Área curricular", "Programa"}


async def test_pendiente_de_acta_nunca_cuenta_como_aprobado(
    db: aiosqlite.Connection,
) -> None:
    await _caso(db, sentido=Sentido.APROBAR, recomendado="2026-10-03 15:00:00")
    aplazado = await _caso(
        db, sentido=Sentido.APROBAR, recomendado="2026-10-03 15:00:00"
    )
    await _acta(db, aplazado, SentidoConsejo.APLAZO)
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    assert r.total.recomendar_aprobar == 2
    assert r.total.consejo_aprobo == 0
    assert r.total.decididos_consejo == 0
    assert r.total.pendientes_acta == 2
    texto = indicadores.formatear(r)
    assert "Decididos por el Consejo: 0 (aprobó 0" in texto
    assert "Pendientes de acta del Consejo: 2" in texto


def _fila(**cambios: object) -> FilaCaso:
    base = {
        "direccion": "ciencias_naturales",
        "programa": "Ingeniería Física",
        "tipo": "Carga mínima",
        "instancia": "Consejo de Facultad",
        "recibido": date(2026, 10, 1),
        "solicitud": date(2026, 10, 1),
        "recomendado": None,
        "sentido": Sentido.PENDIENTE,
        "consejo": None,
        "acta_fecha": None,
    }
    base.update(cambios)
    return FilaCaso(**base)  # type: ignore[arg-type]


def test_fuera_de_plazo_con_plazos_configurados() -> None:
    plazos = solicitudes.Plazos(
        comite=solicitudes.Plazo(dias=10, fuente="plazo ficticio"),
        consejo=solicitudes.Plazo(dias=15, fuente="plazo ficticio"),
    )
    filas = [
        # Recomendado a los 12 días: fuera en el comité; decidido a los 5: a tiempo.
        _fila(
            sentido=Sentido.APROBAR,
            recomendado=date(2026, 10, 13),
            consejo=SentidoConsejo.APROBO,
            acta_fecha=date(2026, 10, 18),
        ),
        # Sin recomendación y ya pasaron 31 días: fuera en el comité.
        _fila(),
        # Recomendado a tiempo y pendiente de acta desde hace 25 días: fuera.
        _fila(sentido=Sentido.APROBAR, recomendado=date(2026, 10, 7)),
        # Recibido ayer: a tiempo en todo.
        _fila(recibido=date(2026, 10, 31), solicitud=date(2026, 10, 31)),
    ]
    i = indicadores.calcular(filas, HOY, plazos)
    assert (i.fuera_plazo_comite, i.fuera_plazo_consejo) == (2, 1)
    sin = indicadores.calcular(filas, HOY, SIN_PLAZOS)
    assert (sin.fuera_plazo_comite, sin.fuera_plazo_consejo) == (None, None)


async def test_fechas_en_hora_de_bogota_y_solicitud_posterior(
    db: aiosqlite.Connection,
) -> None:
    # 1-nov 03:00 UTC es 31-oct en Bogotá: entra en octubre.
    await _caso(db, creado="2026-11-01 03:00:00", fecha_solicitud="2026-12-01")
    filas = await indicadores.leer_casos(db, OCTUBRE)
    assert len(filas) == 1
    # Una fecha de solicitud posterior al registro es mala lectura: se usa el registro.
    assert filas[0].recibido == filas[0].solicitud == date(2026, 10, 31)


async def test_la_recomendacion_guarda_su_fecha(db: aiosqlite.Connection) -> None:
    caso_id = await _caso(db)

    async def recomendado() -> object:
        cur = await db.execute("SELECT recomendado FROM casos WHERE id = ?", (caso_id,))
        fila = await cur.fetchone()
        assert fila is not None
        return fila[0]

    await repo.set_sentido(db, caso_id, Sentido.APROBAR)
    primera = await recomendado()
    assert primera is not None
    await db.execute(
        "UPDATE casos SET recomendado = '2026-10-02 10:00:00' WHERE id = ?", (caso_id,)
    )
    await repo.set_sentido(db, caso_id, Sentido.APROBAR)  # mismo sentido: no cambia
    assert await recomendado() == "2026-10-02 10:00:00"
    await repo.set_sentido(db, caso_id, Sentido.PENDIENTE)
    assert await recomendado() is None


@pytest.mark.parametrize(
    ("texto", "desde", "hasta"),
    [
        ("2026-10", date(2026, 10, 1), date(2026, 11, 1)),
        ("2026-12", date(2026, 12, 1), date(2027, 1, 1)),
        ("2026", date(2026, 1, 1), date(2027, 1, 1)),
        ("", date(2026, 8, 3), date(2026, 11, 2)),
        ("2026-2s", date(2026, 8, 3), date(2026, 11, 2)),
    ],
)
def test_periodos(texto: str, desde: date, hasta: date) -> None:
    p = indicadores.periodo_de(texto, HOY)
    assert (p.desde, p.hasta) == (desde, hasta)


@pytest.mark.parametrize("texto", ["octubre", "2026-13", "1999", "2026-10-01"])
def test_periodo_invalido(texto: str) -> None:
    with pytest.raises(indicadores.PeriodoInvalido):
        indicadores.periodo_de(texto, HOY)


def test_mes_anterior() -> None:
    assert indicadores.mes_anterior(date(2026, 11, 1)) == OCTUBRE
    assert indicadores.mes_anterior(date(2027, 1, 1)).desde == date(2026, 12, 1)


# ---------------------------------------------------------------------------
# Privacidad y fuentes
# ---------------------------------------------------------------------------


async def test_mensaje_sin_datos_personales_y_con_fuentes(
    db: aiosqlite.Connection,
) -> None:
    await _escenario(db)
    texto = indicadores.formatear(await indicadores.reporte(db, OCTUBRE, HOY))
    assert NOMBRE not in texto and DOCUMENTO not in texto
    assert "<b>Fuentes</b>" in texto
    for n in range(1, 8):
        assert f"[{n}]" in texto
    assert "sin plazo configurado" in texto
    assert "pendiente de acta del Consejo y no como aprobado" in texto


async def test_xlsx_con_pestana_de_fuentes(db: aiosqlite.Connection) -> None:
    await _escenario(db)
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    libro = load_workbook(io.BytesIO(indicadores.xlsx(r)))
    assert libro.sheetnames == ["indicadores", "fuentes"]
    hoja = libro["indicadores"]
    columnas = indicadores.columnas()
    encabezados = [c.value for c in hoja[1]]
    assert encabezados == ["Dimensión", "Grupo", *(c.encabezado for c in columnas)]
    total = [c.value for c in hoja[2]]
    assert total[:3] == ["Total", indicadores.FACULTAD, 6]
    fuentes = libro["fuentes"]
    textos = [str(c.value) for fila in fuentes.iter_rows() for c in fila if c.value]
    for c in columnas:
        assert c.encabezado in textos
        assert indicadores.FUENTES[c.clave].fuente in textos
    assert any("plantilla oficial" in t for t in textos)
    todo = textos + [str(c.value) for f in hoja.iter_rows() for c in f if c.value]
    assert not any(NOMBRE in t or DOCUMENTO in t for t in todo)


def test_columnas_configurables(tmp_path: Path) -> None:
    config = indicadores.cargar_columnas(RAIZ / "config" / "indicadores.yaml")
    assert {c.clave for c in config.columnas} == set(indicadores.CLAVES)
    malo = tmp_path / "indicadores.yaml"
    malo.write_text(
        yaml.safe_dump({"columnas": [{"clave": "inventada", "encabezado": "X"}]}),
        encoding="utf-8",
    )
    with pytest.raises(indicadores.ConfigIndicadoresError, match="inventada"):
        indicadores.cargar_columnas(malo)
    with pytest.raises(indicadores.ConfigIndicadoresError):
        indicadores.cargar_columnas(tmp_path / "no-existe.yaml")


def test_columnas_a_la_medida(tmp_path: Path) -> None:
    r = indicadores.Reporte(
        OCTUBRE, indicadores.FACULTAD, HOY, SIN_PLAZOS, indicadores.Indicadores(), []
    )
    cols = [indicadores.Columna(clave="pendientes_acta", encabezado="Pendientes")]
    hoja = load_workbook(io.BytesIO(indicadores.xlsx(r, cols)))["indicadores"]
    assert [c.value for c in hoja[1]] == ["Dimensión", "Grupo", "Pendientes"]


def test_plazos_vacios_en_el_yaml_del_repo() -> None:
    config = solicitudes.cargar(RAIZ / "config" / "solicitudes.yaml")
    assert config.plazos.comite is None and config.plazos.consejo is None


# ---------------------------------------------------------------------------
# Telegram: /indicadores solo para decanatura y admin
# ---------------------------------------------------------------------------

_n = 0


def _update(uid: int, texto: str) -> Update:
    global _n
    _n += 1
    return Update.model_validate(
        {
            "update_id": 70_000 + _n,
            "message": {
                "message_id": _n,
                "chat": {"id": uid, "type": "private"},
                "from": {"id": uid, "first_name": "Persona"},
                "text": texto,
            },
        }
    )


@pytest.mark.parametrize("uid", [DECANA, ADMIN])
async def test_indicadores_en_telegram(
    uid: int, db: aiosqlite.Connection, tg: TelegramFalso
) -> None:
    await _escenario(db)
    await bot.manejar_update(db, _update(uid, "/indicadores 2026-10"))
    assert len(tg.mensajes) == 1
    texto = tg.mensajes[0]["texto"]
    assert "octubre de 2026" in texto and "Casos recibidos: 6" in texto
    assert NOMBRE not in texto
    assert len(tg.documentos) == 1
    doc = tg.documentos[0]
    assert doc["nombre"].endswith(".xlsx")
    assert "fuentes" in load_workbook(io.BytesIO(doc["contenido"])).sheetnames


async def test_periodo_mal_escrito_en_telegram(
    db: aiosqlite.Connection, tg: TelegramFalso
) -> None:
    await bot.manejar_update(db, _update(DECANA, "/indicadores octubre"))
    assert "No reconozco el periodo" in tg.textos()[0]
    assert tg.documentos == []


async def test_administrativa_no_tiene_indicadores_en_telegram(
    db: aiosqlite.Connection, tg: TelegramFalso
) -> None:
    await repo.registrar_usuario(db, ADM_NATURALES, "Administrativa ficticia")
    await repo.asignar_direccion(db, ADM_NATURALES, "ciencias_naturales")
    await _escenario(db)
    await bot.manejar_update(db, _update(ADM_NATURALES, "/indicadores"))
    assert tg.documentos == []
    assert not any("Indicadores del servicio" in t for t in tg.textos())


# ---------------------------------------------------------------------------
# Web: decanatura y admin ven la Facultad; la administrativa, solo su área
# ---------------------------------------------------------------------------


@pytest.fixture
def web(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setattr(web_settings, "session_secret", "secreto-de-prueba")
    monkeypatch.setattr(web_settings, "admins", "admin@unal.edu.co")
    monkeypatch.setattr(settings, "bot_servicio", "actas")
    aplicacion = FastAPI()
    aplicacion.include_router(crear_router("actas"))
    return TestClient(aplicacion)


async def test_web_filtra_por_rol_y_area(
    db: aiosqlite.Connection, web: TestClient
) -> None:
    await _escenario(db)
    ruta = "/api/web/actas/indicadores?periodo=2026-10"

    decana = await _token("decana@unal.edu.co", "decanatura")
    r = web.get(ruta, headers=decana)
    assert r.status_code == 200
    assert r.json()["alcance"] == indicadores.FACULTAD
    assert r.json()["total"]["recibidos"] == 6
    assert "Área curricular" in [d["nombre"] for d in r.json()["dimensiones"]]

    admin = await _token("admin@unal.edu.co", "admin")
    assert web.get(ruta, headers=admin).json()["total"]["recibidos"] == 6

    mate = await _token("mate@unal.edu.co", "administrativa", "matematicas")
    datos = web.get(ruta, headers=mate).json()
    assert datos["alcance"] == "Área Curricular de Matemáticas y Estadística"
    assert datos["total"]["recibidos"] == 2
    programas = next(d for d in datos["dimensiones"] if d["nombre"] == "Programa")
    assert [g["nombre"] for g in programas["grupos"]] == ["Matemáticas"]
    assert "Área curricular" not in [d["nombre"] for d in datos["dimensiones"]]
    assert NOMBRE not in r.text and DOCUMENTO not in r.text

    comite = await _token("comite@unal.edu.co", "comite", "matematicas")
    assert web.get(ruta, headers=comite).status_code == 403
    sin_area = await _token("sinarea@unal.edu.co", "administrativa")
    assert web.get(ruta, headers=sin_area).status_code == 403
    assert web.get(ruta).status_code == 401
    malo = web.get("/api/web/actas/indicadores?periodo=octubre", headers=decana)
    assert malo.status_code == 422


async def test_web_descarga_xlsx_del_area(
    db: aiosqlite.Connection, web: TestClient
) -> None:
    await _escenario(db)
    mate = await _token("mate@unal.edu.co", "administrativa", "matematicas")
    r = web.get("/api/web/actas/indicadores/xlsx?periodo=2026-10", headers=mate)
    assert r.status_code == 200
    assert r.headers["content-type"] == indicadores.MIME_XLSX
    libro = load_workbook(io.BytesIO(r.content))
    assert libro.sheetnames == ["indicadores", "fuentes"]
    assert libro["indicadores"]["C2"].value == 2


# ---------------------------------------------------------------------------
# Reporte mensual
# ---------------------------------------------------------------------------


async def _personas(db: aiosqlite.Connection) -> None:
    for uid, direccion in (
        (ADM_NATURALES, "ciencias_naturales"),
        (ADM_MATE, "matematicas"),
    ):
        await repo.registrar_usuario(db, uid, "Administrativa ficticia")
        await repo.asignar_direccion(db, uid, direccion)


def _cron(monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setattr(settings, "internal_cron_secret", "cron-prueba")
    monkeypatch.setattr(settings, "bot_servicio", "actas")
    return TestClient(app)


async def test_reporte_mensual_por_area_y_consolidado(
    db: aiosqlite.Connection, tg: TelegramFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    await _personas(db)
    await _escenario(db)
    with _cron(monkeypatch) as cliente:
        r = cliente.post(
            "/internal/reporte-mensual", headers={"X-Cron-Secret": "cron-prueba"}
        )
    assert r.status_code == 200
    assert r.json() == {
        "ok": True,
        "activo": True,
        "mes": "2026-10",
        "areas": {"ciencias_naturales": 1, "matematicas": 1},
        "vicedecanatura": 1,
        "enviados": 3,
    }
    por_chat = {m["chat_id"]: m["texto"] for m in tg.mensajes}
    assert "Área Curricular de Ciencias Naturales" in por_chat[ADM_NATURALES]
    assert "Casos recibidos: 4" in por_chat[ADM_NATURALES]
    assert "Matemáticas" not in por_chat[ADM_NATURALES]
    assert "Casos recibidos: 2" in por_chat[ADM_MATE]
    assert indicadores.FACULTAD in por_chat[VICEDECANA]
    assert "Casos recibidos: 6" in por_chat[VICEDECANA]
    assert "Acumulado del periodo 2026-2S" in por_chat[VICEDECANA]
    assert all(NOMBRE not in t for t in por_chat.values())


async def test_reporte_mensual_se_apaga(
    db: aiosqlite.Connection, tg: TelegramFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "reporte_mensual_activo", False)
    await _personas(db)
    await _escenario(db)
    with _cron(monkeypatch) as cliente:
        r = cliente.post(
            "/internal/reporte-mensual", headers={"X-Cron-Secret": "cron-prueba"}
        )
    assert r.json() == {"ok": True, "activo": False, "enviados": 0}
    assert tg.mensajes == []


async def test_reporte_mensual_pide_secreto(
    tg: TelegramFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    with _cron(monkeypatch) as cliente:
        r = cliente.post("/internal/reporte-mensual", headers={"X-Cron-Secret": "x"})
    assert r.status_code == 403
    assert tg.mensajes == []


async def test_reporte_mensual_sigue_si_un_envio_falla(
    db: aiosqlite.Connection, tg: TelegramFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    await _personas(db)

    async def enviar(chat_id: int, texto: str, **kw: object) -> int:
        if chat_id == ADM_NATURALES:
            raise RuntimeError("Telegram caído (ficticio)")
        return await tg.send_message(chat_id, texto)

    monkeypatch.setattr("app.integrations.telegram.send_message", enviar)
    resultado = await reporte_mensual.enviar(db, HOY)
    assert resultado["areas"] == {"ciencias_naturales": 0, "matematicas": 1}
    assert resultado["vicedecanatura"] == 1


def test_workflow_mensual() -> None:
    flujo = yaml.safe_load(
        (RAIZ / ".github" / "workflows" / "reporte-mensual.yml").read_text("utf-8")
    )
    disparo = flujo[True]  # «on» se lee como True en YAML 1.1
    assert disparo["schedule"] == [{"cron": "30 12 1 * *"}]  # 7:30 Bogotá
    paso = flujo["jobs"]["enviar"]["steps"][0]
    assert "/internal/reporte-mensual" in paso["run"]
    assert "INTERNAL_CRON_SECRET" in paso["env"]


# ---------------------------------------------------------------------------
# Resumen semanal: actas del Consejo de la semana
# ---------------------------------------------------------------------------


async def test_resumen_semanal_trae_las_actas_de_la_semana(
    db: aiosqlite.Connection,
) -> None:
    await _escenario(db)
    viejo = await _caso(db, sentido=Sentido.APROBAR)
    await _acta(
        db, viejo, SentidoConsejo.APROBO, "2026-09-01", confirmado="2026-09-02 10:00:00"
    )
    texto = await resumen.armar(db)
    assert "<b>Actas del Consejo registradas esta semana</b>" in texto
    assert "Consejo de Facultad, acta 14 del 2026-10-13: aprobó 1" in texto
    assert (
        "Consejo de Facultad, acta 14 del 2026-10-20: aprobó 0, no aprobó 0, aplazó 1"
        in texto
    )
    assert "acta 15 del 2026-10-17: aprobó 0, no aprobó 1" in texto
    assert "2026-09-01" not in texto
    assert NOMBRE not in texto


# ---------------------------------------------------------------------------
# Revisión del PR #19
# ---------------------------------------------------------------------------


async def _recomendado(db: aiosqlite.Connection, caso_id: int) -> object:
    cur = await db.execute("SELECT recomendado FROM casos WHERE id = ?", (caso_id,))
    fila = await cur.fetchone()
    assert fila is not None
    return fila[0]


async def test_caso_viejo_sin_fecha_no_recibe_la_de_hoy(
    db: aiosqlite.Connection,
) -> None:
    # Recomendado antes de que existiera la columna: sin fecha.
    viejo = await _caso(db, sentido=Sentido.APROBAR, recomendado=None)
    await repo.set_sentido(db, viejo, Sentido.APROBAR)  # mismo botón otra vez
    assert await _recomendado(db, viejo) is None
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    assert r.total.dias_comite.casos == 0
    assert r.total.sin_fecha_recomendacion == 1
    assert "1 recomendados sin fecha quedan fuera" in indicadores.formatear(r)
    # Un cambio real de sentido sí registra la fecha.
    await repo.set_sentido(db, viejo, Sentido.NO_APROBAR)
    assert await _recomendado(db, viejo) is not None


async def test_cambiar_el_sentido_despues_del_acta_no_mueve_la_fecha(
    db: aiosqlite.Connection,
) -> None:
    caso_id = await _caso(
        db, sentido=Sentido.APROBAR, recomendado="2026-10-03 15:00:00"
    )
    await _acta(db, caso_id, SentidoConsejo.APROBO, "2026-10-13")
    await repo.set_sentido(db, caso_id, Sentido.NO_APROBAR)
    assert await _recomendado(db, caso_id) == "2026-10-03 15:00:00"
    await repo.set_sentido(db, caso_id, Sentido.PENDIENTE)
    assert await _recomendado(db, caso_id) == "2026-10-03 15:00:00"
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    assert r.total.dias_consejo.casos == 1


async def test_fechas_inconsistentes_se_cuentan_aparte(
    db: aiosqlite.Connection,
) -> None:
    caso_id = await _caso(
        db, sentido=Sentido.APROBAR, recomendado="2026-10-20 15:00:00"
    )
    await _acta(db, caso_id, SentidoConsejo.APROBO, "2026-10-13")  # antes
    r = await indicadores.reporte(db, OCTUBRE, HOY)
    assert r.total.decididos_consejo == 1
    assert r.total.dias_consejo.casos == 0
    assert r.total.fechas_inconsistentes == 1
    assert "1 con fechas inconsistentes quedan fuera" in indicadores.formatear(r)


async def test_aplazado_sin_recomendacion_es_pendiente_de_acta(
    db: aiosqlite.Connection,
) -> None:
    caso_id = await _caso(db)  # sin recomendación registrada
    await _acta(db, caso_id, SentidoConsejo.APLAZO)
    t = (await indicadores.reporte(db, OCTUBRE, HOY)).total
    assert (t.recomendados, t.consejo_aplazo, t.pendientes_acta) == (0, 1, 1)
    assert t.consejo_aprobo == 0 and t.decididos_consejo == 0


@pytest.mark.parametrize("texto", ["9999-12", "9999", "2101-01", "10000"])
def test_anos_fuera_de_rango_son_periodo_invalido(texto: str) -> None:
    with pytest.raises(indicadores.PeriodoInvalido):
        indicadores.periodo_de(texto, HOY)


async def test_ano_fuera_de_rango_en_telegram_y_web(
    db: aiosqlite.Connection, tg: TelegramFalso, web: TestClient
) -> None:
    await bot.manejar_update(db, _update(DECANA, "/indicadores 9999-12"))
    assert "No reconozco el periodo" in tg.textos()[0]
    decana = await _token("decana2@unal.edu.co", "decanatura")
    r = web.get("/api/web/actas/indicadores?periodo=9999-12", headers=decana)
    assert r.status_code == 422


def test_la_web_descarga_el_periodo_consultado() -> None:
    pagina = (RAIZ / "web" / "src" / "app" / "indicadores" / "page.tsx").read_text(
        "utf-8"
    )
    assert "indicadores/xlsx?periodo=${encodeURIComponent(consultado)}" in pagina
    assert "setConsultado(p)" in pagina


async def test_excel_roto_avisa_en_telegram_y_la_web(
    db: aiosqlite.Connection,
    tg: TelegramFalso,
    web: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    def roto(*a: object, **kw: object) -> bytes:
        raise indicadores.ConfigIndicadoresError("config/indicadores.yaml roto")

    monkeypatch.setattr(indicadores, "xlsx", roto)
    await bot.manejar_update(db, _update(DECANA, "/indicadores 2026-10"))
    textos = tg.textos()
    assert "Indicadores del servicio de actas" in textos[0]
    assert "No pude armar el Excel" in textos[1]
    assert tg.documentos == []
    decana = await _token("decana3@unal.edu.co", "decanatura")
    r = web.get("/api/web/actas/indicadores/xlsx?periodo=2026-10", headers=decana)
    assert r.status_code == 503


def test_reporte_mensual_apagado_por_defecto() -> None:
    assert type(settings).model_fields["reporte_mensual_activo"].default is False


async def test_reporte_mensual_no_se_repite_en_un_rerun(
    db: aiosqlite.Connection, tg: TelegramFalso
) -> None:
    await _personas(db)
    primera = await reporte_mensual.enviar(db, HOY)
    assert primera["enviados"] == 3
    segunda = await reporte_mensual.enviar(db, HOY)
    assert segunda == {
        "activo": True,
        "mes": "2026-10",
        "ya_enviado": True,
        "enviados": 0,
    }
    assert len(tg.mensajes) == 3
    # El mes siguiente sí sale.
    tercera = await reporte_mensual.enviar(db, date(2026, 12, 1))
    assert tercera["mes"] == "2026-11" and tercera["enviados"] == 3


async def test_reporte_mensual_sin_destinatarios_se_puede_reintentar(
    db: aiosqlite.Connection, tg: TelegramFalso, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(settings, "vicedecanatura_telegram_ids", "")
    assert (await reporte_mensual.enviar(db, HOY))["enviados"] == 0
    monkeypatch.setattr(settings, "vicedecanatura_telegram_ids", str(VICEDECANA))
    assert (await reporte_mensual.enviar(db, HOY))["enviados"] == 1


def test_workflow_sin_permisos_y_con_concurrencia() -> None:
    flujo = yaml.safe_load(
        (RAIZ / ".github" / "workflows" / "reporte-mensual.yml").read_text("utf-8")
    )
    assert flujo["permissions"] == {}
    assert flujo["concurrency"]["group"] == "reporte-mensual"
