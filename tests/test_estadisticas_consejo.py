"""Estadísticas del Consejo de Facultad para Calidad. Datos ficticios."""

from __future__ import annotations

import io
import json
from datetime import date

import aiosqlite
import pytest
from openpyxl import load_workbook

from app.db.session import connect
from app.schemas.consejo import SentidoConsejo
from app.services import estadisticas_consejo as est
from app.services.notificaciones import PeriodoInvalido
from tests.web.test_web import Cliente, _token, cliente, web  # noqa: F401

HOY = date(2026, 10, 30)
NOMBRE_FICTICIO = "Persona Ficticia Dos"


async def _caso(
    db: aiosqlite.Connection,
    tipos: list[str],
    sentido: str,
    acta_fecha: str,
    *,
    programa: str = "ing_fisica",
    instancia: str = "facultad",
    estado: str = "cerrado",
) -> None:
    cur = await db.execute(
        "INSERT INTO casos (chat_id, user_id, direccion, programa, estado, sentido,"
        " extraccion_json) VALUES (1, 1, 'fisica', ?, ?, 'aprobar', ?)",
        (
            programa,
            estado,
            json.dumps({"tipos": tipos, "nombre_estudiante": NOMBRE_FICTICIO}),
        ),
    )
    await db.execute(
        "INSERT INTO decisiones_consejo (caso_id, acta_id, decision_id, instancia,"
        " sentido, texto_decision, acta_numero, acta_fecha, fragmento, confirmado_por)"
        " VALUES (?, 1, 1, ?, ?, 'texto ficticio', '35', ?, 'fragmento ficticio',"
        " 'persona@unal.edu.co')",
        (cur.lastrowid, instancia, sentido, acta_fecha),
    )
    await db.commit()


@pytest.fixture
async def db():  # type: ignore[no-untyped-def]
    conexion = await connect()
    await conexion.execute(
        "INSERT INTO actas_consejo (id, instancia, numero, fecha, texto, subido_por)"
        " VALUES (1, 'facultad', '35', '2026-09-17', 'acta ficticia', 'x@unal.edu.co')"
    )
    yield conexion
    await conexion.close()


async def _sembrar(db: aiosqlite.Connection) -> None:
    await _caso(db, ["cancelacion"], "aprobo", "2026-09-17")
    await _caso(
        db, ["inscripcion_extemporanea", "cancelacion"], "no_aprobo", "2026-09-03"
    )
    await _caso(db, ["movilidad"], "aprobo", "2026-10-08", programa="mcf")
    await _caso(db, ["homologacion"], "otro", "2026-10-08", programa="matematicas")
    await _caso(db, [], "aprobo", "2026-07-15", programa="desconocido")
    # Del primer semestre: solo entra en el año completo.
    await _caso(db, ["reingreso"], "aprobo", "2026-03-10")
    # No cuentan: aplazado, de otro año, del Consejo de Sede, descartado.
    await _caso(db, ["cancelacion"], "aplazo", "2026-09-17")
    await _caso(db, ["cancelacion"], "aprobo", "2025-09-17")
    await _caso(db, ["reingreso"], "aprobo", "2026-10-08", instancia="sede")
    await _caso(db, ["cancelacion"], "aprobo", "2026-09-17", estado="descartado")


async def test_cuenta_por_tramite_programa_y_mes(db: aiosqlite.Connection) -> None:
    await _sembrar(db)
    r = await est.calcular(db, est.periodo_de("2026-2S", HOY), HOY)
    assert r.total_casos == 5
    assert r.casos[SentidoConsejo.APROBO] == 3
    assert r.casos[SentidoConsejo.NO_APROBO] == 1
    assert r.casos[SentidoConsejo.OTRO] == 1
    canc = r.por_tramite["cancelacion"]
    assert canc[SentidoConsejo.APROBO] == 1 and canc[SentidoConsejo.NO_APROBO] == 1
    assert r.total_tramite("inscripcion_extemporanea") == 1
    assert r.total_tramite("reingreso") == 0
    assert r.total_tramite(est.SIN_TRAMITE) == 1
    assert r.por_programa["ing_fisica"] == {
        "cancelacion": 2,
        "inscripcion_extemporanea": 1,
    }
    assert r.casos_programa["ing_fisica"] == 2
    assert r.casos_programa[est.SIN_PROGRAMA] == 1
    assert r.por_mes[(2026, 9)]["cancelacion"] == 2
    assert r.casos_mes[(2026, 10)] == 2
    assert r.movilidad_nivel["posgrado"][SentidoConsejo.APROBO] == 1
    assert "pregrado" not in r.movilidad_nivel

    ano = await est.calcular(db, est.periodo_de("2026", HOY), HOY)
    assert ano.total_casos == 6 and ano.total_tramite("reingreso") == 1


async def test_excel(db: aiosqlite.Connection) -> None:
    await _sembrar(db)
    r = await est.calcular(db, est.periodo_de("2026-2S", HOY), HOY)
    libro = load_workbook(io.BytesIO(est.xlsx(r)))
    assert libro.sheetnames == [
        "Resumen",
        "Por programa",
        "Por mes",
        "Movilidades",
        "Léeme",
    ]
    resumen = {
        f[0]: f[1:]
        for f in libro["Resumen"].iter_rows(min_row=5, values_only=True)
        if f[0]
    }
    assert resumen["Cancelación de asignaturas"] == (1, 1, 0, 2)
    assert resumen["Movilidad"] == (1, 0, 0, 1)
    assert resumen["Total de trámites"] == (3, 2, 1, 6)
    assert resumen["Casos distintos"] == (3, 1, 1, 5)

    pp = libro["Por programa"]
    encabezado = [c.value for c in pp[1]]
    filas = {f[0]: f for f in pp.iter_rows(min_row=2, values_only=True)}
    fisica = filas["Ingeniería Física"]
    assert fisica[1] == "pregrado"
    assert fisica[encabezado.index("Cancelación de asignaturas")] == 2
    assert fisica[-2:] == (3, 2)
    assert filas["Total posgrado"][-1] == 1
    assert filas["Total Facultad"][-1] == 5

    pm = libro["Por mes"]
    assert [c.value for c in pm[1]][1:7] == [
        "Jul 2026", "Ago 2026", "Sep 2026", "Oct 2026", "Nov 2026", "Dic 2026",
    ]  # fmt: skip
    casos = next(f for f in pm.iter_rows(values_only=True) if f[0] == "Casos distintos")
    assert casos[1:] == (1, 0, 2, 2, 0, 0, 5)

    mv = list(libro["Movilidades"].iter_rows(values_only=True))
    assert ("posgrado", 1, 0, 0, 1) in mv
    assert ("Maestría en Ciencias - Física", "posgrado", 1) in [f[:3] for f in mv]

    texto = " ".join(str(f[0]) for f in libro["Léeme"].iter_rows(values_only=True))
    assert "aplazamientos" in texto and "2026-10-30" in texto
    assert "hay que tomarlas de su fuente" in texto
    # Solo cifras: ningún nombre en el archivo.
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows(values_only=True):
            assert NOMBRE_FICTICIO not in [str(v) for v in fila]


def test_periodo() -> None:
    assert est.periodo_de("", HOY).texto == "2026-2S"
    ano = est.periodo_de("2025", HOY)
    assert (ano.desde, ano.hasta) == (date(2025, 1, 1), date(2025, 12, 31))
    assert len(est.meses_de(ano)) == 12
    assert est.meses_de(est.periodo_de("2026-1s", HOY))[-1] == (2026, 6)
    with pytest.raises(PeriodoInvalido):
        est.periodo_de("2026-3S", HOY)


async def test_endpoint_y_permisos(cliente: Cliente) -> None:  # noqa: F811
    c = cliente("actas")
    ruta = "/api/web/actas/secretaria/estadisticas/xlsx"
    sec = await _token("secretaria@unal.edu.co", "secretaria")
    dec = await _token("dec@unal.edu.co", "decanatura")
    adm = await _token("adm@unal.edu.co", "administrativa", "fisica")
    r = c.get(f"{ruta}?semestre=2026-2S", headers=sec)
    assert r.status_code == 200, r.text
    assert r.content[:2] == b"PK"
    assert "estadisticas_consejo_fcen_2026-2S.xlsx" in r.headers["content-disposition"]
    assert c.get(f"{ruta}?semestre=2026", headers=dec).status_code == 200
    assert c.get(f"{ruta}?semestre=otoño", headers=dec).status_code == 422
    assert c.get(ruta, headers=adm).status_code == 403
    assert c.get(ruta).status_code == 401
