"""Cuadro de asignaturas del oficio y el acta (§5.4 de la reunión del 28-oct)."""

from __future__ import annotations

import io

from docx import Document

from app.schemas.caso import AsignaturaSolicitada, Sentido
from app.services.documento import construir_acta, construir_oficio
from app.services.ficha import ficha
from app.services.reglas import analizar
from tests.conftest import extraccion_ejemplo
from tests.test_documento import _datos


def _con_asignaturas(sentido: Sentido, asignaturas: list[AsignaturaSolicitada]):
    datos = _datos(sentido)
    item = datos.casos[0]
    item.extraccion = item.extraccion.model_copy(update={"asignaturas": asignaturas})
    item.caso.extraccion = item.extraccion
    return datos


def test_cuadro_con_grupo_y_tipologia_sin_columna_de_tramite() -> None:
    datos = _con_asignaturas(
        Sentido.APROBAR,
        [
            AsignaturaSolicitada(
                codigo="4100849",
                nombre="Termodinámica",
                grupo="2",
                tipologia="DISCIPLINAR OBLIGATORIA",
                creditos=4,
                accion="cancelar",
            ),
            AsignaturaSolicitada(
                codigo="4100850",
                nombre="Óptica",
                grupo="1",
                tipologia="FUND. OPTATIVA",
                creditos=3,
                accion="inscribir",
            ),
        ],
    )
    for contenido in (
        construir_oficio(datos, "2026-2S"),
        construir_acta(datos, "2026-2S"),
    ):
        doc = Document(io.BytesIO(contenido))
        texto = [p.text for p in doc.paragraphs]
        assert "Cancelar las siguientes asignaturas:" in texto
        assert "Inscribir las siguientes asignaturas:" in texto
        assert len(doc.tables) == 2
        for tabla in doc.tables:
            assert [c.text for c in tabla.rows[0].cells] == [
                "Código",
                "Asignatura",
                "Grupo",
                "Tipología",
                "Créditos",
            ]
        assert [c.text for c in doc.tables[0].rows[1].cells] == [
            "4100849",
            "Termodinámica",
            "2",
            "DISCIPLINAR OBLIGATORIA",
            "4",
        ]


def test_grupo_faltante_resaltado_en_amarillo_y_alerta() -> None:
    sin_grupo = AsignaturaSolicitada(
        codigo="4100849", nombre="Termodinámica", creditos=4, accion="cancelar"
    )
    datos = _con_asignaturas(Sentido.APROBAR, [sin_grupo])
    doc = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    celda = doc.tables[0].rows[1].cells[2]
    assert celda.text == ""
    assert 'w:fill="FFFF00"' in celda._tc.xml

    e = datos.casos[0].extraccion
    analisis = analizar(e, datos.programa, periodo_inicio=None, hay_imagenes=False)
    assert any("Falta el grupo" in a.mensaje for a in analisis.alertas)
    assert "Falta el grupo de Termodinámica" in ficha(datos.casos[0].caso, analisis)


def test_con_grupo_no_hay_alerta() -> None:
    e = extraccion_ejemplo()
    e.asignaturas[0].grupo = "1"
    a = analizar(e, None, periodo_inicio=None, hay_imagenes=False)
    assert not any("Falta el grupo" in x.mensaje for x in a.alertas)


def test_zapc_va_en_la_frase_y_no_en_el_cuadro() -> None:
    datos = _con_asignaturas(
        Sentido.APROBAR,
        [
            AsignaturaSolicitada(nombre="Componente ZAPC", accion="apertura"),
            AsignaturaSolicitada(
                codigo="4101167",
                nombre="Mecánica estadística avanzada",
                grupo="1",
                creditos=4,
            ),
        ],
    )
    item = datos.casos[0]
    item.extraccion = item.extraccion.model_copy(
        update={"modalidad_trabajo_grado": "asignaturas de posgrado"}
    )
    doc = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "y, en consecuencia, la apertura del componente ZAPC." in texto
    assert "Apertura de los siguientes componentes" not in texto
    assert all(
        "ZAPC" not in c.text for t in doc.tables for f in t.rows for c in f.cells
    )


def test_sin_posgrado_no_hay_frase_de_zapc() -> None:
    datos = _con_asignaturas(
        Sentido.APROBAR,
        [AsignaturaSolicitada(nombre="Componente ZAPC", accion="apertura")],
    )
    doc = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    assert "ZAPC" not in "\n".join(p.text for p in doc.paragraphs)
    assert not doc.tables


def test_documentos_sin_pantallazos() -> None:
    datos = _datos(Sentido.APROBAR)
    for contenido in (
        construir_oficio(datos, "2026-2S"),
        construir_acta(datos, "2026-2S"),
    ):
        doc = Document(io.BytesIO(contenido))
        assert not doc.inline_shapes
        assert not any("media/" in r.target_ref for r in doc.part.rels.values())


def test_oficio_cita_la_norma_y_el_acta_la_transcribe() -> None:
    datos = _datos(Sentido.APROBAR)
    oficio = "\n".join(
        p.text
        for p in Document(io.BytesIO(construir_oficio(datos, "2026-2S"))).paragraphs
    )
    acta = "\n".join(
        p.text
        for p in Document(io.BytesIO(construir_acta(datos, "2026-2S"))).paragraphs
    )
    assert "Acuerdo 008 de 2008" in oficio
    assert "En cada período académico los estudiantes" not in oficio
    assert "En cada período académico los estudiantes" in acta
