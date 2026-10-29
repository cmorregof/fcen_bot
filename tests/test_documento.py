from __future__ import annotations

import io
import zipfile
from datetime import date

from docx import Document

from app.core.programas import PROGRAMAS
from app.schemas.caso import Caso, Concepto, EstadoCaso, Sentido
from app.services.documento import (
    CasoDocumento,
    DatosOficio,
    construir_acta,
    construir_oficio,
    creditos_texto,
    en_letras,
    fecha_larga,
    titulo_seccion,
)
from app.services.reglas import analizar
from tests.conftest import extraccion_ejemplo


def _datos(sentido: Sentido, conceptos: list[Concepto] | None = None) -> DatosOficio:
    programa = PROGRAMAS["ing_fisica"]
    e = extraccion_ejemplo()
    caso = Caso(
        id=1,
        chat_id=1,
        direccion="ciencias_naturales",
        programa="ing_fisica",
        estado=EstadoCaso.LISTO,
        sentido=sentido,
        texto="x",
        extraccion=e,
        conceptos=conceptos or [],
    )
    item = CasoDocumento(
        caso,
        e,
        analizar(e, programa, periodo_inicio=None, hay_imagenes=False),
        "Buenos días,\n\nSolicito la cancelación.\n\nPersona Ficticia Uno",
    )
    return DatosOficio(
        programa=programa,
        fecha=date(2026, 10, 13),
        casos=[item],
        elaboro="AB",
        miembros=programa.miembros,
    )


def _oficio(sentido: Sentido) -> bytes:
    return construir_oficio(_datos(sentido), "2026-2S")


def _texto_completo(docx: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx)) as z:
        return " ".join(
            z.read(n).decode("utf-8")
            for n in z.namelist()
            if n.startswith("word/") and n.endswith(".xml")
        )


def test_marcadores_reemplazados() -> None:
    xml = _texto_completo(_oficio(Sentido.APROBAR))
    assert "{{" not in xml
    assert "Programa Curricular de Ingeniería Física" in xml
    assert "caringfis_man@unal.edu.co" in xml


def test_cuerpo_del_oficio() -> None:
    doc = Document(io.BytesIO(_oficio(Sentido.APROBAR)))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "CONSEJO DE FACULTAD" in texto
    assert "SOLICITUD DE CARGA MÍNIMA Y CANCELACIÓN DE ASIGNATURAS" in texto
    # En el oficio la recomendación va primero, con el número del caso.
    assert (
        "2.1 RECOMENDACIÓN: El Comité Asesor recomienda al Consejo de Facultad "
        "APROBAR la solicitud de cancelación" in texto
    )
    assert "PERSONA FICTICIA UNO identificada con DNI No. 1000000001" in texto
    assert texto.rstrip().count("a la estudiante Persona Ficticia Uno.") >= 1
    assert "tiene inscritos trece créditos (13 créditos)" in texto
    assert "Acuerdo 008 de 2008" in texto
    assert "Cancelar las siguientes asignaturas:" in texto
    # En el oficio la norma se cita, no se transcribe (§5.4)
    assert "En cada período académico los estudiantes" not in texto
    [tabla] = doc.tables
    assert [c.text for c in tabla.rows[0].cells] == [
        "Código",
        "Asignatura",
        "Grupo",
        "Tipología",
        "Créditos",
    ]
    assert [c.text for c in tabla.rows[1].cells] == [
        "4100849",
        "Termodinámica",
        "",
        "",
        "4",
    ]
    estilos = {p.style.name for p in doc.paragraphs if p.style is not None}
    assert {"FCEN Caso", "FCEN Cita", "FCEN Titulo", "FCEN Fecha"} <= estilos


def test_sentido_pendiente_queda_marcado() -> None:
    doc = Document(io.BytesIO(_oficio(Sentido.PENDIENTE)))
    texto = "\n".join(p.text for p in doc.paragraphs)
    assert "[APROBAR / NO APROBAR]" in texto


def test_numeros_en_letras() -> None:
    assert en_letras(1) == "un"
    assert en_letras(13) == "trece"
    assert en_letras(21) == "veintiún"
    assert en_letras(31) == "treinta y un"
    assert en_letras(47) == "cuarenta y siete"
    assert creditos_texto(1) == "un crédito (1 crédito)"


def test_fecha_larga() -> None:
    assert fecha_larga("2026-09-02") == "02 de septiembre de 2026"
    assert fecha_larga(None) == "____"


def test_titulo_seccion() -> None:
    assert (
        titulo_seccion(extraccion_ejemplo())
        == "SOLICITUD DE CARGA MÍNIMA Y CANCELACIÓN DE ASIGNATURAS"
    )


def test_plantilla_sin_datos_personales() -> None:
    from app.services.documento import PLANTILLA

    xml = _texto_completo(PLANTILLA.read_bytes())
    import re

    cuerpo = re.findall(r"<w:t[^>]*>([^<]+)</w:t>", xml)
    assert not any(re.search(r"\d{8,10}", t) for t in cuerpo)
    with zipfile.ZipFile(PLANTILLA) as z:
        assert not any(n.startswith("word/fonts/") for n in z.namelist())


def test_acta_con_sesion_y_conceptos_y_oficio_sin_ellos() -> None:
    conceptos = [
        Concepto(
            miembro="Elisabeth Restrepo Parra",
            tratamiento="La profesora",
            texto="Recomiendo aprobar.\n\nCumple los requisitos.",
        )
    ]
    datos = _datos(Sentido.APROBAR, conceptos)
    acta = Document(io.BytesIO(construir_acta(datos, "2026-2S")))
    texto = "\n".join(p.text for p in acta.paragraphs)
    assert "COMITÉ ASESOR PROGRAMA CURRICULAR INGENIERÍA FÍSICA" in texto
    assert "ORDEN DEL DÍA" in texto and "2.1 Persona Ficticia Uno" in texto
    assert "Jairo de Jesús Agudelo Calle, Representante Profesoral." in texto
    assert (
        "La profesora Elisabeth Restrepo Parra expresa: “Recomiendo aprobar." in texto
    )
    assert "Cumple los requisitos.”." in texto
    assert (
        "RECOMENDACIÓN: El Comité Asesor recomienda al Consejo de Facultad APROBAR"
        in texto
    )
    assert "CONSEJO DE FACULTAD" not in texto  # el acta no es una carta

    oficio = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    texto_oficio = "\n".join(p.text for p in oficio.paragraphs)
    assert "Recomiendo aprobar" not in texto_oficio
    assert "ORDEN DEL DÍA" not in texto_oficio
