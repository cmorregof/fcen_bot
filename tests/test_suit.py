"""Datos SUIT de la Secretaría de Facultad. Datos ficticios."""

from __future__ import annotations

import io
import json
from datetime import date

import aiosqlite
import pytest
from openpyxl import load_workbook

from app.db.session import connect
from app.services import suit


async def _caso(
    db: aiosqlite.Connection,
    tipos: list[str],
    sentido: str,
    acta_fecha: str,
    *,
    programa: str = "ing_fisica",
    instancia: str = "facultad",
    estado: str = "cerrado",
) -> None:
    cur = await db.execute(
        "INSERT INTO casos (chat_id, user_id, direccion, programa, estado, sentido,"
        " extraccion_json) VALUES (1, 1, 'fisica', ?, ?, 'aprobar', ?)",
        (programa, estado, json.dumps({"tipos": tipos})),
    )
    await db.execute(
        "INSERT INTO decisiones_consejo (caso_id, acta_id, decision_id, instancia, sentido,"
        " texto_decision, acta_numero, acta_fecha, fragmento, confirmado_por)"
        " VALUES (?, 1, 1, ?, ?, 'texto ficticio', '35', ?, 'fragmento ficticio',"
        " 'persona@unal.edu.co')",
        (cur.lastrowid, instancia, sentido, acta_fecha),
    )
    await db.commit()


@pytest.fixture
async def db():  # type: ignore[no-untyped-def]
    conexion = await connect()
    await conexion.execute(
        "INSERT INTO actas_consejo (id, instancia, numero, fecha, texto, subido_por)"
        " VALUES (1, 'facultad', '35', '2026-09-17', 'acta ficticia', 'x@unal.edu.co')"
    )
    yield conexion
    await conexion.close()


async def test_cuenta_por_numeral_y_mes(db: aiosqlite.Connection) -> None:
    await _caso(db, ["cancelacion"], "aprobo", "2026-09-17")
    await _caso(
        db, ["inscripcion_extemporanea", "cancelacion"], "no_aprobo", "2026-09-03"
    )
    await _caso(db, ["cancelacion_periodo"], "aprobo", "2026-10-08")
    await _caso(db, ["reingreso"], "otro", "2026-10-08")
    await _caso(db, ["reserva_cupo"], "aprobo", "2026-01-20")
    # No cuentan: aplazado, de otro año, del Consejo de Sede, descartado.
    await _caso(db, ["cancelacion"], "aplazo", "2026-09-17")
    await _caso(db, ["cancelacion"], "aprobo", "2025-09-17")
    await _caso(db, ["reingreso"], "aprobo", "2026-10-08", instancia="sede")
    await _caso(db, ["cancelacion"], "aprobo", "2026-09-17", estado="descartado")
    r = await suit.reporte(db, 2026, date(2026, 10, 30))
    por = {c.numeral.numero: c for c in r.conteos}
    assert por[13].total_mes(9) == 2 and por[13].total == 2
    assert por[22].total_mes(10) == 1
    assert por[6].total == 1
    assert por[9].total_mes(1) == 1
    assert por[23].total == 0


async def test_traslado_solo_pregrado(db: aiosqlite.Connection) -> None:
    await _caso(db, ["traslado"], "aprobo", "2026-03-10")
    await _caso(db, ["traslado"], "aprobo", "2026-03-10", programa="mcf")
    r = await suit.reporte(db, 2026, date(2026, 10, 30))
    assert {c.numeral.numero: c.total for c in r.conteos}[23] == 1


async def test_excel_y_texto(db: aiosqlite.Connection) -> None:
    await _caso(db, ["cancelacion"], "aprobo", "2026-09-17")
    r = await suit.reporte(db, 2026, date(2026, 10, 30))
    libro = load_workbook(io.BytesIO(suit.xlsx(r)))
    assert libro.sheetnames == ["Para copiar", "Detalle por decisión", "Léeme"]
    ws = libro["Para copiar"]
    fila13 = next(f for f in ws.iter_rows(min_row=5, values_only=True) if f[0] == 13)
    assert fila13[1] == "Registro de Asignaturas"
    assert fila13[3 + 8] == 1 and fila13[-1] == 1  # septiembre
    texto = suit.formatear(r)
    assert "13. Registro de Asignaturas</b>: 1 (Sep 1)" in texto


def test_ano() -> None:
    hoy = date(2026, 10, 30)
    assert suit.ano_de("", hoy) == 2026
    assert suit.ano_de("2025", hoy) == 2025
    with pytest.raises(suit.AnoInvalido):
        suit.ano_de("veinte", hoy)
