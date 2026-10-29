"""Designación de jurados de Trabajo de Grado. Datos ficticios.

El mínimo de jurados no es norma: es instrucción del Consejo de Facultad
(reunión del 28-oct-2026).
"""

from __future__ import annotations

import io
import re

from docx import Document

from app.core.programas import PROGRAMAS
from app.schemas.caso import Jurado, Sentido, TipoSolicitud
from app.services.documento import construir_acta, construir_oficio
from app.services.ficha import ficha
from app.services.reglas import JURADOS_MINIMO, analizar
from tests.conftest import extraccion_ejemplo
from tests.test_documento import _datos
from tests.test_lenguaje_recomendacion import _textos_docx

_OBJETO = "designación de jurados del Trabajo de Grado «Tema Ficticio»"


def _jurados(n: int) -> list[Jurado]:
    letras = "ABCDEFGH"
    return [
        Jurado(
            nombre=f"Profesor Ficticio {letras[i]}",
            vinculacion="Universidad Ficticia" if i % 2 else None,
            correo=f"ficticio{i}@ejemplo.edu.co" if i < 2 else None,
        )
        for i in range(n)
    ]


def _cambios(n: int, objeto: str = _OBJETO) -> dict[str, object]:
    return {
        "tipos": [TipoSolicitud.DESIGNACION_JURADOS],
        "objeto_solicitud": objeto,
        "asignaturas": [],
        "jurados": _jurados(n),
    }


def _datos_jurados(n: int):
    datos = _datos(Sentido.APROBAR)
    item = datos.casos[0]
    item.extraccion = item.extraccion.model_copy(update=_cambios(n))
    item.caso.extraccion = item.extraccion
    return datos


def _mensajes(n: int, *, objeto: str = _OBJETO, minimo: int = JURADOS_MINIMO):
    e = extraccion_ejemplo(**_cambios(n, objeto))
    a = analizar(
        e,
        PROGRAMAS["ing_fisica"],
        periodo_inicio=None,
        hay_imagenes=True,
        jurados_minimo=minimo,
    )
    return a.alertas


def test_minimo_por_defecto_es_seis() -> None:
    assert JURADOS_MINIMO == 6


def test_menos_de_seis_jurados_avisa_y_dice_que_no_es_norma() -> None:
    [alerta] = [a for a in _mensajes(4) if "jurado" in a.mensaje]
    assert alerta.nivel == "aviso"
    assert "trae 4 jurados" in alerta.mensaje
    assert "mínimo 6" in alerta.mensaje
    assert "instrucción del Consejo de Facultad, 28-oct-2026" in alerta.mensaje
    assert "no es norma" in alerta.mensaje


def test_seis_jurados_no_avisa() -> None:
    assert not any("jurado" in a.mensaje for a in _mensajes(6))


def test_el_minimo_es_configurable() -> None:
    assert not any("jurado" in a.mensaje for a in _mensajes(4, minimo=4))
    assert any("mínimo 8" in a.mensaje for a in _mensajes(6, minimo=8))


def test_tipo_sin_jurados_avisa() -> None:
    assert any("trae 0 jurados" in a.mensaje for a in _mensajes(0))


def test_frase_que_propone_jurados_avisa() -> None:
    for objeto in (
        "designación de los jurados propuestos",
        "aval de los jurados que sugiere",
        "jurados sugeridos por el estudiante",
    ):
        assert any(
            "propone o sugiere" in a.mensaje for a in _mensajes(6, objeto=objeto)
        ), objeto
    assert not any("propone o sugiere" in a.mensaje for a in _mensajes(6))


def test_designacion_no_activa_reglas_de_inscripcion_de_tdg() -> None:
    mensajes = [a.mensaje for a in _mensajes(6)]
    assert not any("disciplinares" in m for m in mensajes)


def test_oficio_y_acta_recomiendan_los_jurados_en_un_cuadro() -> None:
    datos = _datos_jurados(6)
    for contenido in (
        construir_oficio(datos, "2026-2S"),
        construir_acta(datos, "2026-2S"),
    ):
        doc = Document(io.BytesIO(contenido))
        parrafos = [p.text for p in doc.paragraphs]
        assert "Se recomiendan los siguientes jurados:" in parrafos
        assert "SOLICITUD DE DESIGNACIÓN DE JURADOS DE TRABAJO DE GRADO" in parrafos
        [tabla] = doc.tables
        assert [c.text for c in tabla.rows[0].cells] == [
            "Nombre",
            "Vinculación o institución",
            "Correo",
        ]
        assert len(tabla.rows) == 7
        assert [c.text for c in tabla.rows[1].cells] == [
            "Profesor Ficticio A",
            "",
            "ficticio0@ejemplo.edu.co",
        ]
        assert [c.text for c in tabla.rows[2].cells] == [
            "Profesor Ficticio B",
            "Universidad Ficticia",
            "ficticio1@ejemplo.edu.co",
        ]
        # Los jurados se recomiendan; nunca «propone» ni «sugiere».
        texto = _textos_docx(contenido)
        assert not re.search(r"propon|propues|sugier|suger", texto, re.IGNORECASE)


def test_el_documento_no_completa_jurados() -> None:
    doc = Document(io.BytesIO(construir_oficio(_datos_jurados(3), "2026-2S")))
    [tabla] = doc.tables
    assert len(tabla.rows) == 4


def test_ficha_muestra_los_jurados() -> None:
    datos = _datos_jurados(2)
    item = datos.casos[0]
    texto = ficha(item.caso, item.analisis)
    assert "Jurados que se recomiendan (2):" in texto
    assert "· Profesor Ficticio A (ficticio0@ejemplo.edu.co)" in texto
    assert (
        "· Profesor Ficticio B (Universidad Ficticia · ficticio1@ejemplo.edu.co)"
        in texto
    )
