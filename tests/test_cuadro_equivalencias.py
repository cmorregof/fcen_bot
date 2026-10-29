"""Cuadro de equivalencias (traslado, doble titulación, homologación). Datos ficticios."""

from __future__ import annotations

import io

from docx import Document

from app.core.programas import PROGRAMAS
from app.schemas.caso import (
    AsignaturaCuadro,
    CuadroEquivalencias,
    FilaCuadro,
    Sentido,
    TipoSolicitud,
)
from app.services.documento import construir_acta, construir_oficio
from app.services.ficha import ficha
from app.services.reglas import analizar
from tests.conftest import extraccion_ejemplo
from tests.test_documento import _datos


def _cuadro(destino_completo: bool = True) -> CuadroEquivalencias:
    return CuadroEquivalencias(
        figura="equivalencia",
        origen="Programa Ficticio A",
        destino="Ingeniería Física",
        filas=[
            FilaCuadro(
                origen=AsignaturaCuadro(
                    codigo="9000001",
                    nombre="Cálculo Ficticio",
                    tipologia="FUNDAMENTACIÓN OBLIGATORIA",
                    creditos=4,
                    nota="4.2",
                ),
                destino=AsignaturaCuadro(
                    codigo="9000101" if destino_completo else None,
                    nombre="Cálculo Diferencial Ficticio",
                    tipologia="FUNDAMENTACIÓN OBLIGATORIA",
                    creditos=4,
                    nota="4.2",
                ),
            ),
            FilaCuadro(
                origen=AsignaturaCuadro(
                    codigo="9000002", nombre="Física Ficticia", creditos=3, nota="AP"
                ),
                destino=AsignaturaCuadro(
                    codigo="9000102", nombre="Mecánica Ficticia", creditos=3, nota="AP"
                ),
            ),
        ],
    )


def _datos_traslado(cuadros: list[CuadroEquivalencias]):
    datos = _datos(Sentido.APROBAR)
    item = datos.casos[0]
    item.extraccion = item.extraccion.model_copy(
        update={
            "tipos": [TipoSolicitud.TRASLADO],
            "asignaturas": [],
            "equivalencias": cuadros,
        }
    )
    item.caso.extraccion = item.extraccion
    return datos


def test_oficio_y_acta_traen_el_cuadro_de_equivalencias() -> None:
    datos = _datos_traslado([_cuadro()])
    for contenido in (
        construir_oficio(datos, "2026-2S"),
        construir_acta(datos, "2026-2S"),
    ):
        doc = Document(io.BytesIO(contenido))
        assert "Cuadro de equivalencias:" in [p.text for p in doc.paragraphs]
        assert len(doc.tables) == 1
        tabla = doc.tables[0]
        encabezado = [c.text for c in tabla.rows[0].cells]
        assert encabezado[0] == "Asignatura cursada en Programa Ficticio A"
        assert encabezado[5] == "Asignatura a equivaler en Ingeniería Física"
        assert [c.text for c in tabla.rows[1].cells][:5] == [
            "Código",
            "Asignatura",
            "Tipología",
            "Cr.",
            "Nota",
        ]
        assert [c.text for c in tabla.rows[2].cells] == [
            "9000001",
            "Cálculo Ficticio",
            "FUNDAMENTACIÓN OBLIGATORIA",
            "4",
            "4.2",
            "9000101",
            "Cálculo Diferencial Ficticio",
            "FUNDAMENTACIÓN OBLIGATORIA",
            "4",
            "4.2",
        ]
        assert len(tabla.rows) == 4


def test_codigo_de_destino_faltante_se_resalta() -> None:
    datos = _datos_traslado([_cuadro(destino_completo=False)])
    doc = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    celda = doc.tables[0].rows[2].cells[5]
    assert 'w:fill="FFFF00"' in celda._tc.xml


def test_sin_cuadro_no_hay_tabla() -> None:
    datos = _datos_traslado([])
    doc = Document(io.BytesIO(construir_oficio(datos, "2026-2S")))
    assert doc.tables == []


def test_homologacion_usa_su_propio_titulo() -> None:
    cuadro = _cuadro().model_copy(update={"figura": "homologacion"})
    doc = Document(io.BytesIO(construir_oficio(_datos_traslado([cuadro]), "2026-2S")))
    assert "Cuadro de homologaciones:" in [p.text for p in doc.paragraphs]
    assert doc.tables[0].rows[0].cells[5].text.startswith("Asignatura a homologar")


def _mensajes(**cambios) -> list[str]:
    e = extraccion_ejemplo(asignaturas=[], **cambios)
    a = analizar(e, PROGRAMAS["ing_fisica"], periodo_inicio=None, hay_imagenes=True)
    return [x.mensaje for x in a.alertas]


def test_traslado_sin_cuadro_informa() -> None:
    for tipo in (
        TipoSolicitud.TRASLADO,
        TipoSolicitud.DOBLE_TITULACION,
        TipoSolicitud.HOMOLOGACION,
    ):
        assert any(
            "no trae cuadro de equivalencias" in m for m in _mensajes(tipos=[tipo])
        )
    assert not any(
        "cuadro de equivalencias" in m
        for m in _mensajes(tipos=[TipoSolicitud.REINGRESO])
    )


def test_traslado_con_cuadro_completo_no_avisa() -> None:
    mensajes = _mensajes(tipos=[TipoSolicitud.TRASLADO], equivalencias=[_cuadro()])
    assert not any("cuadro de equivalencias" in m for m in mensajes)


def test_destino_incompleto_avisa() -> None:
    mensajes = _mensajes(
        tipos=[TipoSolicitud.TRASLADO], equivalencias=[_cuadro(destino_completo=False)]
    )
    assert any("Cálculo Ficticio" in m and "destino" in m for m in mensajes)


def test_ficha_muestra_el_cuadro() -> None:
    datos = _datos_traslado([_cuadro()])
    item = datos.casos[0]
    texto = ficha(item.caso, item.analisis)
    assert "Cuadro de equivalencia: 2 asignaturas" in texto
    assert "9000001 Cálculo Ficticio → 9000101 Cálculo Diferencial Ficticio" in texto
