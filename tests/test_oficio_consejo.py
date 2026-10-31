"""Borrador del oficio de notificación del Consejo de Facultad. Datos ficticios."""

from __future__ import annotations

import io
from typing import Any

import pytest
from docx import Document
from docx.document import Document as DocumentoDocx

from app.core.config import settings
from app.schemas.caso import (
    AsignaturaCuadro,
    AsignaturaSolicitada,
    Caso,
    CuadroEquivalencias,
    EstadoCaso,
    FilaCuadro,
    Jurado,
    ProyectoTesis,
    Sentido,
    TipoSolicitud,
)
from app.schemas.consejo import DecisionConsejo, Instancia, SentidoConsejo
from app.services import oficio_consejo
from app.services.oficio_consejo import (
    M_MODALIDAD,
    M_MOTIVACION,
    M_SECRETARIO,
    NOTIFICACION,
    RECURSOS,
    SinOficio,
    construir_oficio_consejo,
    motivo_sin_oficio,
)
from tests.conftest import extraccion_ejemplo

T = TipoSolicitud


def _decision(
    sentido: SentidoConsejo = SentidoConsejo.APROBO,
    instancia: Instancia = Instancia.FACULTAD,
    **cambios: Any,
) -> DecisionConsejo:
    datos: dict[str, Any] = {
        "caso_id": 1,
        "instancia": instancia,
        "sentido": sentido,
        "texto_decision": "Aprobado"
        if sentido == SentidoConsejo.APROBO
        else "No aprobado",
        "acta_numero": "35",
        "acta_fecha": "2026-09-17",
        "fragmento": "Caso ficticio de la Persona Ficticia Uno: decisión del Consejo.",
        "confirmado_por": "Persona Ficticia Dos",
    }
    datos.update(cambios)
    return DecisionConsejo(**datos)


def _caso(
    tipos: list[TipoSolicitud],
    sentido: SentidoConsejo | None = SentidoConsejo.APROBO,
    *,
    programa: str = "ing_fisica",
    texto: str = "Correo ficticio",
    **extraccion: Any,
) -> Caso:
    return Caso(
        id=1,
        chat_id=1,
        direccion="ciencias_naturales",
        programa=programa,
        estado=EstadoCaso.DECIDIDO_CONSEJO,
        sentido=Sentido.APROBAR,
        texto=texto,
        extraccion=extraccion_ejemplo(tipos=tipos, **extraccion),
        consejo=_decision(sentido) if sentido else None,
    )


def _doc(caso: Caso) -> DocumentoDocx:
    return Document(io.BytesIO(construir_oficio_consejo(caso)))


def _textos(doc: DocumentoDocx) -> list[str]:
    return [p.text for p in doc.paragraphs]


def _amarillos(doc: DocumentoDocx) -> list[str]:
    return [
        r.text
        for p in doc.paragraphs
        for r in p.runs
        if r.font.highlight_color is not None
    ]


def _parrafo(doc: DocumentoDocx, inicio: str) -> str:
    return next(t for t in _textos(doc) if t.startswith(inicio))


# ---------------------------------------------------------------------------
# Textos fijos y estructura
# ---------------------------------------------------------------------------


def test_textos_fijos_literales_y_orden() -> None:
    doc = _doc(_caso([T.CANCELACION]))
    textos = _textos(doc)
    assert NOTIFICACION in textos
    assert RECURSOS in textos
    assert RECURSOS.startswith(
        "De igual manera, le informo que contra esta decisión procede el recurso de "
        "reposición ante el Consejo de Facultad, y directa o subsidiariamente el de "
        "apelación ante el Consejo de Sede, "
    )
    orden = [
        "Manizales, 17 de septiembre de 2026\tMz.FCEN.1.004-___-2026",
        "Estudiante",
        "PERSONA FICTICIA UNO con DNI 1000000001",
        "Programa Curricular de Ingeniería Física",
        "Facultad de Ciencias Exactas y Naturales",
        "Universidad Nacional de Colombia",
        "Sede Manizales",
        "Asunto:\tCancelación Asignatura sin pérdida de créditos 2026-2S.",
        "Reciba un cordial saludo.",
        NOTIFICACION,
        RECURSOS,
        "Sin otro particular,",
        "Secretario(a) de Facultad",
        "Copia: \tComité Asesor Programa Curricular Ingeniería Física, "
        "Oficina de Registro y Matrícula",
    ]
    posiciones = [textos.index(t) for t in orden]
    assert posiciones == sorted(posiciones)


def test_parrafo_de_decision_aprobada() -> None:
    doc = _doc(_caso([T.CANCELACION]))
    assert _parrafo(doc, "El Consejo de la Facultad") == (
        "El Consejo de la Facultad de Ciencias Exactas y Naturales en sesión "
        f"{M_MODALIDAD} del 17 de septiembre de 2026, Acta 35, APRUEBA la solicitud "
        "de cancelación de la asignatura TERMODINÁMICA (4100849) en el periodo "
        "Académico 2026-2S, realizada a través de correo electrónico el 02/09/2026."
    )


def test_negacion_deja_la_motivacion_en_amarillo_y_el_acta_literal() -> None:
    doc = _doc(_caso([T.CANCELACION], SentidoConsejo.NO_APROBO))
    decision = _parrafo(doc, "El Consejo de la Facultad")
    assert ", Acta 35, NO APRUEBA la solicitud de" in decision
    assert decision.endswith(f"02/09/2026, {M_MOTIVACION}.")
    amarillo = _amarillos(doc)
    assert M_MOTIVACION in amarillo
    # El fragmento del acta va literal, para tomar de ahí la motivación.
    nota = next(t for t in _textos(doc) if t.startswith("[Borre este párrafo"))
    assert "“Caso ficticio de la Persona Ficticia Uno: decisión del Consejo.”" in nota


def test_lo_que_falta_sale_en_amarillo() -> None:
    caso = _caso(
        [T.CANCELACION],
        nombre_estudiante=None,
        documento=None,
        fecha_solicitud=None,
        sia=None,
        asignaturas=[
            AsignaturaSolicitada(nombre="Asignatura Ficticia", accion="cancelar")
        ],
    )
    caso = caso.model_copy(
        update={"consejo": _decision(acta_numero=None, acta_fecha=None)}
    )
    amarillo = _amarillos(_doc(caso))
    for marcador in (
        "___",  # número del oficio
        "[NOMBRE DEL ESTUDIANTE]",
        "[DNI]",
        M_MODALIDAD,
        "[fecha de la sesión]",
        "[No.]",
        "[código]",
        "[dd/mm/aaaa]",
        settings.periodo_actual,  # sin periodo del SIA: el configurado, a revisar
        M_SECRETARIO,
    ):
        assert marcador in amarillo, marcador


def test_firma_y_contacto_desde_configuracion(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(settings, "secretaria_facultad_nombre", "Persona Ficticia Tres")
    monkeypatch.setattr(settings, "secretaria_facultad_cargo", "Secretaria de Facultad")
    doc = _doc(_caso([T.CARGA_MINIMA]))
    textos = _textos(doc)
    assert "PERSONA FICTICIA TRES" in textos
    assert "Secretaria de Facultad" in textos
    assert M_SECRETARIO not in _amarillos(doc)


def test_formato_adjunto_deja_el_canal_en_amarillo() -> None:
    doc = _doc(_caso([T.CARGA_MINIMA], texto="--- Formato adjunto ---\nx"))
    assert "[correo electrónico / oficio]" in _amarillos(doc)


# ---------------------------------------------------------------------------
# Cada trámite
# ---------------------------------------------------------------------------

_CUADRO = CuadroEquivalencias(
    figura="convalidacion",
    origen="Programa Ficticio A – Sede Ficticia",
    destino="Ingeniería Física",
    filas=[
        FilaCuadro(
            origen=AsignaturaCuadro(
                codigo="9000001", nombre="Curso Ficticio", nota="4.0"
            ),
            destino=AsignaturaCuadro(
                codigo="9000101", nombre="Asignatura Ficticia", creditos=3
            ),
        )
    ],
)


@pytest.mark.parametrize(
    ("tipo", "asunto", "objeto", "registro"),
    [
        (
            T.CANCELACION,
            "Cancelación Asignatura sin pérdida de créditos 2026-2S.",
            "cancelación de la asignatura TERMODINÁMICA (4100849) en el periodo "
            "Académico 2026-2S",
            True,
        ),
        (
            T.CANCELACION_PERIODO,
            "Cancelación Periodo Académico 2026-2S.",
            "cancelación del periodo Académico 2026-2S",
            True,
        ),
        (
            T.CARGA_MINIMA,
            "Cursar menos de la carga mínima 2026-2S.",
            "cursar menos de la carga mínima en el periodo Académico 2026-2S",
            False,
        ),
        (
            T.INSCRIPCION_EXTEMPORANEA,
            "Inscripción Extemporánea de asignatura.",
            "inscripción extemporánea de la asignatura TRABAJO DE GRADO FICTICIO "
            "(9000877) en el periodo Académico 2026-2S",
            True,
        ),
        (
            T.TRASLADO,
            "Respuesta a solicitud de Traslado 2026-2S.",
            "traslado del Programa Ficticio A – Sede Ficticia al Programa Curricular "
            "de Ingeniería Física – Sede Manizales en el periodo Académico 2026-2S",
            True,
        ),
        (
            T.PERMISO_ACADEMICO,
            "Permiso Académico.",
            "permiso académico para asistir a un evento ficticio",
            False,
        ),
        (
            T.HOMOLOGACION,
            "Convalidación asignatura.",
            "convalidación de las asignaturas que se relacionan a continuación",
            True,
        ),
        (
            T.DESIGNACION_JURADOS,
            "Designación de Jurados – Trabajo de Grado.",
            "designación de los jurados que se relacionan a continuación",
            True,
        ),
        (
            T.PROYECTO_TESIS,
            "Inscripción de proyecto de tesis o trabajo final.",
            "inscripción del proyecto de tesis o trabajo final",
            True,
        ),
        (
            T.REEXPEDICION_RECIBO,
            "Reexpedición de recibo de pago 2026-2S.",
            "reexpedición del recibo de pago en el periodo Académico 2026-2S",
            True,
        ),
    ],
)
def test_cada_tramite(
    tipo: TipoSolicitud, asunto: str, objeto: str, registro: bool
) -> None:
    caso = _caso(
        [tipo],
        objeto_solicitud="permiso académico para asistir a un evento ficticio",
        asignaturas=[
            AsignaturaSolicitada(
                codigo="4100849", nombre="Termodinámica", accion="cancelar"
            ),
            AsignaturaSolicitada(
                codigo="9000877", nombre="Trabajo de Grado Ficticio", accion="inscribir"
            ),
        ],
        equivalencias=[_CUADRO],
        jurados=[Jurado(nombre="Persona Ficticia Cuatro", vinculacion="Profesor")],
        proyecto=ProyectoTesis(
            titulo="Título ficticio", director="Persona Ficticia Cinco"
        ),
    )
    doc = _doc(caso)
    assert f"Asunto:\t{asunto}" in _textos(doc)
    decision = _parrafo(doc, "El Consejo de la Facultad")
    assert f"APRUEBA la solicitud de {objeto}, realizada a través de" in decision
    copia = _parrafo(doc, "Copia:")
    assert copia.endswith("Oficina de Registro y Matrícula") == registro


def test_cuadros_solo_si_aprobo() -> None:
    kw: dict[str, Any] = {
        "equivalencias": [_CUADRO],
        "jurados": [Jurado(nombre="Persona Ficticia Cuatro")],
        "proyecto": ProyectoTesis(
            titulo="Título ficticio", director="Persona Ficticia"
        ),
    }
    tipos = [T.HOMOLOGACION, T.DESIGNACION_JURADOS, T.PROYECTO_TESIS]
    aprobado = _doc(_caso(tipos, **kw))
    assert len(aprobado.tables) == 3
    assert "Cuadro de convalidaciones:" in _textos(aprobado)
    assert "Jurados designados:" in _textos(aprobado)
    negado = _doc(_caso(tipos, SentidoConsejo.NO_APROBO, **kw))
    assert len(negado.tables) == 0


def test_varios_tramites_en_un_oficio() -> None:
    doc = _doc(_caso([T.CANCELACION, T.CARGA_MINIMA]))
    assert (
        "Asunto:\tCancelación Asignatura sin pérdida de créditos y Cursar menos de "
        "la carga mínima 2026-2S."
    ) in _textos(doc)
    decision = _parrafo(doc, "El Consejo de la Facultad")
    # El periodo va una sola vez.
    assert decision.count("periodo Académico") == 1


def test_reexpedicion_trae_el_desbloqueo_en_amarillo() -> None:
    doc = _doc(_caso([T.REEXPEDICION_RECIBO]))
    p = next(
        p
        for p in doc.paragraphs
        if p.text.startswith("En consecuencia, se AUTORIZA el desbloqueo")
    )
    assert all(r.font.highlight_color is not None for r in p.runs)


def test_tramite_no_cubierto_va_en_amarillo() -> None:
    doc = _doc(
        _caso([T.DOBLE_TITULACION], objeto_solicitud="doble titulación ficticia")
    )
    p = next(
        p for p in doc.paragraphs if p.text.startswith("El Consejo de la Facultad")
    )
    assert all(r.font.highlight_color is not None for r in p.runs)
    assert "Doble titulación." in _amarillos(doc)
    assert "[, Oficina de Registro y Matrícula]" in _amarillos(doc)


def test_tramite_cubierto_con_otro_no_cubierto_lo_avisa() -> None:
    doc = _doc(_caso([T.CANCELACION, T.MOVILIDAD]))
    assert any("El caso incluye además: movilidad" in a for a in _amarillos(doc))


# ---------------------------------------------------------------------------
# Resolución de reingreso
# ---------------------------------------------------------------------------


def test_reingreso_aprobado_es_resolucion() -> None:
    caso = _caso([T.REINGRESO], tratamiento="El estudiante")
    assert oficio_consejo.es_resolucion(caso)
    assert oficio_consejo.nombre_archivo(caso).startswith("Borrador_resolucion")
    doc = _doc(caso)
    textos = _textos(doc)
    assert "RESOLUCIÓN [Res.CFCEN-___-AAAA]" in textos
    assert "(17 DE SEPTIEMBRE, ACTA 35)" in textos
    assert "CONSIDERANDO" in textos
    assert "RESUELVE" in textos
    assert "NOTIFÍQUESE Y CÚMPLASE" in textos
    assert "Dada en Manizales, a diecisiete (17) de septiembre de 2026" in textos
    assert any(
        t.startswith("ARTÍCULO 1. Autorizar reingreso al exestudiante") for t in textos
    )
    assert (
        "ARTÍCULO 3. La presente Resolución rige a partir de la fecha de su expedición."
        in textos
    )
    # Las citas normativas salen para revisión.
    amarillo = _amarillos(doc)
    for considerando in oficio_consejo.CONSIDERANDOS_REINGRESO:
        assert considerando in amarillo
    assert "[Nombre del Presidente(a) del Consejo de Facultad]" in amarillo


def test_reingreso_negado_es_oficio() -> None:
    caso = _caso([T.REINGRESO], SentidoConsejo.NO_APROBO)
    assert not oficio_consejo.es_resolucion(caso)
    doc = _doc(caso)
    assert "Asunto:\tRespuesta a solicitud de Reingreso 2026-2S." in _textos(doc)


# ---------------------------------------------------------------------------
# Solo con decisión final del Consejo de Facultad
# ---------------------------------------------------------------------------


def test_sin_decision_no_hay_oficio() -> None:
    caso = _caso([T.CANCELACION], None)
    assert motivo_sin_oficio(caso)
    with pytest.raises(SinOficio):
        construir_oficio_consejo(caso)


@pytest.mark.parametrize("sentido", [SentidoConsejo.APLAZO, SentidoConsejo.OTRO])
def test_aplazado_u_otro_no_tiene_oficio(sentido: SentidoConsejo) -> None:
    caso = _caso([T.CANCELACION], sentido)
    assert motivo_sin_oficio(caso)
    with pytest.raises(SinOficio):
        construir_oficio_consejo(caso)


def test_decision_de_otra_instancia_no_tiene_oficio() -> None:
    caso = _caso([T.CANCELACION]).model_copy(
        update={"consejo": _decision(instancia=Instancia.SEDE)}
    )
    with pytest.raises(SinOficio):
        construir_oficio_consejo(caso)


def test_no_llama_al_modelo(monkeypatch: pytest.MonkeyPatch) -> None:
    import app.services.extraccion as extraccion

    def prohibido(*a: Any, **kw: Any) -> None:
        raise AssertionError("el oficio del Consejo no usa el modelo")

    for nombre in dir(extraccion):
        if nombre.startswith(("extraer", "corregir")):
            monkeypatch.setattr(extraccion, nombre, prohibido)
    construir_oficio_consejo(_caso([T.CANCELACION]))
