"""Inventario determinista de la carpeta de soportes de un plan.

Cada archivo se asigna a una acción por su ruta: la primera carpeta que se
llame «ACCION 1.3» (o «ACCIÓN 1.3») define la acción. Los archivos fuera de
esas carpetas quedan como documentos generales del plan (acción "").
"""

from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterator
from datetime import datetime
from pathlib import Path, PurePosixPath

from app.servicios.calidad.soportes.modelos import Archivo

# «ACCION 1.3» (plan 2016) o «ACCIÓN 1.2.1» (plantilla institucional).
PATRON_ACCION = re.compile(r"^\s*ACCI[OÓ]N\s*(\d+(?:\.\d+)+)", re.IGNORECASE)

# Archivos de sistema que no son soportes.
NOMBRES_OMITIDOS = {"desktop.ini", "thumbs.db", ".ds_store", "icon\r"}
EXTENSIONES_OMITIDAS = {".ini", ".db", ".lnk", ".tmp"}

TIPOS = {
    ".pdf": "pdf",
    ".png": "imagen",
    ".jpg": "imagen",
    ".jpeg": "imagen",
    ".gif": "imagen",
    ".webp": "imagen",
    ".xlsx": "xlsx",
    ".xlsm": "xlsx",
    ".docx": "docx",
    ".zip": "zip",
}


def nfc(texto: str) -> str:
    """macOS guarda los nombres en NFD; se comparan siempre en NFC."""
    return unicodedata.normalize("NFC", texto)


def omitido(nombre: str) -> bool:
    bajo = nombre.lower()
    if bajo in NOMBRES_OMITIDOS or bajo.startswith(("~$", "._")):
        return True
    return Path(bajo).suffix in EXTENSIONES_OMITIDAS


def accion_de_ruta(ruta: str) -> str:
    """'…/ACCION 1.3/sub/x.pdf' -> '1.3'. Sin carpeta de acción -> ''.

    Solo cuentan las carpetas cuyo nombre empieza por «ACCION n.n»; un archivo
    llamado «PLAN DE ACCION 2016.xlsx» no es una acción.
    """
    for parte in PurePosixPath(nfc(ruta)).parent.parts:
        m = PATRON_ACCION.search(parte)
        if m:
            return m.group(1)
    return ""


def tipo_de(nombre: str) -> tuple[str, str]:
    ext = Path(nombre).suffix.lower()
    return TIPOS.get(ext, "otro"), ext


def nombre_miembro_zip(info: zipfile.ZipInfo) -> str:
    """Los zip de Windows sin la bandera UTF-8 traen el nombre en cp437/cp850."""
    if info.flag_bits & 0x800:
        return nfc(info.filename)
    crudo = info.filename.encode("cp437", errors="replace")
    try:
        return nfc(crudo.decode("utf-8"))
    except UnicodeDecodeError:
        return nfc(crudo.decode("cp850", errors="replace"))


def _fecha_pdf(datos: bytes) -> tuple[int | None, str | None, str]:
    from pypdf import PdfReader

    try:
        lector = PdfReader(io.BytesIO(datos))
        paginas = len(lector.pages)
    except Exception as exc:  # PDF dañado o cifrado
        return None, None, f"no se pudo abrir el PDF ({type(exc).__name__})"
    fecha: str | None = None
    try:
        meta = lector.metadata
        creada = meta.creation_date if meta else None
        if creada:
            fecha = creada.date().isoformat()
    except Exception:
        fecha = None
    return paginas, fecha, ""


def _fecha_imagen(datos: bytes) -> tuple[str | None, str]:
    from PIL import Image

    try:
        with Image.open(io.BytesIO(datos)) as img:
            exif = img.getexif()
            crudo = exif.get(36867) or exif.get(306)
            ancho, alto = img.size
    except Exception as exc:
        return None, f"no se pudo abrir la imagen ({type(exc).__name__})"
    if isinstance(crudo, str) and len(crudo) >= 10:
        return crudo[:10].replace(":", "-"), f"{ancho}x{alto} px"
    return None, f"{ancho}x{alto} px"


def _fecha_office(datos: bytes, tipo: str) -> str | None:
    try:
        if tipo == "xlsx":
            from openpyxl import load_workbook

            libro = load_workbook(io.BytesIO(datos), read_only=True)
            creado = libro.properties.created
            libro.close()
        else:
            from docx import Document

            creado = Document(io.BytesIO(datos)).core_properties.created
    except Exception:
        return None
    return creado.date().isoformat() if isinstance(creado, datetime) else None


def _completar_metadatos(archivo: Archivo, datos: bytes) -> None:
    archivo.sha256 = hashlib.sha256(datos).hexdigest()
    if archivo.tipo == "pdf":
        paginas, fecha, nota = _fecha_pdf(datos)
        archivo.paginas, archivo.nota = paginas, nota
        if fecha:
            archivo.fecha, archivo.fuente_fecha = fecha, "metadato PDF"
    elif archivo.tipo == "imagen":
        fecha, nota = _fecha_imagen(datos)
        archivo.paginas = 1
        archivo.nota = nota
        if fecha:
            archivo.fecha, archivo.fuente_fecha = fecha, "EXIF"
    elif archivo.tipo in ("xlsx", "docx"):
        fecha = _fecha_office(datos, archivo.tipo)
        if fecha:
            archivo.fecha, archivo.fuente_fecha = fecha, "propiedades del documento"
        if archivo.tipo == "xlsx":
            from app.servicios.calidad.soportes.contenido import conteos_xlsx

            try:
                archivo.conteos = conteos_xlsx(datos)
            except Exception:
                archivo.conteos = ""


def _recorrer(carpeta: Path) -> Iterator[Path]:
    # Los enlaces simbólicos no se siguen (rglob tampoco entra en carpetas
    # enlazadas): podrían traer archivos de fuera del plan y mandarlos al modelo.
    for ruta in sorted(carpeta.rglob("*"), key=lambda p: nfc(str(p))):
        if ruta.is_symlink():
            continue
        if ruta.is_file() and not omitido(ruta.name):
            yield ruta


def inventariar(carpeta: Path) -> list[Archivo]:
    """Inventario completo, con los zip listados por dentro.

    Los IDs (S001, S002…) siguen el orden alfabético de las rutas, así que son
    estables entre corridas mientras la carpeta no cambie.
    """
    archivos: list[Archivo] = []

    def nuevo_id() -> str:
        return f"S{len(archivos) + 1:03d}"

    for ruta in _recorrer(carpeta):
        relativa = nfc(ruta.relative_to(carpeta).as_posix())
        tipo, ext = tipo_de(ruta.name)
        datos = ruta.read_bytes()
        archivo = Archivo(
            id=nuevo_id(),
            ruta=relativa,
            nombre=nfc(ruta.name),
            accion=accion_de_ruta(relativa),
            tipo=tipo,
            extension=ext,
            tamano=len(datos),
        )
        _completar_metadatos(archivo, datos)
        archivos.append(archivo)
        if tipo != "zip":
            continue
        try:
            zf = zipfile.ZipFile(io.BytesIO(datos))
        except zipfile.BadZipFile:
            archivo.nota = "zip dañado"
            continue
        with zf:
            for info in sorted(zf.infolist(), key=lambda i: i.filename):
                nombre_m = nombre_miembro_zip(info)
                base = PurePosixPath(nombre_m).name
                if info.is_dir() or "__MACOSX" in nombre_m or omitido(base):
                    continue
                archivo.miembros.append(nombre_m)
                tipo_m, ext_m = tipo_de(base)
                datos_m = zf.read(info)
                miembro = Archivo(
                    id=nuevo_id(),
                    ruta=f"{relativa}!/{nombre_m}",
                    nombre=base,
                    # Dentro del zip también cuentan las carpetas «ACCION n.n», y
                    # un zip llamado «ACCION 1.3.zip» es la carpeta de esa acción.
                    accion=archivo.accion
                    or accion_de_ruta(nombre_m)
                    or accion_de_ruta(f"{Path(archivo.nombre).stem}/{nombre_m}"),
                    tipo=tipo_m,
                    extension=ext_m,
                    tamano=len(datos_m),
                    dentro_de_zip=archivo.id,
                )
                _completar_metadatos(miembro, datos_m)
                if tipo_m == "zip":
                    miembro.nota = "zip dentro de un zip: no se abre"
                archivos.append(miembro)
    return archivos


def leer_bytes(carpeta: Path, archivo: Archivo) -> bytes:
    """Bytes del archivo, incluso si está dentro de un zip."""
    if "!/" not in archivo.ruta:
        return (carpeta / archivo.ruta).read_bytes()
    ruta_zip, miembro = archivo.ruta.split("!/", 1)
    with zipfile.ZipFile(carpeta / ruta_zip) as zf:
        for info in zf.infolist():
            if nombre_miembro_zip(info) == miembro:
                return zf.read(info)
    raise FileNotFoundError(archivo.ruta)


def por_id(archivos: list[Archivo]) -> dict[str, Archivo]:
    return {a.id: a for a in archivos}
