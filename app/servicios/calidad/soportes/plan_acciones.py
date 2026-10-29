"""Lector de un plan de mejoramiento de una sola hoja (formato 2016).

Columnas esperadas en la fila de encabezado: «Aspectos a mejorar…»,
«Objetivo», «Acciones», «Fechas Ejecución» (Inicio/Fin) e «Indicadores de
cumplimiento». La columna de responsables (nombres y correos) no se lee.

Las acciones se numeran por aspecto: la tercera acción del cuarto aspecto es
la 4.3, igual que las carpetas «ACCION 4.3» de los soportes.
"""

from __future__ import annotations

import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from app.servicios.calidad.soportes.modelos import Accion


def _clave(valor: Any) -> str:
    texto = unicodedata.normalize("NFKD", str(valor or "")).lower()
    return "".join(c for c in texto if not unicodedata.combining(c)).strip()


def _texto(valor: Any) -> str:
    if valor is None:
        return ""
    if isinstance(valor, datetime | date):
        return valor.strftime("%Y-%m-%d")
    return " ".join(str(valor).split())


def leer_acciones(ruta: Path) -> list[Accion]:
    libro = load_workbook(ruta, data_only=True)
    hoja = libro.worksheets[0]
    filas = [list(f) for f in hoja.iter_rows(values_only=True)]

    columnas: dict[str, int] = {}
    fila_enc = -1
    for i, fila in enumerate(filas[:30]):
        claves = [_clave(c) for c in fila]
        if any(k.startswith("acciones") for k in claves) and any(
            k.startswith("indicador") for k in claves
        ):
            fila_enc = i
            for j, k in enumerate(claves):
                if k.startswith("aspecto"):
                    columnas["aspecto"] = j
                elif k.startswith("objetivo"):
                    columnas["objetivo"] = j
                elif k.startswith("acciones"):
                    columnas["accion"] = j
                elif k.startswith("fecha"):
                    columnas["inicio"] = j
                    columnas["fin"] = j + 1
                elif k.startswith("indicador"):
                    columnas["indicador"] = j
            break
    if fila_enc < 0 or "accion" not in columnas:
        raise ValueError("No encontré la fila de encabezado del plan (Acciones…).")

    def celda(fila: list[Any], nombre: str) -> str:
        j = columnas.get(nombre)
        return _texto(fila[j]) if j is not None and j < len(fila) else ""

    acciones: list[Accion] = []
    n_aspecto = 0
    n_accion = 0
    actual = {"aspecto": "", "objetivo": "", "inicio": "", "fin": ""}
    for fila in filas[fila_enc + 1 :]:
        texto_accion = celda(fila, "accion")
        if not texto_accion:
            continue
        if celda(fila, "accion").lower() in ("inicio", "fin"):
            continue
        aspecto = celda(fila, "aspecto")
        if aspecto:
            n_aspecto += 1
            n_accion = 0
            actual["aspecto"] = aspecto
            actual["objetivo"] = celda(fila, "objetivo")
        if n_aspecto == 0:
            n_aspecto = 1
        for campo in ("inicio", "fin"):
            valor = celda(fila, campo)
            if valor and valor.lower() not in ("inicio", "fin"):
                actual[campo] = valor
        n_accion += 1
        acciones.append(
            Accion(
                codigo=f"{n_aspecto}.{n_accion}",
                aspecto=actual["aspecto"],
                objetivo=actual["objetivo"],
                descripcion=texto_accion,
                indicador=celda(fila, "indicador"),
                inicio=actual["inicio"],
                fin=actual["fin"],
            )
        )
    return acciones
