"""Salidas del borrador de avance: un Word para leer y un Excel para decidir."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.servicios.calidad.avance.borrador import (
    SIN_EVIDENCIA_NUEVA,
    AccionBorrador,
    Borrador,
)
from app.servicios.calidad.soportes.salidas import (
    ENCABEZADO_BORRADOR,
    _ajustar,
    _fila,
    _hoja,
    _leeme,
    limpio,
)

DECISIONES = "Aprobado,Editar,Descartar"
AVISO = (
    "Cada frase cita su fuente: [P] es la hoja y fila del plan, [C] un corte "
    "anterior, [E] un enlace de evidencia registrado y [S] un soporte que el "
    "revisor describió. Las frases sin fuente verificada se marcan con ⚠ y no se "
    "afirman. El borrador es una propuesta: la oficina de calidad lo aprueba, lo "
    "edita o lo descarta."
)


def _fuentes(a: AccionBorrador) -> list[str]:
    catalogo = a.base.por_id()
    return [f"[{c}] {catalogo[c].ref}" for c in a.citas if c in catalogo]


def borrador_docx(ruta: Path, b: Borrador) -> None:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    aviso = doc.add_paragraph().add_run(ENCABEZADO_BORRADOR)
    aviso.bold = True
    aviso.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    doc.sections[0].header.paragraphs[0].text = ENCABEZADO_BORRADOR

    doc.add_heading(
        limpio(f"Descripción de avance del corte {b.corte} — {b.programa}"), 1
    )
    doc.add_paragraph(AVISO)
    r = b.resumen()
    doc.add_paragraph(
        f"Acciones: {r['acciones']}. Con propuesta: {r['con_propuesta']}; "
        f"sin evidencia nueva en el corte: {r['sin_evidencia_nueva']}; sin propuesta "
        f"del modelo: {r['sin_propuesta']}; frases sin fuente verificada: "
        f"{r['frases_rechazadas']}."
    )
    if b.detenido:
        doc.add_paragraph(f"⚠ El borrador no se completó ({b.detenido}).")

    objetivo_actual = ""
    for a in b.acciones:
        ac = a.base
        if ac.objetivo.codigo != objetivo_actual:
            objetivo_actual = ac.objetivo.codigo
            doc.add_heading(
                limpio(f"Objetivo {ac.objetivo.codigo}. {ac.objetivo_texto}"), 2
            )
        doc.add_heading(limpio(f"Acción {ac.accion.codigo}. {ac.descripcion}"), 3)
        meta = doc.add_paragraph()
        meta.add_run("Avance declarado en el plan: ").bold = True
        avance = ac.accion.avance
        meta.add_run("sin dato" if avance is None else f"{avance:.0f} %")
        if a.sin_evidencia_nueva:
            doc.add_paragraph().add_run(SIN_EVIDENCIA_NUEVA).bold = True
        elif a.parrafo:
            doc.add_paragraph(limpio(a.parrafo))
        elif a.error:
            doc.add_paragraph(f"Sin propuesta del modelo: {limpio(a.error)}")
        fuentes = _fuentes(a)
        if fuentes:
            doc.add_paragraph().add_run("Fuentes:").italic = True
            for fuente in fuentes:
                doc.add_paragraph(limpio(fuente), style="List Bullet")
        for f in a.rechazadas:
            doc.add_paragraph(
                limpio(f"⚠ Sin fuente verificada, no se afirma ({f.motivo}): {f.texto}")
            )
    ruta.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(ruta))


def revision_xlsx(ruta: Path, b: Borrador) -> None:
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.title = "Revisión del corte"
    encabezado = [
        "Acción", "Objetivo", "Descripción de la acción", "Avance declarado (%)",
        "Borrador (propuesta)", "Fuentes", "Sin fuente (no se afirma)", "Alertas",
        "Descripción ya escrita en el plan para este corte",
        "Decisión (Aprobado/Editar/Descartar)", "Texto final", "Observaciones",
    ]  # fmt: skip
    anchos = [8, 30, 40, 11, 70, 50, 40, 30, 40, 16, 60, 30]
    _hoja(ws, encabezado, anchos, humanas=3)
    for a in b.acciones:
        ac = a.base
        _fila(
            ws,
            [
                ac.accion.codigo,
                f"{ac.objetivo.codigo}. {ac.objetivo_texto}",
                ac.descripcion,
                ac.accion.avance,
                SIN_EVIDENCIA_NUEVA if a.sin_evidencia_nueva else a.parrafo,
                "\n".join(_fuentes(a)),
                "\n".join(f"({f.motivo}) {f.texto}" for f in a.rechazadas),
                "\n".join(a.alertas),
                ac.existente,
                "",
                "",
                "",
            ],
        )
    validacion = DataValidation(
        type="list", formula1=f'"{DECISIONES}"', allow_blank=True
    )
    ws.add_data_validation(validacion)
    validacion.add(f"J2:J{max(2, ws.max_row)}")
    _ajustar(ws)
    _leeme(
        libro,
        f"Revisión del borrador de avance — corte {b.corte}",
        [
            AVISO,
            "En «Decisión» elija Aprobado (el borrador queda como está), Editar "
            "(escriba el texto en «Texto final») o Descartar.",
            "«Sin evidencia nueva en el corte»: la acción no tiene tareas "
            "terminadas, enlaces registrados ni soportes descritos desde el corte "
            "anterior. El borrador no afirma avance.",
        ],
    )
    ruta.parent.mkdir(parents=True, exist_ok=True)
    libro.save(ruta)
