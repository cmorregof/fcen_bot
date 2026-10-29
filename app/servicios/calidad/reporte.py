"""Reporte en Excel: estado del plan y pestaña «fuentes» con los registros."""

from __future__ import annotations

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font

from app.servicios.calidad.plan import Tarea
from app.servicios.calidad.textos import ORDEN, PLURAL, agrupar


def _encabezado(hoja: object, columnas: list[str]) -> None:
    hoja.append(columnas)  # type: ignore[attr-defined]
    for celda in hoja[1]:  # type: ignore[index]
        celda.font = Font(bold=True)


def generar_reporte(programa: str, tareas: list[Tarea], hoy: date) -> bytes:
    grupos = agrupar(tareas, hoy)
    libro = Workbook()

    resumen = libro.active
    assert resumen is not None
    resumen.title = "estado"
    _encabezado(resumen, ["Cifra", "Valor", "Fecha de corte", "Programa"])
    resumen.append(["Total de tareas", len(tareas), hoy.isoformat(), programa])
    for e in ORDEN:
        resumen.append(
            [f"Tareas {PLURAL[e]}", len(grupos[e]), hoy.isoformat(), programa]
        )

    detalle = libro.create_sheet("tareas")
    _encabezado(
        detalle,
        [
            "Código", "Tarea", "Estado", "Fecha límite", "Cargo responsable",
            "Dependencia responsable", "Evidencia esperada", "Factor CNA",
            "Eje del plan de acción", "Evidencia registrada",
        ],
    )  # fmt: skip
    for t in tareas:
        ev = t.evidencias[-1].valor if t.evidencias else ""
        detalle.append(
            [
                t.codigo, t.tarea, t.estado(hoy).value, t.fecha_limite.isoformat(),
                t.cargo, t.dependencia, t.evidencia_esperada, t.factor_cna,
                t.eje_plan_accion, ev,
            ]
        )  # fmt: skip

    fuentes = libro.create_sheet("fuentes")
    _encabezado(
        fuentes,
        [
            "Cifra", "Código de tarea", "Id registro tarea", "Id importación",
            "Estado calculado", "Fecha límite", "Id evidencia", "Evidencia",
            "Registrada por", "Registrada el",
        ],
    )  # fmt: skip
    cifras: list[tuple[str, list[Tarea]]] = [("Total de tareas", tareas)]
    cifras += [(f"Tareas {PLURAL[e]}", grupos[e]) for e in ORDEN]
    for cifra, lista in cifras:
        for t in lista:
            base = [
                cifra, t.codigo, t.id, t.importacion_id,
                t.estado(hoy).value, t.fecha_limite.isoformat(),
            ]  # fmt: skip
            if t.evidencias:
                for reg in t.evidencias:
                    fuentes.append(
                        [
                            *base,
                            reg.id,
                            reg.valor,
                            reg.registrado_nombre or reg.registrado_por,
                            reg.creado,
                        ]  # fmt: skip
                    )
            else:
                fuentes.append([*base, "", "", "", ""])
    for hoja in libro.worksheets:
        for col in hoja.columns:
            ancho = max(len(str(c.value or "")) for c in col)
            hoja.column_dimensions[col[0].column_letter].width = min(ancho + 2, 60)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()
