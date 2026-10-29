"""Dominio del plan de mejoramiento: lectura del Excel, estados e historial."""

from __future__ import annotations

import io
import json
import re
import unicodedata
from dataclasses import asdict, dataclass, field
from datetime import date, datetime
from enum import StrEnum
from typing import Any

import aiosqlite
from openpyxl import load_workbook

from app.servicios.calidad.config import calidad_settings


def plegar(texto: str) -> str:
    """Minúsculas y sin tildes, para comparar nombres escritos a mano."""
    sin = unicodedata.normalize("NFD", texto)
    sin = "".join(c for c in sin if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", sin).strip().lower()


class Estado(StrEnum):
    PENDIENTE = "pendiente"
    PROXIMA = "próxima a vencer"
    VENCIDA = "vencida"
    CUMPLIDA = "cumplida"


def calcular_estado(
    fecha_limite: date,
    cumplida: bool,
    hoy: date,
    dias_proxima: int | None = None,
) -> Estado:
    """Cumplida si tiene evidencia; vencida si la fecha ya pasó; próxima si
    faltan `dias_proxima` días o menos (4 semanas por defecto)."""
    if cumplida:
        return Estado.CUMPLIDA
    plazo = calidad_settings.dias_proxima if dias_proxima is None else dias_proxima
    dias = (fecha_limite - hoy).days
    if dias < 0:
        return Estado.VENCIDA
    if dias <= plazo:
        return Estado.PROXIMA
    return Estado.PENDIENTE


# ---------------------------------------------------------------------------
# Lectura del Excel
# ---------------------------------------------------------------------------

CAMPOS_OBLIGATORIOS = (
    "programa",
    "codigo",
    "tarea",
    "responsable_cargo",
    "responsable_dependencia",
    "fecha_limite",
    "evidencia_esperada",
)
CAMPOS_OPCIONALES = ("factor_cna", "eje_plan_accion")

_ALIAS = {
    "cargo": "responsable_cargo",
    "dependencia": "responsable_dependencia",
    "fecha_limite": "fecha_limite",
    "fecha": "fecha_limite",
    "eje": "eje_plan_accion",
    "factor": "factor_cna",
}


def _clave_columna(valor: object) -> str:
    texto = plegar(str(valor or "")).replace(" ", "_")
    return _ALIAS.get(texto, texto)


@dataclass(frozen=True)
class FilaPlan:
    programa: str
    codigo: str
    tarea: str
    dependencia: str
    cargo: str
    fecha_limite: date
    evidencia_esperada: str
    factor_cna: str = ""
    eje_plan_accion: str = ""

    @property
    def programa_clave(self) -> str:
        return plegar(self.programa)


@dataclass
class Lectura:
    filas: list[FilaPlan] = field(default_factory=list)
    errores: list[str] = field(default_factory=list)


def _como_fecha(valor: object) -> date | None:
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def leer_plan(contenido: bytes) -> Lectura:
    lectura = Lectura()
    try:
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        filas = list(libro.worksheets[0].iter_rows(values_only=True))
    except Exception:
        lectura.errores.append("No pude abrir el archivo: debe ser un .xlsx válido.")
        return lectura
    if not filas:
        lectura.errores.append("La primera hoja está vacía.")
        return lectura

    columnas = {_clave_columna(v): i for i, v in enumerate(filas[0]) if v}
    faltantes = [c for c in CAMPOS_OBLIGATORIOS if c not in columnas]
    if faltantes:
        lectura.errores.append("Faltan columnas: " + ", ".join(faltantes) + ".")
        return lectura

    def celda(fila: tuple[Any, ...], campo: str) -> Any:
        i = columnas.get(campo)
        return fila[i] if i is not None and i < len(fila) else None

    vistos: dict[tuple[str, str], int] = {}
    for n, fila in enumerate(filas[1:], start=2):
        if all(v in (None, "") for v in fila):
            continue
        textos = {
            c: str(celda(fila, c) or "").strip()
            for c in (*CAMPOS_OBLIGATORIOS, *CAMPOS_OPCIONALES)
            if c != "fecha_limite"
        }
        vacios = [
            c for c in CAMPOS_OBLIGATORIOS if c != "fecha_limite" and not textos[c]
        ]
        fecha = _como_fecha(celda(fila, "fecha_limite"))
        if fecha is None:
            vacios.append("fecha_limite (use AAAA-MM-DD o DD/MM/AAAA)")
        if vacios:
            lectura.errores.append(
                f"Fila {n}: falta o es inválido {', '.join(vacios)}."
            )
            continue
        assert fecha is not None
        nueva = FilaPlan(
            programa=textos["programa"],
            codigo=textos["codigo"].upper(),
            tarea=textos["tarea"],
            dependencia=textos["responsable_dependencia"],
            cargo=textos["responsable_cargo"],
            fecha_limite=fecha,
            evidencia_esperada=textos["evidencia_esperada"],
            factor_cna=textos["factor_cna"],
            eje_plan_accion=textos["eje_plan_accion"],
        )
        llave = (nueva.programa_clave, nueva.codigo)
        if llave in vistos:
            lectura.errores.append(
                f"Fila {n}: el código {nueva.codigo} ya aparece en la fila "
                f"{vistos[llave]} para ese programa."
            )
            continue
        vistos[llave] = n
        lectura.filas.append(nueva)
    if not lectura.filas and not lectura.errores:
        lectura.errores.append("El archivo no tiene tareas.")
    return lectura


def filas_a_json(filas: list[FilaPlan]) -> str:
    return json.dumps(
        [{**asdict(f), "fecha_limite": f.fecha_limite.isoformat()} for f in filas]
    )


def filas_de_json(texto: str) -> list[FilaPlan]:
    return [
        FilaPlan(**{**d, "fecha_limite": date.fromisoformat(d["fecha_limite"])})
        for d in json.loads(texto)
    ]


# ---------------------------------------------------------------------------
# Tareas, evidencias e importaciones
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Evidencia:
    id: int
    tarea_id: int
    tipo: str
    valor: str
    referencia: str | None
    registrado_por: int
    registrado_nombre: str | None
    creado: str


@dataclass(frozen=True)
class Tarea:
    id: int
    programa: str
    codigo: str
    tarea: str
    dependencia: str
    cargo: str
    fecha_limite: date
    evidencia_esperada: str
    factor_cna: str
    eje_plan_accion: str
    importacion_id: int | None
    evidencias: tuple[Evidencia, ...] = ()

    @property
    def cumplida(self) -> bool:
        return bool(self.evidencias)

    def estado(self, hoy: date) -> Estado:
        return calcular_estado(self.fecha_limite, self.cumplida, hoy)


_COLUMNAS_TAREA = (
    "id, programa, codigo, tarea, dependencia, cargo, fecha_limite, "
    "evidencia_esperada, factor_cna, eje_plan_accion, importacion_id"
)
_CAMPOS_COMPARABLES = (
    "tarea",
    "dependencia",
    "cargo",
    "fecha_limite",
    "evidencia_esperada",
    "factor_cna",
    "eje_plan_accion",
)


async def listar_tareas(
    db: aiosqlite.Connection, programa: str | None = None
) -> list[Tarea]:
    sql = f"SELECT {_COLUMNAS_TAREA} FROM cal_tareas"
    params: tuple[str, ...] = ()
    if programa is not None:
        sql += " WHERE programa_clave = ?"
        params = (plegar(programa),)
    cur = await db.execute(sql + " ORDER BY programa, fecha_limite, codigo", params)
    filas = await cur.fetchall()
    cur = await db.execute(
        "SELECT id, tarea_id, tipo, valor, referencia, registrado_por, "
        "registrado_nombre, creado FROM cal_evidencias ORDER BY id"
    )
    por_tarea: dict[int, list[Evidencia]] = {}
    for r in await cur.fetchall():
        ev = Evidencia(*tuple(r))
        por_tarea.setdefault(ev.tarea_id, []).append(ev)
    return [
        Tarea(
            id=r["id"],
            programa=r["programa"],
            codigo=r["codigo"],
            tarea=r["tarea"],
            dependencia=r["dependencia"],
            cargo=r["cargo"],
            fecha_limite=date.fromisoformat(r["fecha_limite"]),
            evidencia_esperada=r["evidencia_esperada"],
            factor_cna=r["factor_cna"],
            eje_plan_accion=r["eje_plan_accion"],
            importacion_id=r["importacion_id"],
            evidencias=tuple(por_tarea.get(r["id"], [])),
        )
        for r in filas
    ]


async def listar_programas(db: aiosqlite.Connection) -> list[str]:
    cur = await db.execute("SELECT DISTINCT programa FROM cal_tareas ORDER BY programa")
    return [str(r[0]) for r in await cur.fetchall()]


async def buscar_programa(db: aiosqlite.Connection, texto: str) -> str | None:
    """Coincidencia exacta o, si no, por subcadena única (sin tildes)."""
    clave = plegar(texto)
    if not clave:
        return None
    programas = await listar_programas(db)
    for p in programas:
        if plegar(p) == clave:
            return p
    parecidos = [p for p in programas if clave in plegar(p)]
    return parecidos[0] if len(parecidos) == 1 else None


async def buscar_tareas_por_codigo(
    db: aiosqlite.Connection, codigo: str
) -> list[Tarea]:
    codigo = codigo.strip().upper()
    return [t for t in await listar_tareas(db) if t.codigo == codigo]


@dataclass
class Resumen:
    nuevas: list[FilaPlan] = field(default_factory=list)
    actualizadas: list[tuple[FilaPlan, dict[str, tuple[str, str]]]] = field(
        default_factory=list
    )
    sin_cambios: int = 0
    ausentes: list[tuple[str, str]] = field(default_factory=list)


def _valores(fila: FilaPlan) -> dict[str, str]:
    return {
        "tarea": fila.tarea,
        "dependencia": fila.dependencia,
        "cargo": fila.cargo,
        "fecha_limite": fila.fecha_limite.isoformat(),
        "evidencia_esperada": fila.evidencia_esperada,
        "factor_cna": fila.factor_cna,
        "eje_plan_accion": fila.eje_plan_accion,
    }


async def resumir(db: aiosqlite.Connection, filas: list[FilaPlan]) -> Resumen:
    existentes = {(plegar(t.programa), t.codigo): t for t in await listar_tareas(db)}
    resumen = Resumen()
    llegaron: set[tuple[str, str]] = set()
    for f in filas:
        llave = (f.programa_clave, f.codigo)
        llegaron.add(llave)
        previa = existentes.get(llave)
        if previa is None:
            resumen.nuevas.append(f)
            continue
        antes = {
            "tarea": previa.tarea,
            "dependencia": previa.dependencia,
            "cargo": previa.cargo,
            "fecha_limite": previa.fecha_limite.isoformat(),
            "evidencia_esperada": previa.evidencia_esperada,
            "factor_cna": previa.factor_cna,
            "eje_plan_accion": previa.eje_plan_accion,
        }
        despues = _valores(f)
        cambios = {
            c: (antes[c], despues[c])
            for c in _CAMPOS_COMPARABLES
            if antes[c] != despues[c]
        }
        if cambios:
            resumen.actualizadas.append((f, cambios))
        else:
            resumen.sin_cambios += 1
    programas = {f.programa_clave for f in filas}
    resumen.ausentes = [
        (t.programa, t.codigo)
        for k, t in existentes.items()
        if k[0] in programas and k not in llegaron
    ]
    return resumen


async def aplicar_importacion(
    db: aiosqlite.Connection,
    filas: list[FilaPlan],
    *,
    archivo: str,
    usuario_id: int,
    creado: str,
) -> int:
    """Guarda las filas. Reimportar actualiza por (programa, código) y deja
    en cal_historial qué cambió; nunca duplica ni borra tareas."""
    resumen = await resumir(db, filas)
    cambios_por_llave = {
        (f.programa_clave, f.codigo): c for f, c in resumen.actualizadas
    }
    nuevas = {(f.programa_clave, f.codigo) for f in resumen.nuevas}
    cur = await db.execute(
        "INSERT INTO cal_importaciones (archivo, usuario_id, creadas, actualizadas, "
        "sin_cambios, creado) VALUES (?, ?, ?, ?, ?, ?)",
        (
            archivo,
            usuario_id,
            len(resumen.nuevas),
            len(resumen.actualizadas),
            resumen.sin_cambios,
            creado,
        ),
    )
    imp_id = int(cur.lastrowid or 0)
    for f in filas:
        llave = (f.programa_clave, f.codigo)
        if llave in nuevas:
            cur = await db.execute(
                "INSERT INTO cal_tareas (programa, programa_clave, codigo, tarea, "
                "dependencia, cargo, fecha_limite, evidencia_esperada, factor_cna, "
                "eje_plan_accion, importacion_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
                (
                    f.programa,
                    f.programa_clave,
                    f.codigo,
                    f.tarea,
                    f.dependencia,
                    f.cargo,
                    f.fecha_limite.isoformat(),
                    f.evidencia_esperada,
                    f.factor_cna,
                    f.eje_plan_accion,
                    imp_id,
                ),
            )
            await db.execute(
                "INSERT INTO cal_historial (tarea_id, importacion_id, accion, "
                "cambios_json) VALUES (?, ?, 'creada', '{}')",
                (cur.lastrowid, imp_id),
            )
        elif llave in cambios_por_llave:
            v = _valores(f)
            cur = await db.execute(
                "UPDATE cal_tareas SET tarea=?, dependencia=?, cargo=?, "
                "fecha_limite=?, evidencia_esperada=?, factor_cna=?, "
                "eje_plan_accion=?, importacion_id=? "
                "WHERE programa_clave=? AND codigo=? RETURNING id",
                (*(v[c] for c in _CAMPOS_COMPARABLES), imp_id, *llave),
            )
            fila = await cur.fetchone()
            await db.execute(
                "INSERT INTO cal_historial (tarea_id, importacion_id, accion, "
                "cambios_json) VALUES (?, ?, 'actualizada', ?)",
                (
                    fila["id"] if fila else 0,
                    imp_id,
                    json.dumps(cambios_por_llave[llave], ensure_ascii=False),
                ),
            )
    await db.commit()
    return imp_id


async def registrar_evidencia(
    db: aiosqlite.Connection,
    tarea_id: int,
    *,
    tipo: str,
    valor: str,
    referencia: str | None,
    usuario_id: int,
    usuario_nombre: str | None,
    creado: str,
) -> None:
    await db.execute(
        "INSERT INTO cal_evidencias (tarea_id, tipo, valor, referencia, "
        "registrado_por, registrado_nombre, creado) VALUES (?, ?, ?, ?, ?, ?, ?)",
        (tarea_id, tipo, valor, referencia, usuario_id, usuario_nombre, creado),
    )
    await db.commit()
