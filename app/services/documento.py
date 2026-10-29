"""Oficio al Consejo de Facultad y acta del comité, en .docx, con la plantilla única.

El oficio lleva solo la recomendación del comité. El acta lleva además el
encabezado de la sesión (fecha, hora, lugar, asistentes, orden del día) y los
conceptos de los integrantes del comité, copiados tal como los enviaron.

El formato es el del oficio del Comité Asesor de Ciencias de la Computación:
los estilos «FCEN …» de la plantilla copian sus párrafos (Ancízar Sans 11 pt,
interlineado 1,15, sangría francesa en cada caso, la cita del correo dentro del
mismo párrafo). Al caso se le agregan el ANÁLISIS y la norma textual que usa
Ingeniería Física. Todo el texto sale de los datos ya registrados; aquí no se
llama al modelo.
"""

from __future__ import annotations

import io
import re
from collections import OrderedDict
from dataclasses import dataclass
from datetime import date
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor
from docx.table import _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.core.programas import Miembro, Programa
from app.normativa.base import NormaCitada, normas_para
from app.schemas.caso import (
    TITULOS_TIPO,
    AsignaturaCuadro,
    AsignaturaSolicitada,
    Caso,
    Concepto,
    ExtraccionCaso,
    Sentido,
)
from app.services.reglas import Analisis, lleva_zapc

PLANTILLA = Path(__file__).resolve().parents[1] / "templates" / "oficio_base.docx"

_MESES = (
    "enero febrero marzo abril mayo junio julio agosto septiembre octubre "
    "noviembre diciembre"
).split()

_UNIDADES = (
    "cero uno dos tres cuatro cinco seis siete ocho nueve diez once doce trece "
    "catorce quince dieciséis diecisiete dieciocho diecinueve veinte veintiuno "
    "veintidós veintitrés veinticuatro veinticinco veintiséis veintisiete "
    "veintiocho veintinueve"
).split()
_DECENAS = {
    3: "treinta",
    4: "cuarenta",
    5: "cincuenta",
    6: "sesenta",
    7: "setenta",
    8: "ochenta",
    9: "noventa",
}


def en_letras(n: int) -> str:
    """Número de créditos en letras, como en las actas («trece créditos»)."""
    if 0 <= n < 30:
        palabra = _UNIDADES[n]
        return "un" if palabra == "uno" else palabra.replace("veintiuno", "veintiún")
    if 30 <= n < 100:
        decena, unidad = divmod(n, 10)
        base = _DECENAS[decena]
        if unidad == 0:
            return base
        resto = "un" if unidad == 1 else _UNIDADES[unidad]
        return f"{base} y {resto}"
    return str(n)


def creditos_texto(n: int) -> str:
    unidad = "crédito" if n == 1 else "créditos"
    return f"{en_letras(n)} {unidad} ({n} {unidad})"


def fecha_larga(valor: date | str | None) -> str:
    if valor is None:
        return "____"
    if isinstance(valor, str):
        try:
            dia = date.fromisoformat(valor[:10])
        except ValueError:
            return valor
    else:
        dia = valor
    return f"{dia.day:02d} de {_MESES[dia.month - 1]} de {dia.year}"


_NO_XML = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def limpiar(texto: str) -> str:
    """Quita caracteres que Word no acepta (saltos suaves de Outlook, etc.)."""
    return _NO_XML.sub(" ", texto)


def _num(valor: float) -> str:
    return f"{valor:.1f}".replace(".", ",")


@dataclass
class CasoDocumento:
    caso: Caso
    extraccion: ExtraccionCaso
    analisis: Analisis
    cuerpo_correo: str
    # «correo» o «formato»: por dónde llegó la solicitud.
    medio: str = "correo"


@dataclass
class DatosOficio:
    programa: Programa
    fecha: date
    casos: list[CasoDocumento]
    radicado: str | None = None
    numero_acta: str | None = None
    fecha_sesion: str | None = None
    elaboro: str = ""
    borrador: bool = True
    # Solo para el acta.
    hora: str | None = None
    lugar: str | None = None
    miembros: tuple[Miembro, ...] = ()


# ---------------------------------------------------------------------------
# Utilidades de formato
# ---------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _p(doc: DocumentoDocx, estilo: str, texto: str = "", **formato: bool) -> Paragraph:
    p = doc.add_paragraph(style=estilo)
    if texto:
        _r(p, texto, **formato)
    return p


def _r(
    p: Paragraph,
    texto: str,
    *,
    negrita: bool = False,
    cursiva: bool = False,
    subrayado: bool = False,
) -> Run:
    r = p.add_run(limpiar(texto))
    if negrita:
        r.bold = True
    if cursiva:
        r.italic = True
    if subrayado:
        r.underline = True
    return r


def _marcador(p: Paragraph, texto: str) -> Run:
    """Dato que el comité completa en Word: resaltado en amarillo."""
    r = _r(p, texto, negrita=True)
    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return r


def _nota_gris(p: Paragraph, texto: str) -> Run:
    r = _r(p, texto)
    r.font.size = Pt(8)
    r.italic = False
    r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    return r


def _vacio(doc: DocumentoDocx, estilo: str = "FCEN Destinatario", n: int = 1) -> None:
    for _ in range(n):
        doc.add_paragraph(style=estilo)


def _reemplazar_marcadores(doc: DocumentoDocx, valores: dict[str, str]) -> None:
    """Reemplaza {{MARCADOR}} en encabezados y pies (incluidos cuadros de texto)."""
    vistos: set[int] = set()
    for seccion in doc.sections:
        for parte in (
            seccion.header,
            seccion.footer,
            seccion.first_page_header,
            seccion.first_page_footer,
        ):
            elemento = parte._element
            if id(elemento) in vistos:
                continue
            vistos.add(id(elemento))
            for t in elemento.iter(f"{_W}t"):
                if t.text and "{{" in t.text:
                    for clave, valor in valores.items():
                        t.text = t.text.replace("{{" + clave + "}}", valor)


# ---------------------------------------------------------------------------
# Bloques del oficio
# ---------------------------------------------------------------------------


def _encabezado_oficio(doc: DocumentoDocx, datos: DatosOficio) -> None:
    if datos.borrador:
        p = _p(doc, "FCEN Destinatario")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = _r(p, "BORRADOR · pendiente de revisión del Comité Asesor", negrita=True)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
        _vacio(doc)

    radicado = datos.radicado or (
        f"{datos.programa.radicado_prefijo or 'Mz.FCEN.1.0__'}-___-{datos.fecha:%y}"
    )
    p = _p(doc, "FCEN Fecha", f"Manizales, {fecha_larga(datos.fecha)}")
    _r(p, f"\t[{radicado.upper()}]", negrita=True)
    _vacio(doc, n=2)

    _p(doc, "FCEN Destinatario", "Señores")
    _p(doc, "FCEN Destinatario", "CONSEJO DE FACULTAD", negrita=True)
    for linea in (
        "Facultad de Ciencias Exactas y Naturales",
        "Universidad Nacional de Colombia",
        "Sede Manizales",
    ):
        _p(doc, "FCEN Destinatario", linea)
    _vacio(doc, n=2)

    acta = datos.numero_acta or "____"
    _p(
        doc,
        "FCEN Texto",
        f"Asunto: {datos.programa.comite} Acta {acta} de {datos.fecha.year}",
    )
    _vacio(doc, "FCEN Texto")
    _p(doc, "FCEN Texto", "Reciban un cordial y respetuoso saludo.")
    _vacio(doc, "FCEN Texto")
    fecha_sesion = datos.fecha_sesion or "____"
    _p(
        doc,
        "FCEN Texto",
        f"El {datos.programa.comite}, en sesión realizada el día {fecha_sesion}, "
        f"según consta en el Acta No. {acta}, llevó a cabo el estudio, análisis y "
        "revisión de los casos relacionados a continuación. Luego de la "
        "correspondiente deliberación y de acuerdo con las disposiciones académicas "
        "y normativas vigentes, el Comité emitió las recomendaciones que se "
        "presentan en el presente documento:",
    )
    _vacio(doc, "FCEN Cuerpo")
    _p(doc, "FCEN Titulo", "CASOS ESTUDIANTILES")
    _vacio(doc, "FCEN Cuerpo")


def titulo_seccion(extraccion: ExtraccionCaso) -> str:
    if not extraccion.tipos:
        return "OTRAS SOLICITUDES"
    titulos = [TITULOS_TIPO[t] for t in extraccion.tipos]
    if len(titulos) == 1:
        return f"SOLICITUD DE {titulos[0]}".upper()
    return ("SOLICITUD DE " + ", ".join(titulos[:-1]) + " Y " + titulos[-1]).upper()


def _bloques(cuerpo: str) -> list[str]:
    return [b.strip() for b in cuerpo.replace("\r", "").split("\n\n") if b.strip()]


def _texto_con_saltos(p: Paragraph, texto: str, *, cursiva: bool) -> None:
    lineas = [ln.strip() for ln in texto.split("\n")]
    for j, linea in enumerate(lineas):
        r = _r(p, linea, cursiva=cursiva)
        if j < len(lineas) - 1:
            r.add_break()


def _bullets_analisis(item: CasoDocumento, periodo: str) -> list[str]:
    e, a = item.extraccion, item.analisis
    vinetas: list[str] = []
    sia = e.sia
    if sia and sia.plan_codigo:
        nombre = (sia.plan_nombre or "").strip().title()
        vinetas.append(f"Plan de estudios [{sia.plan_codigo}] {nombre}".strip())
    if sia and sia.convocatoria:
        vinetas.append(f"Estudiante de la convocatoria {sia.convocatoria}")
    if a.creditos_inscritos is not None:
        vinetas.append(
            f"En el periodo académico {sia.periodo if sia and sia.periodo else periodo} "
            f"tiene inscritos {creditos_texto(a.creditos_inscritos)}"
        )
    if (
        a.creditos_tras_solicitud is not None
        and a.creditos_tras_solicitud != a.creditos_inscritos
    ):
        vinetas.append(
            f"Con la solicitud quedaría con {creditos_texto(a.creditos_tras_solicitud)}"
        )
    if a.disciplinares_aprobados is not None:
        if a.disciplinares_exigidos_tdg:
            vinetas.append(
                f"En el componente disciplinar tiene aprobados {a.disciplinares_aprobados} "
                f"créditos de los {a.disciplinares_exigidos_tdg} que se requieren para "
                "inscribir Trabajo de Grado"
            )
        else:
            vinetas.append(
                f"Tiene aprobados {a.disciplinares_aprobados} créditos del componente "
                "disciplinar"
            )
    metricas = []
    if a.papa is not None:
        metricas.append(f"P.A.P.A. {_num(a.papa)}")
    if a.porcentaje_avance is not None:
        metricas.append(f"porcentaje de avance {_num(a.porcentaje_avance)} %")
    if metricas:
        texto = "; ".join(metricas)
        vinetas.append(texto[0].upper() + texto[1:])
    if a.semana is not None:
        vinetas.append(
            f"La solicitud se presentó en la semana {a.semana} del periodo académico"
        )
    if e.modalidad_trabajo_grado:
        vinetas.append(
            f"Modalidad de Trabajo de Grado solicitada: {e.modalidad_trabajo_grado}"
        )
    for anexo in e.anexos:
        vinetas.append(f"Se anexa {anexo}")
    return vinetas


def texto_recomendacion(
    programa: Programa, item: CasoDocumento
) -> tuple[str, str, str]:
    """(antes, sentido, después) para poder resaltar el sentido si está pendiente."""
    e = item.extraccion
    sentido = {
        Sentido.APROBAR: "APROBAR",
        Sentido.NO_APROBAR: "NO APROBAR",
        Sentido.PENDIENTE: "[APROBAR / NO APROBAR]",
    }[item.caso.sentido]
    objeto = e.objeto_solicitud.strip().rstrip(".") or "lo solicitado"
    destinatario = (
        "al estudiante" if e.tratamiento == "El estudiante" else "a la estudiante"
    )
    nombre = (e.nombre_estudiante or "").title()
    final = f" {destinatario} {nombre}" if nombre else ""
    zapc = (
        " y, en consecuencia, la apertura del componente ZAPC" if lleva_zapc(e) else ""
    )
    return (
        "El Comité Asesor recomienda al Consejo de Facultad ",
        sentido,
        f" la solicitud de {objeto}{final}{zapc}.",
    )


def _resaltar_celda_amarillo(celda: _Cell) -> None:
    """Resalta la celda en amarillo mediante sombreado XML."""
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFF00"/>')
    celda._tc.get_or_add_tcPr().append(shd)


def _asignaturas(doc: DocumentoDocx, item: CasoDocumento) -> None:
    """Cuadro por trámite: párrafo introductorio + tabla Código | Asignatura | Grupo | Tipología | Créditos."""
    # ZAPC no es una asignatura: va en la frase de la recomendación (§5.4).
    asignaturas = [a for a in item.extraccion.asignaturas if a.accion != "apertura"]
    if not asignaturas:
        return
    # Agrupar por acción (trámite).
    por_accion: OrderedDict[str, list[AsignaturaSolicitada]] = OrderedDict()
    for asig in asignaturas:
        por_accion.setdefault(asig.accion, []).append(asig)

    _INTRO = {
        "inscribir": "Inscribir las siguientes asignaturas:",
        "cancelar": "Cancelar las siguientes asignaturas:",
    }

    for accion, grupo_asigs in por_accion.items():
        _p(doc, "FCEN Cuerpo", _INTRO[accion])
        tabla = doc.add_table(rows=1, cols=5)
        try:
            tabla.style = doc.styles["Table Grid"]
        except KeyError:
            pass
        for celda, titulo in zip(
            tabla.rows[0].cells,
            ("Código", "Asignatura", "Grupo", "Tipología", "Créditos"),
            strict=True,
        ):
            celda.text = ""
            _r(celda.paragraphs[0], titulo, negrita=True)
        for asig in grupo_asigs:
            fila = tabla.add_row().cells
            valores = (
                asig.codigo or "",
                asig.nombre,
                asig.grupo or "",
                asig.tipologia or "",
                "" if asig.creditos is None else str(asig.creditos),
            )
            for celda, valor in zip(fila, valores, strict=True):
                celda.text = ""
                r = _r(celda.paragraphs[0], valor)
                # Grupo faltante: resaltado en amarillo de la celda y del texto.
                if celda is fila[2] and not valor:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    _resaltar_celda_amarillo(celda)
        anchos = (2.2, 6.0, 1.5, 3.5, 1.8)
        tabla.autofit = False
        for columna, ancho_col in zip(tabla.columns, anchos, strict=True):
            columna.width = Cm(ancho_col)
        for fila_t in tabla.rows:
            for celda, ancho in zip(fila_t.cells, anchos, strict=True):
                celda.width = Cm(ancho)
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.font.size = Pt(10)
        _vacio(doc, "FCEN Cuerpo")


_FIGURA = {
    "equivalencia": ("Cuadro de equivalencias", "Asignatura a equivaler"),
    "homologacion": ("Cuadro de homologaciones", "Asignatura a homologar"),
    "convalidacion": ("Cuadro de convalidaciones", "Asignatura a convalidar"),
}
_COLUMNAS_LADO = ("Código", "Asignatura", "Tipología", "Créditos", "Nota")
_ANCHOS_CUADRO = (1.4, 3.3, 1.6, 1.2, 1.0) * 2


def _lado(a: AsignaturaCuadro) -> tuple[str, ...]:
    return (
        a.codigo or "",
        a.nombre or "",
        a.tipologia or "",
        "" if a.creditos is None else str(a.creditos),
        a.nota or "",
    )


def _equivalencias(doc: DocumentoDocx, item: CasoDocumento) -> None:
    """Cuadro comparativo como el del formato de homologación, convalidación o
    equivalencia: lo cursado a la izquierda y lo que se reconoce a la derecha.
    Las celdas vacías del lado de destino se resaltan: sin ellas no se ejecuta."""
    for cuadro in item.extraccion.equivalencias:
        if not cuadro.filas:
            continue
        titulo, encabezado_destino = _FIGURA[cuadro.figura]
        _p(doc, "FCEN Cuerpo", f"{titulo}:")
        tabla = doc.add_table(rows=2, cols=10)
        try:
            tabla.style = doc.styles["Table Grid"]
        except KeyError:
            pass
        izquierda = tabla.cell(0, 0).merge(tabla.cell(0, 4))
        derecha = tabla.cell(0, 5).merge(tabla.cell(0, 9))
        for celda, texto in (
            (
                izquierda,
                "Asignatura cursada"
                + (f" en {cuadro.origen}" if cuadro.origen else ""),
            ),
            (
                derecha,
                encabezado_destino
                + (f" en {cuadro.destino}" if cuadro.destino else ""),
            ),
        ):
            celda.text = ""
            _r(celda.paragraphs[0], texto, negrita=True)
        for celda, titulo_col in zip(
            tabla.rows[1].cells, _COLUMNAS_LADO * 2, strict=True
        ):
            celda.text = ""
            _r(celda.paragraphs[0], titulo_col, negrita=True)
        for fila in cuadro.filas:
            celdas = tabla.add_row().cells
            valores = _lado(fila.origen) + _lado(fila.destino)
            for i, (celda, valor) in enumerate(zip(celdas, valores, strict=True)):
                celda.text = ""
                r = _r(celda.paragraphs[0], valor)
                # Código, asignatura o créditos de destino faltantes.
                if i in (5, 6, 8) and not valor:
                    r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    _resaltar_celda_amarillo(celda)
        tabla.autofit = False
        for columna, ancho in zip(tabla.columns, _ANCHOS_CUADRO, strict=True):
            columna.width = Cm(ancho)
        for fila_t in tabla.rows[1:]:
            for celda, ancho in zip(fila_t.cells, _ANCHOS_CUADRO, strict=True):
                celda.width = Cm(ancho)
        for fila_t in tabla.rows:
            for celda in fila_t.cells:
                for parrafo in celda.paragraphs:
                    for run in parrafo.runs:
                        run.font.size = Pt(8)
        _vacio(doc, "FCEN Cuerpo")


def _normas(
    doc: DocumentoDocx,
    normas: list[NormaCitada],
    borrador: bool,
    *,
    solo_cita: bool = False,
) -> None:
    if not normas:
        return
    _p(doc, "FCEN Cuerpo", "Basados en la normativa:")
    for norma in normas:
        _p(doc, "FCEN Norma", f"{norma.norma}, {norma.articulo}", negrita=True)
        if solo_cita:
            # En el oficio la norma se cita, no se transcribe (§5.4).
            continue
        lineas = norma.texto.split("\n")
        for k, linea in enumerate(lineas):
            texto = (
                ("“" if k == 0 else "") + linea + ("”" if k == len(lineas) - 1 else "")
            )
            _p(doc, "FCEN Norma", texto)
        if borrador and not norma.verificada:
            p = _p(doc, "FCEN Norma")
            r = _r(
                p, "(Texto transcrito de un acta anterior; verificar contra la norma.)"
            )
            r.font.size = Pt(8)
            r.italic = False
            r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)


def _conceptos(doc: DocumentoDocx, conceptos: list[Concepto], borrador: bool) -> None:
    """Como en el Acta 35 de CC: «El profesor X expresa: “…”». Texto literal."""
    if not conceptos:
        if borrador:
            _nota_gris(
                _p(doc, "FCEN Cuerpo"), "(Sin conceptos de los miembros del comité.)"
            )
        return
    for concepto in conceptos:
        bloques = _bloques(concepto.texto)
        if not bloques:
            continue
        sujeto = (
            f"{concepto.tratamiento} {concepto.miembro}"
            if concepto.tratamiento
            else concepto.miembro
        )
        p = _p(doc, "FCEN Cuerpo")
        _r(p, sujeto, negrita=True)
        _r(p, " expresa: “")
        _texto_con_saltos(
            p, bloques[0] + ("”." if len(bloques) == 1 else ""), cursiva=False
        )
        for i, bloque in enumerate(bloques[1:], start=2):
            cita = _p(doc, "FCEN Cita")
            _texto_con_saltos(
                cita, bloque + ("”." if i == len(bloques) else ""), cursiva=False
            )
    _vacio(doc, "FCEN Cuerpo")


def _caso(
    doc: DocumentoDocx,
    datos: DatosOficio,
    item: CasoDocumento,
    numero: str,
    periodo: str,
    *,
    acta: bool = False,
) -> None:
    e = item.extraccion
    identificado = (
        "identificado" if e.tratamiento == "El estudiante" else "identificada"
    )
    objeto = e.objeto_solicitud.strip().rstrip(".") or "lo que se transcribe"
    bloques = _bloques(item.cuerpo_correo)

    p = _p(doc, "FCEN Caso")
    _r(p, f"{numero} {(e.nombre_estudiante or '____').upper()}", negrita=True)
    _r(p, f" {identificado} con ")
    _r(p, f"DNI No. {e.documento or '____'}", negrita=True)
    _r(
        p,
        f" mediante {'formato de solicitud' if item.medio == 'formato' else 'correo'}"
        f" del día {fecha_larga(e.fecha_solicitud)}, solicita "
        f"{objeto}, justificando que “",
    )
    # La cita empieza en el mismo párrafo, como en el oficio original.
    if bloques:
        primero = bloques[0] + ("”." if len(bloques) == 1 else "")
        _texto_con_saltos(p, primero, cursiva=True)
        for i, bloque in enumerate(bloques[1:], start=2):
            cita = _p(doc, "FCEN Cita")
            if "\n" in bloque:  # firmas y listas: sin justificar
                cita.alignment = WD_ALIGN_PARAGRAPH.LEFT
            final = "”." if i == len(bloques) else ""
            _texto_con_saltos(cita, bloque + final, cursiva=True)
    else:
        _r(p, "…”.", cursiva=True)

    vinetas = _bullets_analisis(item, periodo)
    if vinetas:
        _p(doc, "FCEN Rotulo", "ANÁLISIS:")
        for vineta in vinetas:
            _p(doc, "FCEN Vineta", f"•\t{vineta}")
    _vacio(doc, "FCEN Cuerpo")

    if acta:
        _vacio(doc, "FCEN Cuerpo")
        _conceptos(doc, item.caso.conceptos, datos.borrador)

    antes, sentido, despues = texto_recomendacion(datos.programa, item)
    p = _p(doc, "FCEN Cuerpo")
    _r(p, "RECOMENDACIÓN", negrita=True)
    _r(p, f": {antes}")
    r = _r(p, sentido, negrita=item.caso.sentido == Sentido.PENDIENTE)
    if item.caso.sentido == Sentido.PENDIENTE:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    _r(p, despues)
    _vacio(doc, "FCEN Cuerpo")
    _asignaturas(doc, item)
    _equivalencias(doc, item)

    _normas(
        doc,
        normas_para(e.tipos, item.caso.programa),
        datos.borrador,
        solo_cita=not acta,
    )
    _vacio(doc, "FCEN Cuerpo")


def _cierre(doc: DocumentoDocx, datos: DatosOficio) -> None:
    _p(doc, "FCEN Firma", "Agradecemos su valiosa gestión")
    _vacio(doc, "FCEN Firma")
    _p(doc, "FCEN Firma", "Atentamente,")
    _vacio(doc, "FCEN Firma")
    if datos.borrador:
        p = _p(doc, "FCEN Firma")
        r = _r(p, "[la firma se agrega al firmar]")
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
    else:
        _vacio(doc, "FCEN Firma")
    _p(doc, "FCEN Firma", "NOMBRE DEL FIRMANTE", negrita=True)
    _p(doc, "FCEN Firma", "Cargo")
    if datos.programa.correo:
        p = _p(doc, "FCEN Firma")
        r = _r(p, datos.programa.correo, subrayado=True)
        r.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)


def _documento(datos: DatosOficio) -> DocumentoDocx:
    doc = Document(str(PLANTILLA))
    _reemplazar_marcadores(
        doc,
        {
            "ENCABEZADO": datos.programa.encabezado,
            "CORREO": datos.programa.correo or "",
            "EXTENSION": datos.programa.extension or "",
            "ELABORO": datos.elaboro,
        },
    )
    # El párrafo vacío que trae la plantilla se elimina.
    if doc.paragraphs and not doc.paragraphs[0].text:
        primero = doc.paragraphs[0]._element
        primero.getparent().remove(primero)
    return doc


def _secciones(datos: DatosOficio) -> OrderedDict[str, list[CasoDocumento]]:
    secciones: OrderedDict[str, list[CasoDocumento]] = OrderedDict()
    for item in datos.casos:
        secciones.setdefault(titulo_seccion(item.extraccion), []).append(item)
    return secciones


def _casos(doc: DocumentoDocx, datos: DatosOficio, periodo: str, *, acta: bool) -> None:
    # Como en las actas: los casos son el punto 2 del orden del día (2.1, 2.2…).
    numero = 0
    for titulo, items in _secciones(datos).items():
        _p(doc, "FCEN Titulo", titulo)
        for item in items:
            numero += 1
            _caso(doc, datos, item, f"2.{numero}", periodo, acta=acta)


def _guardar(doc: DocumentoDocx) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def construir_oficio(datos: DatosOficio, periodo: str) -> bytes:
    """Oficio al Consejo de Facultad: solo la recomendación, sin conceptos."""
    doc = _documento(datos)
    _encabezado_oficio(doc, datos)
    _casos(doc, datos, periodo, acta=False)
    _cierre(doc, datos)
    return _guardar(doc)


# ---------------------------------------------------------------------------
# Acta del comité
# ---------------------------------------------------------------------------


def _linea_dato(doc: DocumentoDocx, rotulo: str, valor: str | None, falta: str) -> None:
    p = _p(doc, "FCEN Destinatario")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(3.5))
    _r(p, f"{rotulo}:\t", negrita=True)
    if valor:
        _r(p, valor)
    else:
        _marcador(p, falta)


def _encabezado_acta(doc: DocumentoDocx, datos: DatosOficio) -> None:
    if datos.borrador:
        p = _p(doc, "FCEN Destinatario")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = _r(p, "BORRADOR · pendiente de revisión del Comité Asesor", negrita=True)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
        _vacio(doc)

    for linea in (
        "UNIVERSIDAD NACIONAL DE COLOMBIA",
        "SEDE MANIZALES",
        "FACULTAD DE CIENCIAS EXACTAS Y NATURALES",
        f"COMITÉ ASESOR PROGRAMA CURRICULAR {datos.programa.nombre.upper()}",
        "PREGRADO",
    ):
        _p(
            doc, "FCEN Destinatario", linea, negrita=True
        ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    _vacio(doc)
    p = _p(doc, "FCEN Destinatario")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _r(p, "ACTA ", negrita=True)
    if datos.numero_acta:
        _r(p, datos.numero_acta, negrita=True)
    else:
        _marcador(p, "[No.]")
    _r(p, f" DE {datos.fecha.year}", negrita=True)
    _vacio(doc, "FCEN Texto")

    _linea_dato(
        doc,
        "FECHA",
        datos.fecha_sesion and f"Manizales, {datos.fecha_sesion}",
        f"Manizales, {fecha_larga(datos.fecha)}",
    )
    _linea_dato(doc, "HORA", datos.hora, "[hora]")
    _linea_dato(doc, "LUGAR", datos.lugar, "[lugar o «sesión ad referéndum»]")
    _vacio(doc, "FCEN Texto")

    _p(doc, "FCEN Texto", "ASISTENTES:", negrita=True)
    if datos.miembros:
        for m in datos.miembros:
            _p(doc, "FCEN Vineta", f"•\t{m.nombre}, {m.cargo}.")
    else:
        _marcador(_p(doc, "FCEN Vineta"), "[Integrantes del comité, con su cargo]")
    p = _p(doc, "FCEN Destinatario")
    p.paragraph_format.tab_stops.add_tab_stop(Cm(3.5))
    _r(p, "INVITADOS:\t", negrita=True)
    _marcador(p, "[Ninguno]")
    _vacio(doc, "FCEN Texto")

    _p(doc, "FCEN Texto", "ORDEN DEL DÍA", negrita=True)
    _p(doc, "FCEN Vineta", "1.\tVerificación del quórum")
    _p(doc, "FCEN Vineta", "2.\tCasos estudiantiles")
    numero = 0
    for titulo, items in _secciones(datos).items():
        _p(doc, "FCEN Caso", titulo.capitalize(), negrita=True)
        for item in items:
            numero += 1
            nombre = (item.extraccion.nombre_estudiante or "____").title()
            _p(doc, "FCEN Caso", f"2.{numero} {nombre}")
    _p(doc, "FCEN Vineta", "3.\tVarios")
    _vacio(doc, "FCEN Texto")

    _p(doc, "FCEN Titulo", "DESARROLLO DE LA SESIÓN")
    _p(doc, "FCEN Texto", "1. VERIFICACIÓN DEL QUÓRUM", negrita=True)
    _marcador(_p(doc, "FCEN Texto"), "[Se verificó el quórum con la asistencia de …]")
    _vacio(doc, "FCEN Texto")
    _p(doc, "FCEN Texto", "2. CASOS ESTUDIANTILES", negrita=True)
    _vacio(doc, "FCEN Cuerpo")


def _cierre_acta(doc: DocumentoDocx, datos: DatosOficio) -> None:
    _p(doc, "FCEN Texto", "3. VARIOS", negrita=True)
    _marcador(_p(doc, "FCEN Texto"), "[Sin asuntos varios]")
    _vacio(doc, "FCEN Firma", n=2)
    if datos.borrador:
        _nota_gris(_p(doc, "FCEN Firma"), "[la firma se agrega al firmar]")
    else:
        _vacio(doc, "FCEN Firma")
    presidente = datos.miembros[0] if datos.miembros else None
    _p(
        doc,
        "FCEN Firma",
        presidente.nombre.upper() if presidente else "NOMBRE DEL FIRMANTE",
        negrita=True,
    )
    _p(doc, "FCEN Firma", presidente.cargo.split(",")[0] if presidente else "Cargo")
    if datos.programa.correo:
        p = _p(doc, "FCEN Firma")
        r = _r(p, datos.programa.correo, subrayado=True)
        r.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)


def construir_acta(datos: DatosOficio, periodo: str) -> bytes:
    """Acta interna del comité: la sesión, los casos y los conceptos."""
    doc = _documento(datos)
    _encabezado_acta(doc, datos)
    _casos(doc, datos, periodo, acta=True)
    _cierre_acta(doc, datos)
    return _guardar(doc)
