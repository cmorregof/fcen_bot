"""Indicadores de tiempos de respuesta del servicio de actas.

Solo lectura y solo cifras: nunca nombres, documentos ni textos de los casos.
Las definiciones están en docs/indicadores.md y cada cifra lleva su fuente
(`FUENTES`).

El comité asesor solo recomienda; lo decidido sale del acta del Consejo
confirmada por una persona en MemorIA (`decisiones_consejo`). Un caso
recomendado sin esa confirmación es «pendiente de acta del Consejo», nunca
aprobado. Un aplazamiento no es una decisión: el caso sigue pendiente.

Cada indicador se calcula sobre los casos *recibidos* en el periodo (fecha de
registro en MemorIA, hora de Bogotá), con su estado a la fecha del reporte.
"""

from __future__ import annotations

import io
import json
import math
import re
import statistics
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime, timedelta
from functools import lru_cache
from html import escape
from pathlib import Path
from zoneinfo import ZoneInfo

import aiosqlite
import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from pydantic import BaseModel, ConfigDict, Field, ValidationError, field_validator

from app.core import solicitudes
from app.core.config import settings
from app.core.programas import DIRECCIONES, PROGRAMAS, normalizar
from app.schemas.caso import TITULOS_TIPO, EstadoCaso, Sentido, TipoSolicitud
from app.schemas.consejo import NOMBRE_INSTANCIA, Instancia, SentidoConsejo

MESES = (
    "enero",
    "febrero",
    "marzo",
    "abril",
    "mayo",
    "junio",
    "julio",
    "agosto",
    "septiembre",
    "octubre",
    "noviembre",
    "diciembre",
)
SIN_PROGRAMA = "Sin programa asignado"
SIN_FICHA = "Sin ficha (caso sin leer)"
FACULTAD = "Facultad de Ciencias Exactas y Naturales"
_POR_CONFIRMAR = "Instancia por confirmar"
_DE_CONFIG = {
    "consejo_facultad": NOMBRE_INSTANCIA[Instancia.FACULTAD],
    "consejo_sede": NOMBRE_INSTANCIA[Instancia.SEDE],
    "consejo_superior": NOMBRE_INSTANCIA[Instancia.SUPERIOR],
    "por_confirmar": _POR_CONFIRMAR,
}
_DECISION = (SentidoConsejo.APROBO, SentidoConsejo.NO_APROBO, SentidoConsejo.OTRO)


# ---------------------------------------------------------------------------
# Definiciones y fuentes
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Fuente:
    nombre: str
    definicion: str
    fuente: str


_MEMORIA = "Casos de MemorIA (servicio de actas)"
_BOTON = "Botón «Recomendar aprobar / no aprobar» de la administrativa en MemorIA"
_ACTA = (
    "Acta del Consejo (Facultad, Sede o Superior) leída en MemorIA y confirmada "
    "por una persona, caso por caso"
)
_PLAZOS = "config/solicitudes.yaml, sección «plazos»"

FUENTES: dict[str, Fuente] = {
    "recibidos": Fuente(
        "Casos recibidos",
        "Casos registrados en MemorIA en el periodo, sin los descartados.",
        f"{_MEMORIA}: fecha de registro del caso (hora de Bogotá).",
    ),
    "recomendados": Fuente(
        "Recomendados por el comité",
        "Casos con recomendación del comité asesor registrada: recomendar "
        "aprobar más recomendar no aprobar.",
        _BOTON + ".",
    ),
    "recomendar_aprobar": Fuente(
        "Recomendar aprobar",
        "Casos en que el comité recomendó aprobar.",
        _BOTON + ".",
    ),
    "recomendar_no_aprobar": Fuente(
        "Recomendar no aprobar",
        "Casos en que el comité recomendó no aprobar.",
        _BOTON + ".",
    ),
    "decididos_consejo": Fuente(
        "Decididos por el Consejo",
        "Casos cuya decisión consta en un acta del Consejo confirmada en "
        "MemorIA: aprobó, no aprobó u otra decisión. Un aplazamiento no cuenta.",
        _ACTA + ".",
    ),
    "consejo_aprobo": Fuente(
        "El Consejo aprobó",
        "Casos que el acta del Consejo confirmada registra como aprobados.",
        _ACTA + ".",
    ),
    "consejo_no_aprobo": Fuente(
        "El Consejo no aprobó",
        "Casos que el acta del Consejo confirmada registra como no aprobados.",
        _ACTA + ".",
    ),
    "consejo_aplazo": Fuente(
        "El Consejo aplazó",
        "Casos que el acta del Consejo confirmada registra como aplazados. "
        "Siguen pendientes de acta del Consejo.",
        _ACTA + ".",
    ),
    "consejo_otro": Fuente(
        "Otra decisión del Consejo",
        "Decisiones del acta que no son aprobó, no aprobó ni aplazó.",
        _ACTA + ".",
    ),
    "pendientes_acta": Fuente(
        "Pendientes de acta del Consejo",
        "Casos recomendados por el comité sin decisión confirmada desde el acta "
        "del Consejo, más los aplazados. Nunca se cuentan como aprobados.",
        f"{_BOTON}, cruzado con {_ACTA[0].lower() + _ACTA[1:]}.",
    ),
    "dias_comite_mediana": Fuente(
        "Días de la solicitud a la recomendación (mediana)",
        "Días calendario entre la fecha de la solicitud y la fecha en que se "
        "registró la recomendación del comité. Mediana de los casos con ambas "
        "fechas.",
        "Fecha del correo o del formato de solicitud (si no se leyó o es "
        "posterior al registro, la fecha de registro del caso); fecha del clic "
        "en el botón de recomendación en MemorIA.",
    ),
    "dias_comite_p90": Fuente(
        "Días de la solicitud a la recomendación (percentil 90)",
        "Percentil 90 por rango más cercano: el valor en la posición "
        "⌈0,9·n⌉ de los días ordenados.",
        "Las mismas fechas de la mediana.",
    ),
    "dias_comite_casos": Fuente(
        "Casos con días de la solicitud a la recomendación",
        "Cuántos casos entran en la mediana y el percentil 90. Los casos "
        "recomendados antes de que MemorIA guardara la fecha de la "
        "recomendación no entran.",
        "Las mismas fechas de la mediana.",
    ),
    "dias_consejo_mediana": Fuente(
        "Días de la recomendación a la decisión del Consejo (mediana)",
        "Días calendario entre la fecha en que se registró la recomendación y "
        "la fecha de la sesión que consta en el acta del Consejo. Solo casos "
        "decididos (no aplazados) con ambas fechas.",
        "Fecha del clic en el botón de recomendación; fecha de la sesión "
        "leída del acta del Consejo confirmada.",
    ),
    "dias_consejo_p90": Fuente(
        "Días de la recomendación a la decisión del Consejo (percentil 90)",
        "Percentil 90 por rango más cercano.",
        "Las mismas fechas de la mediana.",
    ),
    "dias_consejo_casos": Fuente(
        "Casos con días de la recomendación a la decisión",
        "Cuántos casos entran en la mediana y el percentil 90. No entran los "
        "casos sin fecha de recomendación, sin fecha en el acta o con la fecha "
        "del acta anterior al registro de la recomendación.",
        "Las mismas fechas de la mediana.",
    ),
    "fuera_plazo_comite": Fuente(
        "Fuera de plazo en el comité",
        "Casos cuya recomendación llegó después del plazo, o que siguen sin "
        "recomendación y ya lo superaron a la fecha del reporte. Vacío si el "
        "plazo no está configurado.",
        f"{_PLAZOS} («comite»), con las fechas de los días de la solicitud a la "
        "recomendación.",
    ),
    "fuera_plazo_consejo": Fuente(
        "Fuera de plazo en el Consejo",
        "Casos decididos después del plazo desde la recomendación, o pendientes "
        "de acta que ya lo superaron a la fecha del reporte. Vacío si el plazo "
        "no está configurado.",
        f"{_PLAZOS} («consejo»), con las fechas de los días de la recomendación "
        "a la decisión.",
    ),
}

DIMENSIONES: dict[str, Fuente] = {
    "Mes": Fuente("Mes", "Mes de registro del caso en MemorIA.", _MEMORIA + "."),
    "Área curricular": Fuente(
        "Área curricular",
        "Área a la que pertenece la administrativa que registró el caso.",
        _MEMORIA + ".",
    ),
    "Programa": Fuente(
        "Programa",
        "Programa curricular reconocido en la ficha (código del plan en el "
        "pantallazo del SIA) o asignado a mano.",
        _MEMORIA + ".",
    ),
    "Tipo de solicitud": Fuente(
        "Tipo de solicitud",
        "Trámite principal de la ficha (el primero). Un caso cuenta una sola vez.",
        _MEMORIA + ": ficha del caso.",
    ),
    "Instancia": Fuente(
        "Instancia",
        "Si el caso ya tiene acta confirmada, el Consejo que la expidió; si no, "
        "el Consejo que decide el trámite principal según la ruta configurada.",
        f"{_ACTA}; si no hay acta, config/solicitudes.yaml (campo «decide»).",
    ),
}

CLAVES = tuple(FUENTES)


# ---------------------------------------------------------------------------
# Columnas del Excel (config/indicadores.yaml)
# ---------------------------------------------------------------------------


class ConfigIndicadoresError(RuntimeError):
    """El archivo de columnas falta o no es válido."""


class Columna(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    clave: str
    encabezado: str = Field(min_length=1)

    @field_validator("clave")
    @classmethod
    def _conocida(cls, valor: str) -> str:
        if valor not in FUENTES:
            raise ValueError(
                f"«{valor}» no es un indicador (válidos: {', '.join(CLAVES)})"
            )
        return valor


class ConfigIndicadores(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)

    columnas: list[Columna] = Field(min_length=1)


def cargar_columnas(ruta: Path) -> ConfigIndicadores:
    try:
        datos = yaml.safe_load(ruta.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ConfigIndicadoresError(
            f"No encuentro {ruta}: dice qué columnas lleva el Excel de "
            "indicadores (docs/indicadores.md)."
        ) from exc
    except yaml.YAMLError as exc:
        raise ConfigIndicadoresError(f"{ruta} no es un YAML válido: {exc}") from exc
    try:
        return ConfigIndicadores.model_validate(datos)
    except ValidationError as exc:
        problemas = "\n".join(
            f"  - {'.'.join(str(p) for p in e['loc']) or '(raíz)'}: {e['msg']}"
            for e in exc.errors()
        )
        raise ConfigIndicadoresError(f"{ruta} tiene errores:\n{problemas}") from exc


@lru_cache(maxsize=4)
def _columnas_cacheadas(ruta: Path) -> ConfigIndicadores:
    return cargar_columnas(ruta)


def columnas() -> list[Columna]:
    return list(_columnas_cacheadas(settings.indicadores_path).columnas)


# ---------------------------------------------------------------------------
# Periodo
# ---------------------------------------------------------------------------


class PeriodoInvalido(ValueError):
    """El texto no es un periodo reconocible; el mensaje se muestra tal cual."""


@dataclass(frozen=True)
class Periodo:
    etiqueta: str
    desde: date
    hasta: date  # exclusivo

    def contiene(self, dia: date) -> bool:
        return self.desde <= dia < self.hasta

    @property
    def texto(self) -> str:
        ultimo = self.hasta - timedelta(days=1)
        return f"{self.etiqueta} ({fecha_larga(self.desde)} a {fecha_larga(ultimo)})"


def ayuda_periodo() -> str:
    return (
        "Escriba el periodo así: «2026-10» (un mes), «2026» (un año) o "
        f"«{settings.periodo_actual}» o nada (el periodo académico en curso)."
    )


_RE_MES = re.compile(r"^(\d{4})-(\d{1,2})$")
_RE_ANO = re.compile(r"^(\d{4})$")


def fecha_larga(d: date) -> str:
    return f"{d.day} de {MESES[d.month - 1]} de {d.year}"


def mes(ano: int, numero: int) -> Periodo:
    desde = date(ano, numero, 1)
    hasta = date(ano + 1, 1, 1) if numero == 12 else date(ano, numero + 1, 1)
    return Periodo(f"{MESES[numero - 1]} de {ano}", desde, hasta)


def mes_anterior(hoy: date) -> Periodo:
    ultimo = hoy.replace(day=1) - timedelta(days=1)
    return mes(ultimo.year, ultimo.month)


def periodo_academico(hoy: date) -> Periodo | None:
    """Del primer día de clases (PERIODO_INICIO) a hoy, inclusive."""
    inicio = settings.periodo_inicio
    if inicio is None or inicio > hoy:
        return None
    return Periodo(
        f"periodo {settings.periodo_actual}", inicio, hoy + timedelta(days=1)
    )


def periodo_de(texto: str, hoy: date) -> Periodo:
    t = texto.strip()
    if t.lower() in ("", "actual", settings.periodo_actual.lower()):
        periodo = periodo_academico(hoy)
        if periodo is None:
            raise PeriodoInvalido(
                "No está configurada la fecha de inicio del periodo académico "
                "(PERIODO_INICIO). Pida un mes, por ejemplo «2026-10»."
            )
        return periodo
    if m := _RE_MES.match(t):
        ano, numero = int(m.group(1)), int(m.group(2))
        if 1 <= numero <= 12 and ano >= 2000:
            return mes(ano, numero)
    elif m := _RE_ANO.match(t):
        ano = int(m.group(1))
        if ano >= 2000:
            return Periodo(f"año {ano}", date(ano, 1, 1), date(ano + 1, 1, 1))
    raise PeriodoInvalido(f"No reconozco el periodo «{t}». {ayuda_periodo()}")


# ---------------------------------------------------------------------------
# Lectura: un caso, solo con lo que necesitan las cifras
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class FilaCaso:
    """Lo único que se lee de cada caso. Sin nombres ni documentos."""

    direccion: str
    programa: str
    tipo: str
    instancia: str
    recibido: date
    solicitud: date
    recomendado: date | None
    sentido: Sentido
    consejo: SentidoConsejo | None
    acta_fecha: date | None


def _zona() -> ZoneInfo:
    return ZoneInfo(settings.timezone)


def _dia_local(valor: object) -> date | None:
    """SQLite guarda CURRENT_TIMESTAMP en UTC («AAAA-MM-DD HH:MM:SS»)."""
    if not valor:
        return None
    try:
        momento = datetime.fromisoformat(str(valor))
    except ValueError:
        return None
    if momento.tzinfo is None:
        momento = momento.replace(tzinfo=UTC)
    return momento.astimezone(_zona()).date()


def _fecha(valor: object) -> date | None:
    try:
        return date.fromisoformat(str(valor)[:10]) if valor else None
    except ValueError:
        return None


def _ficha(extraccion_json: str | None) -> tuple[TipoSolicitud | None, date | None]:
    """Del JSON de la ficha solo se leen el trámite principal y la fecha."""
    if not extraccion_json:
        return None, None
    try:
        datos = json.loads(extraccion_json)
    except ValueError:
        return None, None
    tipos = datos.get("tipos") if isinstance(datos, dict) else None
    tipo: TipoSolicitud | None = None
    if isinstance(tipos, list) and tipos:
        try:
            tipo = TipoSolicitud(tipos[0])
        except ValueError:
            tipo = TipoSolicitud.OTRO
    elif isinstance(datos, dict):
        tipo = TipoSolicitud.OTRO
    fecha = _fecha(datos.get("fecha_solicitud")) if isinstance(datos, dict) else None
    return tipo, fecha


def _instancia(
    tipo: TipoSolicitud | None, acta: str | None, config: solicitudes.ConfigSolicitudes
) -> str:
    if acta:
        try:
            return NOMBRE_INSTANCIA[Instancia(acta)]
        except ValueError:
            pass
    if tipo is None:
        return SIN_FICHA
    return _DE_CONFIG[config.ruta(tipo).decide]


async def leer_casos(
    db: aiosqlite.Connection, periodo: Periodo, direccion: str | None = None
) -> list[FilaCaso]:
    """Casos recibidos en el periodo (sin descartados), de un área o de todas."""
    config = solicitudes.actual()
    # Margen de un día por la diferencia entre UTC y la hora de Bogotá.
    desde = (periodo.desde - timedelta(days=1)).isoformat()
    hasta = (periodo.hasta + timedelta(days=1)).isoformat()
    sql = (
        "SELECT c.direccion, c.programa, c.extraccion_json, c.sentido, c.creado,"
        " c.recomendado, d.sentido, d.instancia, d.acta_fecha"
        " FROM casos c LEFT JOIN decisiones_consejo d ON d.caso_id = c.id"
        " WHERE c.estado != ? AND c.creado >= ? AND c.creado < ?"
    )
    parametros: list[object] = [EstadoCaso.DESCARTADO.value, desde, hasta]
    if direccion is not None:
        sql += " AND c.direccion = ?"
        parametros.append(direccion)
    cur = await db.execute(sql + " ORDER BY c.id", parametros)
    filas: list[FilaCaso] = []
    for r in await cur.fetchall():
        recibido = _dia_local(r[4])
        if recibido is None or not periodo.contiene(recibido):
            continue
        tipo, fecha_solicitud = _ficha(r[2])
        # Una fecha leída posterior al registro es una mala lectura.
        solicitud = (
            fecha_solicitud
            if fecha_solicitud is not None and fecha_solicitud <= recibido
            else recibido
        )
        filas.append(
            FilaCaso(
                direccion=str(r[0]),
                programa=PROGRAMAS[r[1]].nombre if r[1] in PROGRAMAS else SIN_PROGRAMA,
                tipo=TITULOS_TIPO[tipo] if tipo is not None else SIN_FICHA,
                instancia=_instancia(tipo, r[7], config),
                recibido=recibido,
                solicitud=solicitud,
                recomendado=_dia_local(r[5]),
                sentido=Sentido(r[3]),
                consejo=SentidoConsejo(r[6]) if r[6] else None,
                acta_fecha=_fecha(r[8]),
            )
        )
    return filas


# ---------------------------------------------------------------------------
# Cálculo
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class Estadistica:
    casos: int = 0
    mediana: float | None = None
    p90: int | None = None


def percentil(valores: list[int], p: float) -> int | None:
    """Percentil por rango más cercano: el valor en la posición ⌈p·n⌉."""
    if not valores:
        return None
    orden = sorted(valores)
    return orden[max(1, math.ceil(p * len(orden))) - 1]


def estadistica(valores: list[int]) -> Estadistica:
    if not valores:
        return Estadistica()
    return Estadistica(
        len(valores), float(statistics.median(valores)), percentil(valores, 0.9)
    )


@dataclass(frozen=True)
class Indicadores:
    recibidos: int = 0
    recomendar_aprobar: int = 0
    recomendar_no_aprobar: int = 0
    consejo_aprobo: int = 0
    consejo_no_aprobo: int = 0
    consejo_aplazo: int = 0
    consejo_otro: int = 0
    pendientes_acta: int = 0
    dias_comite: Estadistica = field(default_factory=Estadistica)
    dias_consejo: Estadistica = field(default_factory=Estadistica)
    # None: el plazo no está configurado en config/solicitudes.yaml.
    fuera_plazo_comite: int | None = None
    fuera_plazo_consejo: int | None = None

    @property
    def recomendados(self) -> int:
        return self.recomendar_aprobar + self.recomendar_no_aprobar

    @property
    def decididos_consejo(self) -> int:
        return self.consejo_aprobo + self.consejo_no_aprobo + self.consejo_otro

    def valor(self, clave: str) -> int | float | None:
        valores: dict[str, int | float | None] = {
            "recibidos": self.recibidos,
            "recomendados": self.recomendados,
            "recomendar_aprobar": self.recomendar_aprobar,
            "recomendar_no_aprobar": self.recomendar_no_aprobar,
            "decididos_consejo": self.decididos_consejo,
            "consejo_aprobo": self.consejo_aprobo,
            "consejo_no_aprobo": self.consejo_no_aprobo,
            "consejo_aplazo": self.consejo_aplazo,
            "consejo_otro": self.consejo_otro,
            "pendientes_acta": self.pendientes_acta,
            "dias_comite_mediana": self.dias_comite.mediana,
            "dias_comite_p90": self.dias_comite.p90,
            "dias_comite_casos": self.dias_comite.casos,
            "dias_consejo_mediana": self.dias_consejo.mediana,
            "dias_consejo_p90": self.dias_consejo.p90,
            "dias_consejo_casos": self.dias_consejo.casos,
            "fuera_plazo_comite": self.fuera_plazo_comite,
            "fuera_plazo_consejo": self.fuera_plazo_consejo,
        }
        return valores[clave]


def _recomendado(f: FilaCaso) -> bool:
    return f.sentido in (Sentido.APROBAR, Sentido.NO_APROBAR)


def _decidido(f: FilaCaso) -> bool:
    return f.consejo in _DECISION


def calcular(
    filas: Iterable[FilaCaso], hoy: date, plazos: solicitudes.Plazos
) -> Indicadores:
    filas = list(filas)
    dias_comite: list[int] = []
    dias_consejo: list[int] = []
    fuera_comite = 0
    fuera_consejo = 0
    for f in filas:
        if _recomendado(f) and f.recomendado is not None:
            dias_comite.append((f.recomendado - f.solicitud).days)
        if (
            _decidido(f)
            and f.recomendado is not None
            and f.acta_fecha is not None
            and f.acta_fecha >= f.recomendado
        ):
            dias_consejo.append((f.acta_fecha - f.recomendado).days)
        if plazos.comite:
            # Sin recomendación, el reloj del comité sigue corriendo hasta hoy.
            fin: date | None
            if _recomendado(f):
                fin = f.recomendado  # None: recomendado sin fecha, no se mide
            elif _decidido(f):
                fin = None  # decidido sin recomendación registrada
            else:
                fin = hoy
            if fin is not None and (fin - f.solicitud).days > plazos.comite.dias:
                fuera_comite += 1
        if plazos.consejo and _recomendado(f) and f.recomendado is not None:
            # Pendiente de acta (o aplazado): el reloj del Consejo sigue corriendo.
            fin = f.acta_fecha if _decidido(f) else hoy
            if fin is not None and (fin - f.recomendado).days > plazos.consejo.dias:
                fuera_consejo += 1
    return Indicadores(
        recibidos=len(filas),
        recomendar_aprobar=sum(f.sentido == Sentido.APROBAR for f in filas),
        recomendar_no_aprobar=sum(f.sentido == Sentido.NO_APROBAR for f in filas),
        consejo_aprobo=sum(f.consejo == SentidoConsejo.APROBO for f in filas),
        consejo_no_aprobo=sum(f.consejo == SentidoConsejo.NO_APROBO for f in filas),
        consejo_aplazo=sum(f.consejo == SentidoConsejo.APLAZO for f in filas),
        consejo_otro=sum(f.consejo == SentidoConsejo.OTRO for f in filas),
        # Sin decisión confirmada desde el acta, la recomendación sigue pendiente.
        pendientes_acta=sum(_recomendado(f) and not _decidido(f) for f in filas),
        dias_comite=estadistica(dias_comite),
        dias_consejo=estadistica(dias_consejo),
        fuera_plazo_comite=fuera_comite if plazos.comite else None,
        fuera_plazo_consejo=fuera_consejo if plazos.consejo else None,
    )


@dataclass(frozen=True)
class Grupo:
    dimension: str
    nombre: str
    indicadores: Indicadores


@dataclass(frozen=True)
class Reporte:
    periodo: Periodo
    alcance: str
    hoy: date
    plazos: solicitudes.Plazos
    total: Indicadores
    grupos: list[Grupo]

    def de(self, dimension: str) -> list[Grupo]:
        return [g for g in self.grupos if g.dimension == dimension]


def _agrupar(
    filas: list[FilaCaso],
    dimension: str,
    clave: Callable[[FilaCaso], str],
    hoy: date,
    plazos: solicitudes.Plazos,
) -> list[Grupo]:
    grupos: dict[str, list[FilaCaso]] = {}
    for f in filas:
        grupos.setdefault(clave(f), []).append(f)
    # Alfabético, con «Sin programa» o «Sin ficha» al final.
    nombres = sorted(grupos, key=lambda n: (n.startswith("Sin "), n))
    return [Grupo(dimension, n, calcular(grupos[n], hoy, plazos)) for n in nombres]


def _nombre_mes(f: FilaCaso) -> str:
    return f"{f.recibido.year}-{f.recibido.month:02d}"


def nombre_area(direccion: str | None) -> str:
    if direccion is None:
        return FACULTAD
    return DIRECCIONES[direccion].nombre if direccion in DIRECCIONES else direccion


async def reporte(
    db: aiosqlite.Connection,
    periodo: Periodo,
    hoy: date,
    direccion: str | None = None,
) -> Reporte:
    """Indicadores del periodo para un área (o toda la Facultad si es None)."""
    plazos = solicitudes.actual().plazos
    filas = await leer_casos(db, periodo, direccion)
    grupos: list[Grupo] = []
    if direccion is None:
        grupos += _agrupar(
            filas, "Área curricular", lambda f: nombre_area(f.direccion), hoy, plazos
        )
    grupos += _agrupar(filas, "Programa", lambda f: f.programa, hoy, plazos)
    grupos += _agrupar(filas, "Tipo de solicitud", lambda f: f.tipo, hoy, plazos)
    grupos += _agrupar(filas, "Instancia", lambda f: f.instancia, hoy, plazos)
    meses = {_nombre_mes(f) for f in filas}
    if len(meses) > 1:
        grupos += _agrupar(filas, "Mes", _nombre_mes, hoy, plazos)
    return Reporte(
        periodo=periodo,
        alcance=nombre_area(direccion),
        hoy=hoy,
        plazos=plazos,
        total=calcular(filas, hoy, plazos),
        grupos=grupos,
    )


def hoy_local() -> date:
    return datetime.now(_zona()).date()


# ---------------------------------------------------------------------------
# Texto para Telegram
# ---------------------------------------------------------------------------

# Notas de fuente del mensaje: número → claves de FUENTES que explica.
_NOTAS: tuple[tuple[str, ...], ...] = (
    ("recibidos",),
    ("recomendados",),
    ("decididos_consejo",),
    ("pendientes_acta",),
    ("dias_comite_mediana", "dias_comite_p90"),
    ("dias_consejo_mediana", "dias_consejo_p90"),
    ("fuera_plazo_comite", "fuera_plazo_consejo"),
)


def _dias(valor: float | None) -> str:
    if valor is None:
        return "—"
    return (
        str(int(valor))
        if float(valor).is_integer()
        else f"{valor:.1f}".replace(".", ",")
    )


def _est(e: Estadistica) -> str:
    if not e.casos:
        return "sin casos con ambas fechas"
    return (
        f"mediana {_dias(e.mediana)} · percentil 90 {_dias(e.p90)} "
        f"({e.casos} {'caso' if e.casos == 1 else 'casos'})"
    )


def _plazo(valor: int | None, plazo: solicitudes.Plazo | None) -> str:
    if plazo is None or valor is None:
        return "sin plazo configurado"
    return f"{valor} (plazo {plazo.dias} días)"


def lineas_total(r: Reporte) -> list[str]:
    """Las cifras del total, cada una con el número de su fuente."""
    i, plazos = r.total, r.plazos
    return [
        f"• Casos recibidos: {i.recibidos} [1]",
        f"• Recomendados por el comité: {i.recomendados} "
        f"({i.recomendar_aprobar} recomendar aprobar, "
        f"{i.recomendar_no_aprobar} recomendar no aprobar) [2]",
        f"• Decididos por el Consejo: {i.decididos_consejo} (aprobó "
        f"{i.consejo_aprobo}, no aprobó {i.consejo_no_aprobo}"
        + (f", otra decisión {i.consejo_otro}" if i.consejo_otro else "")
        + f"); aplazó {i.consejo_aplazo} [3]",
        f"• Pendientes de acta del Consejo: {i.pendientes_acta} [4]",
        f"• Días de la solicitud a la recomendación: {_est(i.dias_comite)} [5]",
        f"• Días de la recomendación a la decisión del Consejo: "
        f"{_est(i.dias_consejo)} [6]",
        f"• Fuera de plazo: comité {_plazo(i.fuera_plazo_comite, plazos.comite)}; "
        f"Consejo {_plazo(i.fuera_plazo_consejo, plazos.consejo)} [7]",
    ]


def _linea_grupo(g: Grupo) -> str:
    i = g.indicadores
    return (
        f"• {escape(g.nombre)}: {i.recibidos} recibidos · {i.recomendados} "
        f"recomendados · {i.decididos_consejo} decididos · {i.pendientes_acta} "
        f"pendientes de acta · mediana {_dias(i.dias_comite.mediana)} / "
        f"{_dias(i.dias_consejo.mediana)} días"
    )


def _fuentes() -> list[str]:
    lineas = ["<b>Fuentes</b>"]
    for n, claves in enumerate(_NOTAS, start=1):
        texto = FUENTES[claves[0]].fuente
        if claves[0] == "fuera_plazo_comite":
            texto = f"{_PLAZOS} (días calendario), con las fechas de [5] y [6]."
        lineas.append(f"[{n}] {escape(texto)}")
    return lineas


def formatear(
    r: Reporte,
    *,
    titulo: str = "Indicadores del servicio de actas",
    extra: list[str] | None = None,
) -> str:
    """Mensaje HTML de Telegram. `extra` va antes de las fuentes."""
    lineas = [
        f"<b>{escape(titulo)}</b>",
        f"{escape(r.alcance)} · {escape(r.periodo.texto)}",
        f"Corte: {fecha_larga(r.hoy)}",
        "",
    ]
    if not r.total.recibidos:
        lineas.append("No hay casos recibidos en el periodo.")
    else:
        lineas += ["<b>Total</b>", *lineas_total(r)]
        for dimension in DIMENSIONES:
            grupos = r.de(dimension)
            if not grupos:
                continue
            lineas += ["", f"<b>Por {dimension.lower()}</b>"]
            lineas += [_linea_grupo(g) for g in grupos]
        lineas += [
            "",
            "<i>Por grupo: recibidos · recomendados · decididos · pendientes "
            "de acta · mediana de días solicitud→recomendación / "
            "recomendación→Consejo.</i>",
        ]
    if extra:
        lineas += ["", *extra]
    lineas += ["", *_fuentes(), "", _NOTA_FINAL]
    return "\n".join(lineas)


_NOTA_FINAL = (
    "<i>El comité asesor solo recomienda; decide el Consejo. Mientras su "
    "decisión no conste en el acta del Consejo confirmada en MemorIA, el caso "
    "figura como pendiente de acta del Consejo y no como aprobado. Solo cifras: "
    "sin nombres ni documentos.</i>"
)


# ---------------------------------------------------------------------------
# JSON (web) y Excel
# ---------------------------------------------------------------------------


def _json_indicadores(i: Indicadores) -> dict[str, int | float | None]:
    return {clave: i.valor(clave) for clave in CLAVES}


def _json_plazo(p: solicitudes.Plazo | None) -> dict[str, object] | None:
    return None if p is None else {"dias": p.dias, "fuente": p.fuente}


def a_json(r: Reporte) -> dict[str, object]:
    dimensiones: dict[str, list[dict[str, object]]] = {}
    for g in r.grupos:
        dimensiones.setdefault(g.dimension, []).append(
            {"nombre": g.nombre, "indicadores": _json_indicadores(g.indicadores)}
        )
    return {
        "periodo": {
            "etiqueta": r.periodo.etiqueta,
            "texto": r.periodo.texto,
            "desde": r.periodo.desde.isoformat(),
            "hasta": (r.periodo.hasta - timedelta(days=1)).isoformat(),
        },
        "alcance": r.alcance,
        "corte": r.hoy.isoformat(),
        "plazos": {
            "comite": _json_plazo(r.plazos.comite),
            "consejo": _json_plazo(r.plazos.consejo),
        },
        "total": _json_indicadores(r.total),
        "dimensiones": [
            {
                "nombre": d,
                "definicion": DIMENSIONES[d].definicion,
                "fuente": DIMENSIONES[d].fuente,
                "grupos": dimensiones[d],
            }
            for d in DIMENSIONES
            if d in dimensiones
        ],
        "fuentes": {
            clave: {"nombre": f.nombre, "definicion": f.definicion, "fuente": f.fuente}
            for clave, f in FUENTES.items()
        },
        "nota": _NOTA_FINAL.removeprefix("<i>").removesuffix("</i>"),
    }


def _celda(valor: int | float | None) -> int | float | str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return int(valor)
    return valor


def xlsx(r: Reporte, cols: list[Columna] | None = None) -> bytes:
    """Excel con la pestaña «indicadores» y la pestaña «fuentes»."""
    cols = cols if cols is not None else columnas()
    libro = Workbook()
    hoja = libro.active
    assert hoja is not None
    hoja.title = "indicadores"
    negrita = Font(bold=True)
    hoja.append(["Dimensión", "Grupo", *(c.encabezado for c in cols)])
    for celda in hoja[1]:
        celda.font = negrita
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    filas: list[tuple[str, str, Indicadores]] = [("Total", r.alcance, r.total)]
    filas += [(g.dimension, g.nombre, g.indicadores) for g in r.grupos]
    for dimension, nombre, i in filas:
        hoja.append([dimension, nombre, *(_celda(i.valor(c.clave)) for c in cols)])
    hoja.column_dimensions["A"].width = 18
    hoja.column_dimensions["B"].width = 42
    hoja.freeze_panes = "C2"

    fuentes = libro.create_sheet("fuentes")
    info: list[list[str]] = [
        ["Reporte", "Indicadores del servicio de actas (MemorIA)"],
        ["Alcance", r.alcance],
        ["Periodo", r.periodo.texto],
        ["Corte", r.hoy.isoformat()],
        [
            "Casos incluidos",
            "Recibidos en el periodo, con su estado a la fecha de corte",
        ],
    ]
    for etiqueta, plazo in (
        ("Plazo del comité", r.plazos.comite),
        ("Plazo del Consejo", r.plazos.consejo),
    ):
        info.append(
            [
                etiqueta,
                f"{plazo.dias} días calendario ({plazo.fuente})"
                if plazo
                else "Sin plazo configurado en config/solicitudes.yaml",
            ]
        )
    info += [
        [
            "Plantilla",
            "Columnas provisionales (config/indicadores.yaml): la plantilla "
            "oficial de indicadores de Calidad todavía no llega.",
        ],
        ["Nota", _NOTA_FINAL.removeprefix("<i>").removesuffix("</i>")],
        [],
    ]
    for fila in info:
        fuentes.append(fila)
    fuentes.append(["Columna", "Indicador", "Definición", "Fuente"])
    for celda in fuentes[fuentes.max_row]:
        celda.font = negrita
    for c in cols:
        f = FUENTES[c.clave]
        fuentes.append([c.encabezado, c.clave, f.definicion, f.fuente])
    fuentes.append([])
    fuentes.append(["Dimensión", "", "Definición", "Fuente"])
    for celda in fuentes[fuentes.max_row]:
        celda.font = negrita
    for nombre, f in DIMENSIONES.items():
        fuentes.append([nombre, "", f.definicion, f.fuente])
    for letra, ancho in (("A", 34), ("B", 24), ("C", 70), ("D", 70)):
        fuentes.column_dimensions[letra].width = ancho
    for fila_celdas in fuentes.iter_rows():
        for celda in fila_celdas:
            celda.alignment = Alignment(wrap_text=True, vertical="top")
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def nombre_xlsx(r: Reporte) -> str:
    alcance = "facultad"
    if r.alcance != FACULTAD:
        alcance = re.sub(r"[^a-z0-9]+", "-", normalizar(r.alcance)).strip("-")
    return f"indicadores-actas-{alcance}-{r.periodo.desde:%Y-%m-%d}.xlsx"
