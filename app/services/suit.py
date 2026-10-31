"""Datos de operación (SUIT) que la Secretaría de Facultad reporta cada trimestre.

La Secretaría diligencia en Drive («Datos de Operación 2026») un Excel por
trámite con el número de solicitudes resueltas por mes y por Facultad. De esos
numerales, MemorIA puede contar los que pasan por los comités asesores y el
Consejo de Facultad (ver `NUMERALES`).

Reglas:
- Solo cifras: ni nombres ni documentos.
- «Resuelta» = decidida por el Consejo de Facultad en un acta confirmada en
  MemorIA (aprobó, no aprobó u otra decisión). Un aplazamiento no la resuelve.
- El mes es el de la sesión del Consejo que consta en el acta.
- Un caso con varios trámites cuenta una vez en cada numeral que le aplique.
- Solo cuenta lo que pasó por MemorIA: lo que llegó por otra vía no aparece.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.conexion import Conexion
from app.core.programas import PROGRAMAS
from app.schemas.caso import EstadoCaso, TipoSolicitud
from app.schemas.consejo import Instancia, SentidoConsejo

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FACULTAD = "Ciencias Exactas y Naturales"
MESES = ("Ene", "Feb", "Mar", "Abr", "May", "Jun",
         "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")  # fmt: skip


@dataclass(frozen=True)
class Numeral:
    numero: int
    titulo: str  # como en el nombre del Excel de Drive
    tipos: frozenset[TipoSolicitud]
    solo_pregrado: bool = False


NUMERALES: tuple[Numeral, ...] = (
    Numeral(
        6, "Reingreso a un Programa Académico", frozenset({TipoSolicitud.REINGRESO})
    ),
    Numeral(9, "Aplazamiento de Semestre", frozenset({TipoSolicitud.RESERVA_CUPO})),
    Numeral(
        13,
        "Registro de Asignaturas",
        frozenset({TipoSolicitud.INSCRIPCION_EXTEMPORANEA, TipoSolicitud.CANCELACION}),
    ),
    Numeral(
        22,
        "Cancelación de la Matrícula Académica",
        frozenset({TipoSolicitud.CANCELACION_PERIODO}),
    ),
    Numeral(
        23,
        "Transferencia de Estudiantes de Pregrado",
        frozenset({TipoSolicitud.TRASLADO}),
        solo_pregrado=True,
    ),
)

# Numerales del listado de la Secretaría que no pasan por los comités: salen
# de otros sistemas (SIA, admisiones, registro) y MemorIA no los cuenta.
FUERA_DE_MEMORIA = (
    "2. Contenido del programa académico (certificado)",
    "3. Renovación de matrícula de estudiantes",
    "8. Inscripción de aspirantes a programas de pregrado",
    "12. Matrícula de admitidos a programas de posgrado",
    "14. Grado de pregrado y posgrado",
    "15. Certificados de notas",
    "16. Certificados y constancias de estudio",
    "18. Inscripción de aspirantes a programas de posgrado",
    "19. Duplicaciones de diplomas y actas",
    "20. Cursos intersemestrales",
    "21. Matrícula de admitidos a programas de pregrado",
)

RESUELTAS = (SentidoConsejo.APROBO, SentidoConsejo.NO_APROBO, SentidoConsejo.OTRO)


@dataclass
class Conteo:
    numeral: Numeral
    # mes (1-12) -> sentido -> número de solicitudes
    por_mes: dict[int, dict[SentidoConsejo, int]] = field(default_factory=dict)

    def total_mes(self, mes: int) -> int:
        return sum(self.por_mes.get(mes, {}).values())

    @property
    def total(self) -> int:
        return sum(self.total_mes(m) for m in range(1, 13))


@dataclass
class ReporteSuit:
    ano: int
    conteos: list[Conteo]
    corte: date


class AnoInvalido(ValueError):
    pass


def ano_de(texto: str, hoy: date) -> int:
    texto = texto.strip()
    if not texto:
        return hoy.year
    if not texto.isdigit() or not 2000 <= int(texto) <= 2100:
        raise AnoInvalido("Escriba el año con cuatro cifras, por ejemplo /suit 2026.")
    return int(texto)


def trimestre(mes: int) -> int:
    return (mes - 1) // 3 + 1


def _tipos(extraccion_json: str | None) -> set[TipoSolicitud]:
    try:
        datos = json.loads(extraccion_json or "")
    except ValueError:
        return set()
    salida: set[TipoSolicitud] = set()
    for t in datos.get("tipos", []) if isinstance(datos, dict) else []:
        try:
            salida.add(TipoSolicitud(t))
        except ValueError:
            continue
    return salida


async def reporte(db: Conexion, ano: int, hoy: date) -> ReporteSuit:
    conteos = [Conteo(n) for n in NUMERALES]
    cur = await db.execute(
        "SELECT c.programa, c.extraccion_json, d.sentido, d.acta_fecha"
        " FROM casos c JOIN decisiones_consejo d ON d.caso_id = c.id"
        " WHERE c.estado != ? AND d.instancia = ? AND d.acta_fecha LIKE ?",
        (EstadoCaso.DESCARTADO.value, Instancia.FACULTAD.value, f"{ano}-%"),
    )
    for programa, extraccion, sentido_txt, acta_fecha in await cur.fetchall():
        try:
            sentido = SentidoConsejo(sentido_txt)
            mes = date.fromisoformat(str(acta_fecha)[:10]).month
        except ValueError:
            continue
        if sentido not in RESUELTAS:
            continue
        tipos = _tipos(extraccion)
        prog = PROGRAMAS.get(programa)
        for c in conteos:
            if not tipos & c.numeral.tipos:
                continue
            if c.numeral.solo_pregrado and (prog is None or prog.nivel != "pregrado"):
                continue
            mes_c = c.por_mes.setdefault(mes, {})
            mes_c[sentido] = mes_c.get(sentido, 0) + 1
    return ReporteSuit(ano, conteos, hoy)


def formatear(r: ReporteSuit) -> str:
    lineas = [
        f"<b>Datos SUIT {r.ano}</b> — Facultad de {FACULTAD}",
        "Solicitudes resueltas por el Consejo de Facultad (acta confirmada en "
        "MemorIA), por mes de la sesión.",
        "",
    ]
    for c in r.conteos:
        meses = [
            f"{MESES[m - 1]} {c.total_mes(m)}" for m in range(1, 13) if c.total_mes(m)
        ]
        detalle = ", ".join(meses) if meses else "sin solicitudes"
        lineas.append(
            f"<b>{c.numeral.numero}. {c.numeral.titulo}</b>: {c.total} ({detalle})"
        )
    lineas += [
        "",
        "Solo cuenta lo que pasó por MemorIA con el acta del Consejo confirmada. "
        "El Excel trae la fila lista para copiar y el detalle por decisión.",
    ]
    return "\n".join(lineas)


_GRIS = PatternFill("solid", fgColor="D9D9D9")


def xlsx(r: ReporteSuit) -> bytes:
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.title = "Para copiar"
    ws.append([f"DATOS DE OPERACIÓN {r.ano} — Facultad de {FACULTAD}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(
        ["NÚMERO DE SOLICITUDES RESUELTAS (decididas por el Consejo de Facultad)"]
    )
    ws.append([])
    encabezado = ["Numeral", "Trámite", "Facultad"]
    encabezado += [f"{m} (T{trimestre(i)})" for i, m in enumerate(MESES, start=1)]
    encabezado.append("Total")
    ws.append(encabezado)
    for celda in ws[4]:
        celda.font = Font(bold=True)
        celda.fill = _GRIS
        celda.alignment = Alignment(wrap_text=True, horizontal="center")
    for c in r.conteos:
        ws.append(
            [c.numeral.numero, c.numeral.titulo, FACULTAD]
            + [c.total_mes(m) for m in range(1, 13)]
            + [c.total]
        )
    ws.column_dimensions["A"].width = 9
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 28
    ws.freeze_panes = "D5"

    det = libro.create_sheet("Detalle por decisión")
    det.append(
        ["Numeral", "Trámite", "Mes", "Aprobó", "No aprobó", "Otra decisión", "Total"]
    )
    for celda in det[1]:
        celda.font = Font(bold=True)
        celda.fill = _GRIS
    for c in r.conteos:
        for m in range(1, 13):
            s = c.por_mes.get(m)
            if not s:
                continue
            det.append(
                [
                    c.numeral.numero,
                    c.numeral.titulo,
                    MESES[m - 1],
                    s.get(SentidoConsejo.APROBO, 0),
                    s.get(SentidoConsejo.NO_APROBO, 0),
                    s.get(SentidoConsejo.OTRO, 0),
                    c.total_mes(m),
                ]
            )
    det.column_dimensions["B"].width = 42

    leeme = libro.create_sheet("Léeme")
    leeme.column_dimensions["A"].width = 120
    for linea in [
        f"Datos SUIT {r.ano} — generado por MemorIA el {r.corte.isoformat()}",
        "Cuenta las solicitudes estudiantiles resueltas por el Consejo de Facultad: "
        "las que tienen la decisión registrada desde un acta del Consejo confirmada "
        "en MemorIA (aprobó, no aprobó u otra decisión). Los aplazamientos no "
        "cuentan. El mes es el de la sesión del Consejo.",
        "Correspondencia con los numerales: 6 = reingreso; 9 = reserva de cupo "
        "(aplazamiento de semestre); 13 = inscripción extemporánea y cancelación de "
        "asignaturas; 22 = cancelación del periodo académico; 23 = traslado (solo "
        "programas de pregrado). Un caso con varios trámites cuenta en cada numeral.",
        "La hoja «Para copiar» va en el bloque «Número de solicitudes resueltas "
        "parcialmente en línea» o «de forma presencial», según defina la Secretaría.",
        "Solo aparece lo que pasó por MemorIA. Revise contra sus registros antes de "
        "copiar las cifras al Excel de Drive.",
        "Numerales que MemorIA no cuenta (no pasan por los comités): "
        + "; ".join(FUERA_DE_MEMORIA)
        + ".",
    ]:
        leeme.append([linea])
    leeme["A1"].font = Font(bold=True, size=13)
    for fila in leeme.iter_rows(min_row=2):
        fila[0].alignment = Alignment(wrap_text=True)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def nombre_xlsx(r: ReporteSuit) -> str:
    return f"datos_suit_fcen_{r.ano}.xlsx"
