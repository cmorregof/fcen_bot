"""Tiempos de respuesta de las notificaciones oficiales de la Facultad.

La oficina de Calidad pide cada semestre, para el indicador de divulgación de
la información oficial, el tiempo entre:

1. el ingreso del caso a la Secretaría de Facultad,
2. la expedición del oficio o acto administrativo de respuesta, y
3. su notificación.

La Secretaría registra en MemorIA, por cada caso decidido por el Consejo de
Facultad, el acto (oficio o resolución), su número y esas tres fechas. Como
sugerencia de la fecha de ingreso se muestra la de la recomendación del comité
asesor; solo cuenta lo que la Secretaría guarda.

El reporte es de solo cifras: casos por número, sin nombres ni documentos.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass
from datetime import date
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from pydantic import BaseModel, field_validator

from app.core.conexion import Conexion
from app.core.programas import PROGRAMAS
from app.schemas.caso import TITULOS_TIPO, EstadoCaso, TipoSolicitud
from app.schemas.consejo import (
    ETIQUETA_SENTIDO_CONSEJO,
    Instancia,
    SentidoConsejo,
)
from app.services import flujo
from app.services.indicadores import Estadistica, estadistica

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
Acto = Literal["oficio", "resolucion"]
NOMBRE_ACTO = {"oficio": "Oficio", "resolucion": "Resolución"}


class PeriodoInvalido(ValueError):
    pass


@dataclass(frozen=True)
class Semestre:
    texto: str
    desde: date
    hasta: date


def semestre_de(texto: str, hoy: date) -> Semestre:
    """«2026-1S» (enero a junio) o «2026-2S» (julio a diciembre); vacío = actual.

    El semestre es calendario y se toma por la fecha de la sesión del Consejo.
    """
    texto = texto.strip().upper()
    if not texto:
        texto = f"{hoy.year}-{1 if hoy.month <= 6 else 2}S"
    partes = texto.split("-")
    if (
        len(partes) != 2
        or not partes[0].isdigit()
        or partes[1] not in ("1S", "2S")
        or not 2000 <= int(partes[0]) <= 2100
    ):
        raise PeriodoInvalido("Escriba el semestre así: 2026-1S o 2026-2S.")
    ano = int(partes[0])
    if partes[1] == "1S":
        return Semestre(texto, date(ano, 1, 1), date(ano, 6, 30))
    return Semestre(texto, date(ano, 7, 1), date(ano, 12, 31))


class Registro(BaseModel):
    """Lo que guarda la Secretaría para un caso."""

    acto: Acto = "oficio"
    numero: str = ""
    fecha_ingreso: date | None = None
    fecha_expedicion: date | None = None
    fecha_notificacion: date | None = None

    @field_validator("numero")
    @classmethod
    def _numero(cls, v: str) -> str:
        return v.strip()[:80]

    def error(self, hoy: date) -> str | None:
        fechas = [
            ("ingreso", self.fecha_ingreso),
            ("expedición", self.fecha_expedicion),
            ("notificación", self.fecha_notificacion),
        ]
        for nombre, f in fechas:
            if f is not None and f > hoy:
                return f"La fecha de {nombre} no puede ser posterior a hoy."
        conocidas = [(n, f) for n, f in fechas if f is not None]
        for (n1, f1), (n2, f2) in zip(conocidas, conocidas[1:], strict=False):
            if f2 < f1:
                return f"La fecha de {n2} no puede ser anterior a la de {n1}."
        return None


@dataclass
class Fila:
    caso_id: int
    programa: str
    tramite: str
    estudiante: str  # solo para la pantalla; nunca va al Excel
    decision: str
    acta_numero: str
    acta_fecha: date
    sugerida_ingreso: date | None
    registro: Registro | None
    registrado_por: str = ""

    def dias(self) -> tuple[int | None, int | None, int | None]:
        r = self.registro
        if r is None:
            return None, None, None

        def entre(a: date | None, b: date | None) -> int | None:
            return (b - a).days if a and b else None

        return (
            entre(r.fecha_ingreso, r.fecha_expedicion),
            entre(r.fecha_expedicion, r.fecha_notificacion),
            entre(r.fecha_ingreso, r.fecha_notificacion),
        )

    @property
    def completo(self) -> bool:
        r = self.registro
        return bool(
            r and r.fecha_ingreso and r.fecha_expedicion and r.fecha_notificacion
        )


def _dia(valor: object) -> date | None:
    try:
        return date.fromisoformat(str(valor)[:10]) if valor else None
    except ValueError:
        return None


async def listar(db: Conexion, semestre: Semestre) -> list[Fila]:
    """Casos decididos por el Consejo de Facultad en el semestre (sin aplazados)."""
    cur = await db.execute(
        "SELECT c.id, c.programa, c.extraccion_json, c.recomendado,"
        " d.sentido, d.acta_numero, d.acta_fecha,"
        " n.acto, n.numero, n.fecha_ingreso, n.fecha_expedicion,"
        " n.fecha_notificacion, n.registrado_por"
        " FROM casos c JOIN decisiones_consejo d ON d.caso_id = c.id"
        " LEFT JOIN notificaciones n ON n.caso_id = c.id"
        " WHERE c.estado != ? AND d.instancia = ? AND d.sentido != ?"
        " ORDER BY d.acta_fecha, c.id",
        (
            EstadoCaso.DESCARTADO.value,
            Instancia.FACULTAD.value,
            SentidoConsejo.APLAZO.value,
        ),
    )
    filas: list[Fila] = []
    for r in await cur.fetchall():
        acta_fecha = _dia(r[6])
        if acta_fecha is None or not semestre.desde <= acta_fecha <= semestre.hasta:
            continue
        try:
            datos = json.loads(r[2] or "{}")
        except ValueError:
            datos = {}
        tipos = datos.get("tipos") or []
        try:
            tramite = TITULOS_TIPO[TipoSolicitud(tipos[0])] if tipos else "sin ficha"
        except ValueError:
            tramite = TITULOS_TIPO[TipoSolicitud.OTRO]
        registro = None
        if r[7] is not None:
            registro = Registro(
                acto=r[7],
                numero=r[8] or "",
                fecha_ingreso=_dia(r[9]),
                fecha_expedicion=_dia(r[10]),
                fecha_notificacion=_dia(r[11]),
            )
        filas.append(
            Fila(
                caso_id=int(r[0]),
                programa=PROGRAMAS[r[1]].nombre if r[1] in PROGRAMAS else "—",
                tramite=tramite,
                estudiante=str(datos.get("nombre_estudiante") or ""),
                decision=ETIQUETA_SENTIDO_CONSEJO[SentidoConsejo(r[4])],
                acta_numero=str(r[5] or ""),
                acta_fecha=acta_fecha,
                sugerida_ingreso=_dia(r[3]),
                registro=registro,
                registrado_por=str(r[12] or ""),
            )
        )
    return filas


async def guardar(
    db: Conexion,
    caso_id: int,
    registro: Registro,
    usuario: str,
    *,
    rol: str = "secretaria",
) -> None:
    """Guarda el acto y sus fechas, y lleva la línea del caso a expedida o
    notificada (app/services/flujo.py). Si la línea no lo admite, no guarda."""

    def iso(d: date | None) -> str | None:
        return d.isoformat() if d else None

    await db.execute(
        "INSERT INTO notificaciones (caso_id, acto, numero, fecha_ingreso,"
        " fecha_expedicion, fecha_notificacion, registrado_por)"
        " VALUES (?, ?, ?, ?, ?, ?, ?)"
        " ON CONFLICT(caso_id) DO UPDATE SET acto = excluded.acto,"
        " numero = excluded.numero, fecha_ingreso = excluded.fecha_ingreso,"
        " fecha_expedicion = excluded.fecha_expedicion,"
        " fecha_notificacion = excluded.fecha_notificacion,"
        " registrado_por = excluded.registrado_por,"
        " actualizado = CURRENT_TIMESTAMP",
        (
            caso_id,
            registro.acto,
            registro.numero,
            iso(registro.fecha_ingreso),
            iso(registro.fecha_expedicion),
            iso(registro.fecha_notificacion),
            usuario,
        ),
    )
    try:
        await flujo.sincronizar_secretaria(
            db,
            caso_id,
            numero=registro.numero,
            fecha_expedicion=iso(registro.fecha_expedicion),
            fecha_notificacion=iso(registro.fecha_notificacion),
            rol=rol,
            persona=usuario,
        )
    except flujo.TransicionNoPermitida:
        await db.rollback()
        raise
    await db.commit()


async def decidido_por_facultad(db: Conexion, caso_id: int) -> bool:
    cur = await db.execute(
        "SELECT 1 FROM decisiones_consejo d JOIN casos c ON c.id = d.caso_id"
        " WHERE d.caso_id = ? AND d.instancia = ? AND d.sentido != ?"
        " AND c.estado != ?",
        (
            caso_id,
            Instancia.FACULTAD.value,
            SentidoConsejo.APLAZO.value,
            EstadoCaso.DESCARTADO.value,
        ),
    )
    return await cur.fetchone() is not None


@dataclass
class Resumen:
    decididos: int
    completos: int
    ingreso_expedicion: Estadistica
    expedicion_notificacion: Estadistica
    total: Estadistica


def resumir(filas: list[Fila]) -> Resumen:
    a: list[int] = []
    b: list[int] = []
    c: list[int] = []
    for f in filas:
        d1, d2, d3 = f.dias()
        if d1 is not None:
            a.append(d1)
        if d2 is not None:
            b.append(d2)
        if d3 is not None:
            c.append(d3)
    return Resumen(
        decididos=len(filas),
        completos=sum(1 for f in filas if f.completo),
        ingreso_expedicion=estadistica(a),
        expedicion_notificacion=estadistica(b),
        total=estadistica(c),
    )


def _json_est(e: Estadistica) -> dict[str, object]:
    return {"casos": e.casos, "mediana": e.mediana, "p90": e.p90}


def a_json(semestre: Semestre, filas: list[Fila]) -> dict[str, object]:
    r = resumir(filas)
    return {
        "semestre": semestre.texto,
        "desde": semestre.desde.isoformat(),
        "hasta": semestre.hasta.isoformat(),
        "resumen": {
            "decididos": r.decididos,
            "completos": r.completos,
            "ingreso_expedicion": _json_est(r.ingreso_expedicion),
            "expedicion_notificacion": _json_est(r.expedicion_notificacion),
            "total": _json_est(r.total),
        },
        "casos": [
            {
                "caso_id": f.caso_id,
                "programa": f.programa,
                "tramite": f.tramite,
                "estudiante": f.estudiante,
                "decision": f.decision,
                "acta_numero": f.acta_numero,
                "acta_fecha": f.acta_fecha.isoformat(),
                "sugerida_ingreso": (
                    f.sugerida_ingreso.isoformat() if f.sugerida_ingreso else None
                ),
                "registro": f.registro.model_dump(mode="json") if f.registro else None,
                "registrado_por": f.registrado_por,
            }
            for f in filas
        ],
    }


_GRIS = PatternFill("solid", fgColor="D9D9D9")
_AMARILLO = PatternFill("solid", fgColor="FFFF00")


def xlsx(semestre: Semestre, filas: list[Fila], corte: date) -> bytes:
    r = resumir(filas)
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.title = "Resumen"
    ws.append([f"Tiempos de respuesta de las notificaciones — {semestre.texto}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(
        [
            f"Casos decididos por el Consejo de Facultad entre {semestre.desde} y {semestre.hasta}. Corte: {corte}."
        ]
    )
    ws.append([])
    ws.append(["Indicador", "Casos", "Mediana (días)", "Percentil 90 (días)"])
    for celda in ws[4]:
        celda.font = Font(bold=True)
        celda.fill = _GRIS
    for nombre, e in [
        ("Ingreso a la Secretaría → expedición del acto", r.ingreso_expedicion),
        ("Expedición → notificación", r.expedicion_notificacion),
        ("Ingreso a la Secretaría → notificación", r.total),
    ]:
        ws.append([nombre, e.casos, e.mediana, e.p90])
    ws.append([])
    ws.append(["Casos decididos en el semestre", r.decididos])
    ws.append(["Con las tres fechas registradas", r.completos])
    ws.column_dimensions["A"].width = 48
    for col in "BCD":
        ws.column_dimensions[col].width = 18

    casos = libro.create_sheet("Casos")
    encabezado = [
        "Caso", "Programa", "Trámite", "Decisión", "Acta", "Fecha de la sesión",
        "Acto", "Número", "Ingreso a la Secretaría", "Expedición", "Notificación",
        "Días ingreso → expedición", "Días expedición → notificación",
        "Días ingreso → notificación",
    ]  # fmt: skip
    casos.append(encabezado)
    for celda in casos[1]:
        celda.font = Font(bold=True)
        celda.fill = _GRIS
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    for f in filas:
        reg = f.registro
        d1, d2, d3 = f.dias()
        casos.append(
            [
                f.caso_id,
                f.programa,
                f.tramite,
                f.decision,
                f.acta_numero,
                f.acta_fecha,
                NOMBRE_ACTO[reg.acto] if reg else "",
                reg.numero if reg else "",
                reg.fecha_ingreso if reg else None,
                reg.fecha_expedicion if reg else None,
                reg.fecha_notificacion if reg else None,
                d1,
                d2,
                d3,
            ]
        )
        if not f.completo:
            for n in (9, 10, 11):
                if casos.cell(casos.max_row, n).value is None:
                    casos.cell(casos.max_row, n).fill = _AMARILLO
    for col, ancho in zip(
        "ABCDEFGHIJKLMN",
        [7, 30, 34, 13, 7, 12, 11, 22, 12, 12, 12, 12, 12, 12],
        strict=True,
    ):
        casos.column_dimensions[col].width = ancho
    for fila in casos.iter_rows(min_row=2):
        for celda in fila:
            if isinstance(celda.value, date):
                celda.number_format = "DD/MM/YYYY"
    casos.freeze_panes = "B2"

    leeme = libro.create_sheet("Léeme")
    leeme.column_dimensions["A"].width = 120
    for linea in [
        "Indicador de divulgación de la información oficial — tiempos de respuesta",
        "Una fila por caso decidido por el Consejo de Facultad en el semestre (por "
        "la fecha de la sesión). Los aplazados no entran. Sin nombres ni documentos.",
        "Las fechas las registra la Secretaría de Facultad en MemorIA: ingreso del "
        "caso a la Secretaría, expedición del oficio o resolución y notificación.",
        "Días calendario. Mediana y percentil 90 por rango más cercano, solo con "
        "los casos que tienen las dos fechas de cada tramo.",
        "En amarillo: fechas que faltan por registrar.",
    ]:
        leeme.append([linea])
    leeme["A1"].font = Font(bold=True, size=13)
    for fila in leeme.iter_rows(min_row=2):
        fila[0].alignment = Alignment(wrap_text=True)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def nombre_xlsx(semestre: Semestre) -> str:
    return f"tiempos_notificaciones_{semestre.texto}.xlsx"
