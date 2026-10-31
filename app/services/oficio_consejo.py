"""Borrador del oficio con que la Secretaría de Facultad notifica una decisión
del Consejo de Facultad al estudiante (y, para el reingreso aprobado, de la
resolución).

Sale de un caso con la decisión del Consejo ya confirmada por una persona
(`DecisionConsejo`, leída del acta). Es determinista: no se llama al modelo. El
sentido es el que consta en el acta; aquí no se decide nada. Lo que MemorIA no
sabe (número del oficio, modalidad de la sesión, motivación de una negación,
firma) sale resaltado en amarillo para que la Secretaría lo complete antes de
firmar. Los textos fijos (notificación y recursos) son los de los oficios de
la Secretaría y no se parafrasean. Ver docs/oficios-consejo.md.
"""

from __future__ import annotations

import io
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date

from docx import Document
from docx.document import Document as DocumentoDocx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from app.core.config import settings
from app.core.programas import PROGRAMAS
from app.schemas.caso import (
    TITULOS_TIPO,
    AsignaturaSolicitada,
    Caso,
    ExtraccionCaso,
    Sentido,
    TipoSolicitud,
)
from app.schemas.consejo import DecisionConsejo, Instancia, SentidoConsejo
from app.services.documento import (
    PLANTILLA,
    _equivalencias,
    _jurados,
    _p,
    _proyecto,
    _r,
    _reemplazar_marcadores,
    _vacio,
    en_letras,
    fecha_larga,
)

# ---------------------------------------------------------------------------
# Textos fijos (literales de los oficios de la Secretaría de Facultad)
# ---------------------------------------------------------------------------

NOTIFICACION = (
    "Se le notifica a través de la Secretaría de Facultad de Ciencias Exactas y "
    "Naturales del contenido del presente oficio, de conformidad con los "
    "artículos 6 y ss de la Resolución No. 1194 de 2013, expedida por la "
    "Rectoría, por la cual se reglamenta las notificaciones de los actos "
    "académicos de carácter particular en la Universidad Nacional de Colombia."
)

RECURSOS = (
    "De igual manera, le informo que contra esta decisión procede el recurso de "
    "reposición ante el Consejo de Facultad, y directa o subsidiariamente el de "
    "apelación ante el Consejo de Sede, recursos que deben ser presentados con "
    "el lleno de los requisitos exigidos por el Artículo 77 del Código de "
    "Procedimiento Administrativo y de lo Contencioso Administrativo (Ley 1437 de "
    "2011), dentro de los diez (10) días siguientes a la notificación del "
    "presente acto."
)

RADICADO_PREFIJO = "Mz.FCEN.1.004"
CONSEJO = "El Consejo de la Facultad de Ciencias Exactas y Naturales"

# Marcadores (en amarillo) de lo que MemorIA no sabe.
M_NUMERO = "___"
M_MODALIDAD = "[presencial / no presencial]"
M_MOTIVACION = "[motivación de la decisión según el acta]"
M_SECRETARIO = "[Nombre del Secretario(a) de Facultad]"
M_PRESIDENTE = "[Nombre del Presidente(a) del Consejo de Facultad]"

_MESES = (
    "enero febrero marzo abril mayo junio julio agosto septiembre octubre "
    "noviembre diciembre"
).split()


class SinOficio(ValueError):
    """El caso no tiene una decisión del Consejo de Facultad que notificar."""


# ---------------------------------------------------------------------------
# Texto con partes resaltadas
# ---------------------------------------------------------------------------

# Cada parte es (texto, estilo): "" normal, "b" negrita, "y" amarillo (falta).
Parte = tuple[str, str]


def _n(texto: str) -> Parte:
    return (texto, "")


def _b(texto: str) -> Parte:
    return (texto, "b")


def _y(texto: str) -> Parte:
    return (texto, "y")


def _escribir(
    p: Paragraph, partes: Sequence[Parte], *, todo_amarillo: bool = False
) -> None:
    for texto, estilo in partes:
        if not texto:
            continue
        r = _r(p, texto, negrita=estilo in ("b", "y"))
        if estilo == "y" or todo_amarillo:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW


def _parrafo(
    doc: DocumentoDocx,
    partes: Sequence[Parte],
    estilo: str = "FCEN Texto",
    *,
    todo_amarillo: bool = False,
) -> Paragraph:
    p = _p(doc, estilo)
    _escribir(p, partes, todo_amarillo=todo_amarillo)
    return p


def _o(valor: str | None, falta: str) -> Parte:
    """El valor si se conoce; si no, el marcador en amarillo."""
    return _n(valor) if valor else _y(falta)


def _ob(valor: str | None, falta: str) -> Parte:
    return _b(valor) if valor else _y(falta)


def _y_lista(partes: list[list[Parte]]) -> list[Parte]:
    """«a, b y c» (o «e» ante palabra que suena /i/), con partes."""
    salida: list[Parte] = []
    for i, grupo in enumerate(partes):
        if i:
            if i < len(partes) - 1:
                salida.append(_n(", "))
            else:
                primera = grupo[0][0].lower() if grupo and grupo[0][0] else ""
                ie = primera.startswith(("i", "hi")) and not primera.startswith("hie")
                salida.append(_n(" e " if ie else " y "))
        salida.extend(grupo)
    return salida


# ---------------------------------------------------------------------------
# Datos del caso
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class _Contexto:
    caso: Caso
    decision: DecisionConsejo
    extraccion: ExtraccionCaso
    programa: str | None
    posgrado: bool
    periodo: str | None
    periodo_config: str
    fecha: date

    @property
    def aprobo(self) -> bool:
        return self.decision.sentido == SentidoConsejo.APROBO

    def periodo_partes(self) -> Parte:
        """El periodo del SIA; si no aparece, el configurado, en amarillo."""
        return _n(self.periodo) if self.periodo else _y(self.periodo_config)


def motivo_sin_oficio(
    caso: Caso, decision: DecisionConsejo | None = None
) -> str | None:
    """Por qué no se puede armar el borrador; None si se puede."""
    d = decision or caso.consejo
    if d is None:
        return "El caso no tiene una decisión del Consejo confirmada."
    if d.instancia != Instancia.FACULTAD:
        return (
            "La decisión registrada no es del Consejo de Facultad: la notifica "
            "otra instancia."
        )
    if d.sentido not in (SentidoConsejo.APROBO, SentidoConsejo.NO_APROBO):
        return (
            "El Consejo no aprobó ni negó la solicitud (aplazó u otra decisión): "
            "no hay nada que notificar todavía."
        )
    if caso.extraccion is None:
        return "El caso no se ha leído."
    return None


def _fecha_iso(valor: str | None) -> date | None:
    if not valor:
        return None
    try:
        return date.fromisoformat(valor[:10])
    except ValueError:
        return None


def _contexto(
    caso: Caso,
    decision: DecisionConsejo | None,
    fecha: date | None,
    periodo: str | None,
) -> _Contexto:
    d = decision or caso.consejo
    motivo = motivo_sin_oficio(caso, d)
    if motivo:
        raise SinOficio(motivo)
    assert d is not None and caso.extraccion is not None
    e = caso.extraccion
    programa = PROGRAMAS.get(caso.programa or "")
    return _Contexto(
        caso=caso,
        decision=d,
        extraccion=e,
        programa=programa.nombre if programa else None,
        posgrado=bool(programa and programa.nivel == "posgrado"),
        periodo=e.sia.periodo if e.sia and e.sia.periodo else None,
        periodo_config=periodo or settings.periodo_actual,
        # Como en los ejemplos de la Secretaría: el oficio lleva la fecha de la
        # sesión; sin fecha de acta, la del día.
        fecha=fecha or _fecha_iso(d.acta_fecha) or date.today(),
    )


# ---------------------------------------------------------------------------
# Trámites
# ---------------------------------------------------------------------------


def _asignaturas(e: ExtraccionCaso, accion: str) -> list[Parte]:
    """«la asignatura NOMBRE (código)» / «las asignaturas A (c1) y B (c2)»."""
    lista: list[AsignaturaSolicitada] = [a for a in e.asignaturas if a.accion == accion]
    if not lista:
        return [_n("de la asignatura "), _y("[NOMBRE DE LA ASIGNATURA (código)]")]
    grupos: list[list[Parte]] = []
    for a in lista:
        grupo: list[Parte] = [_b(f"{a.nombre.strip().upper()} (")]
        grupo.append(_b(a.codigo) if a.codigo else _y("[código]"))
        grupo.append(_b(")"))
        grupos.append(grupo)
    articulo = "de la asignatura " if len(lista) == 1 else "de las asignaturas "
    return [_n(articulo), *_y_lista(grupos)]


def _en_periodo(ctx: _Contexto) -> list[Parte]:
    return [_n(" en el periodo Académico "), ctx.periodo_partes()]


@dataclass(frozen=True)
class _Tramite:
    asunto: str
    lleva_periodo: bool
    registro: bool  # copia a la Oficina de Registro y Matrícula


_TRAMITES: dict[TipoSolicitud, _Tramite] = {
    TipoSolicitud.CANCELACION: _Tramite(
        "Cancelación Asignatura sin pérdida de créditos", True, True
    ),
    TipoSolicitud.CANCELACION_PERIODO: _Tramite(
        "Cancelación Periodo Académico", True, True
    ),
    TipoSolicitud.CARGA_MINIMA: _Tramite(
        "Cursar menos de la carga mínima", True, False
    ),
    TipoSolicitud.INSCRIPCION_EXTEMPORANEA: _Tramite(
        "Inscripción Extemporánea de asignatura", False, True
    ),
    TipoSolicitud.TRASLADO: _Tramite("Respuesta a solicitud de Traslado", True, True),
    TipoSolicitud.PERMISO_ACADEMICO: _Tramite("Permiso Académico", False, False),
    TipoSolicitud.HOMOLOGACION: _Tramite("Equivalencia de asignatura", False, True),
    TipoSolicitud.DESIGNACION_JURADOS: _Tramite("Designación de Jurados", False, True),
    TipoSolicitud.PROYECTO_TESIS: _Tramite(
        "Inscripción de proyecto de tesis o trabajo final", False, True
    ),
    TipoSolicitud.REEXPEDICION_RECIBO: _Tramite(
        "Reexpedición de recibo de pago", True, True
    ),
    # Reingreso negado: oficio. El aprobado va en resolución.
    TipoSolicitud.REINGRESO: _Tramite("Respuesta a solicitud de Reingreso", True, True),
}

TRAMITES_CUBIERTOS: frozenset[TipoSolicitud] = frozenset(_TRAMITES)

_ASUNTO_FIGURA = {
    "equivalencia": "Equivalencia de asignatura",
    "convalidacion": "Convalidación asignatura",
    "homologacion": "Homologación de asignatura",
}
_OBJETO_FIGURA = {
    "equivalencia": "equivalencia",
    "convalidacion": "convalidación",
    "homologacion": "homologación",
}


def _figura(e: ExtraccionCaso) -> str:
    return e.equivalencias[0].figura if e.equivalencias else "equivalencia"


def _asunto(tipo: TipoSolicitud, ctx: _Contexto) -> str:
    if tipo == TipoSolicitud.HOMOLOGACION:
        return _ASUNTO_FIGURA[_figura(ctx.extraccion)]
    if tipo == TipoSolicitud.DESIGNACION_JURADOS:
        destino = "Tesis o Trabajo Final" if ctx.posgrado else "Trabajo de Grado"
        return f"Designación de Jurados – {destino}"
    return _TRAMITES[tipo].asunto


# Trámites cuyo párrafo termina «… en el periodo Académico 2026-2S» (una sola
# vez aunque el caso traiga varios).
_PERIODO_AL_FINAL = frozenset(
    {
        TipoSolicitud.CANCELACION,
        TipoSolicitud.CARGA_MINIMA,
        TipoSolicitud.INSCRIPCION_EXTEMPORANEA,
        TipoSolicitud.TRASLADO,
        TipoSolicitud.REEXPEDICION_RECIBO,
    }
)


def _objeto(tipo: TipoSolicitud, ctx: _Contexto) -> list[Parte]:
    """Lo que va después de «la solicitud de» (sin el periodo)."""
    e = ctx.extraccion
    if tipo == TipoSolicitud.CANCELACION:
        return [_n("cancelación "), *_asignaturas(e, "cancelar")]
    if tipo == TipoSolicitud.CANCELACION_PERIODO:
        return [_n("cancelación del periodo Académico "), ctx.periodo_partes()]
    if tipo == TipoSolicitud.CARGA_MINIMA:
        return [_n("cursar menos de la carga mínima")]
    if tipo == TipoSolicitud.INSCRIPCION_EXTEMPORANEA:
        return [_n("inscripción extemporánea "), *_asignaturas(e, "inscribir")]
    if tipo == TipoSolicitud.TRASLADO:
        origen = e.equivalencias[0].origen if e.equivalencias else None
        return [
            _n("traslado del "),
            _o(origen, "[programa y sede de origen]"),
            _n(" al Programa Curricular de "),
            _o(ctx.programa, "[programa]"),
            _n(" – Sede Manizales"),
        ]
    if tipo == TipoSolicitud.PERMISO_ACADEMICO:
        objeto = e.objeto_solicitud.strip().rstrip(".")
        if objeto and "permiso" in objeto.lower():
            return [_n(objeto)]
        return [_n("permiso académico "), _y("[motivo y fechas del permiso]")]
    if tipo == TipoSolicitud.HOMOLOGACION:
        figura = _OBJETO_FIGURA[_figura(e)]
        return [_n(f"{figura} de las asignaturas que se relacionan a continuación")]
    if tipo == TipoSolicitud.DESIGNACION_JURADOS:
        return [_n("designación de los jurados que se relacionan a continuación")]
    if tipo == TipoSolicitud.PROYECTO_TESIS:
        return [_n("inscripción del proyecto de tesis o trabajo final")]
    if tipo == TipoSolicitud.REEXPEDICION_RECIBO:
        return [_n("reexpedición del recibo de pago")]
    if tipo == TipoSolicitud.REINGRESO:
        return [
            _n("reingreso al Programa Curricular de "),
            _o(ctx.programa, "[programa]"),
        ]
    raise KeyError(tipo)


def _canal(caso: Caso) -> Parte:
    """Por dónde llegó la solicitud. El formato adjunto puede haber llegado por
    correo o como oficio: se deja para que la Secretaría lo confirme."""
    if caso.texto.startswith("--- Formato adjunto"):
        return _y("[correo electrónico / oficio]")
    return _n("correo electrónico")


def _fecha_corta(valor: str | None) -> Parte:
    dia = _fecha_iso(valor)
    return _n(f"{dia:%d/%m/%Y}") if dia else _y("[dd/mm/aaaa]")


# ---------------------------------------------------------------------------
# Oficio
# ---------------------------------------------------------------------------


def _documento(encabezado: str, elaboro: str) -> DocumentoDocx:
    doc = Document(str(PLANTILLA))
    _reemplazar_marcadores(
        doc,
        {
            "ENCABEZADO": encabezado,
            "CORREO": settings.secretaria_facultad_correo,
            "EXTENSION": settings.secretaria_facultad_extension or "____",
            "ELABORO": elaboro,
        },
    )
    if doc.paragraphs and not doc.paragraphs[0].text:
        primero = doc.paragraphs[0]._element
        primero.getparent().remove(primero)
    return doc


def _aviso_borrador(doc: DocumentoDocx) -> None:
    p = _p(doc, "FCEN Destinatario")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = _r(
        p,
        "BORRADOR · pendiente de revisión, número y firma de la Secretaría de Facultad",
        negrita=True,
    )
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(0xB0, 0x00, 0x20)
    _vacio(doc)


def _nombre_dni(e: ExtraccionCaso) -> list[Parte]:
    return [
        _ob(
            e.nombre_estudiante.strip().upper() if e.nombre_estudiante else None,
            "[NOMBRE DEL ESTUDIANTE]",
        ),
        _b(" con DNI "),
        _ob(e.documento, "[DNI]"),
    ]


def _encabezado(doc: DocumentoDocx, ctx: _Contexto) -> None:
    _parrafo(
        doc,
        [
            _n(f"Manizales, {fecha_larga(ctx.fecha)}\t"),
            _b(f"{RADICADO_PREFIJO}-"),
            _y(M_NUMERO),
            _b(f"-{ctx.fecha.year}"),
        ],
        "FCEN Fecha",
    )
    _vacio(doc, n=2)
    _p(doc, "FCEN Destinatario", "Estudiante")
    _parrafo(doc, _nombre_dni(ctx.extraccion), "FCEN Destinatario")
    _parrafo(
        doc,
        [_n("Programa Curricular de "), _o(ctx.programa, "[programa]")],
        "FCEN Destinatario",
    )
    for linea in (
        "Facultad de Ciencias Exactas y Naturales",
        "Universidad Nacional de Colombia",
        "Sede Manizales",
    ):
        _p(doc, "FCEN Destinatario", linea)
    _vacio(doc, n=2)


def _sesion(d: DecisionConsejo) -> list[Parte]:
    """«en sesión … del 17 de septiembre de 2026, Acta 35»."""
    fecha = _fecha_iso(d.acta_fecha)
    return [
        _n(" en sesión "),
        _y(M_MODALIDAD),
        _n(" del "),
        _n(fecha_larga(fecha)) if fecha else _y("[fecha de la sesión]"),
        _n(", Acta "),
        _o(d.acta_numero, "[No.]"),
    ]


def _firma(doc: DocumentoDocx, cierre: str = "Sin otro particular,") -> None:
    _p(doc, "FCEN Firma", cierre)
    _vacio(doc, "FCEN Firma", n=3)
    nombre = settings.secretaria_facultad_nombre.strip()
    _parrafo(doc, [_ob(nombre.upper() if nombre else None, M_SECRETARIO)], "FCEN Firma")
    _p(
        doc,
        "FCEN Firma",
        settings.secretaria_facultad_cargo or "Secretario(a) de Facultad",
    )
    _vacio(doc, "FCEN Firma", n=2)


def _nota_acta(doc: DocumentoDocx, d: DecisionConsejo) -> None:
    """Lo que dice el acta, literal, para que la Secretaría tome de ahí la
    motivación. Va en amarillo: se borra antes de firmar."""
    p = _p(doc, "FCEN Cuerpo")
    _escribir(
        p,
        [
            _n(
                "[Borre este párrafo antes de firmar] Texto de la decisión en el acta: "
            ),
            _n(f"“{d.texto_decision}”. Fragmento del acta: “{d.fragmento}”"),
        ],
        todo_amarillo=True,
    )
    for r in p.runs:
        r.font.size = Pt(9)
    _vacio(doc, "FCEN Cuerpo")


def _tablas(doc: DocumentoDocx, tipos: list[TipoSolicitud], ctx: _Contexto) -> None:
    """Lo que el Consejo aprobó, cuando va en cuadro (solo si aprobó)."""
    if not ctx.aprobo:
        return
    e = ctx.extraccion
    if TipoSolicitud.HOMOLOGACION in tipos or TipoSolicitud.TRASLADO in tipos:
        _equivalencias(doc, e)
    if TipoSolicitud.DESIGNACION_JURADOS in tipos:
        _jurados(doc, e, intro="Jurados designados:")
    if TipoSolicitud.PROYECTO_TESIS in tipos:
        _proyecto(doc, e)


def _copia(doc: DocumentoDocx, ctx: _Contexto, registro: bool | None) -> None:
    partes: list[Parte] = [
        _n("Copia: \tComité Asesor Programa Curricular "),
        _o(ctx.programa, "[programa]"),
    ]
    if registro:
        partes.append(_n(", Oficina de Registro y Matrícula"))
    elif registro is None:
        partes.append(_y("[, Oficina de Registro y Matrícula]"))
    _parrafo(doc, partes, "FCEN Destinatario")


def _oficio(doc: DocumentoDocx, ctx: _Contexto) -> None:
    e, d = ctx.extraccion, ctx.decision
    cubiertos = [t for t in e.tipos if t in _TRAMITES]
    otros = [t for t in e.tipos if t not in _TRAMITES]
    generico = not cubiertos

    _encabezado(doc, ctx)

    # Asunto
    if generico:
        titulo = TITULOS_TIPO[otros[0]] if otros else "solicitud"
        asunto: list[Parte] = [_y(titulo[0].upper() + titulo[1:] + ".")]
    else:
        nombres = [[_b(_asunto(t, ctx))] for t in cubiertos]
        if any(_TRAMITES[t].lleva_periodo for t in cubiertos):
            nombres[-1] += [_b(" "), ctx.periodo_partes()]
        # En el asunto todo va en negrita, como en los oficios de la Secretaría.
        asunto = [(t, s or "b") for t, s in _y_lista(nombres)]
        asunto.append(_b("."))
    _parrafo(doc, [_b("Asunto:"), _b("\t"), *asunto])
    _vacio(doc, "FCEN Texto", n=2)
    _p(doc, "FCEN Texto", "Reciba un cordial saludo.")
    _vacio(doc, "FCEN Texto")

    # Decisión
    if generico:
        objeto = e.objeto_solicitud.strip().rstrip(".") or "[objeto de la solicitud]"
        objetos: list[Parte] = [_n(objeto)]
    else:
        objetos = _y_lista([_objeto(t, ctx) for t in cubiertos])
        if _PERIODO_AL_FINAL.intersection(cubiertos):
            objetos += _en_periodo(ctx)
    verbo = "APRUEBA" if ctx.aprobo else "NO APRUEBA"
    partes: list[Parte] = [
        _n(CONSEJO),
        *_sesion(d),
        _n(", "),
        _b(verbo),
        _n(" la solicitud de "),
        *objetos,
        _n(", realizada a través de "),
        _canal(ctx.caso),
        _n(" el "),
        _fecha_corta(e.fecha_solicitud),
    ]
    if not ctx.aprobo:
        partes += [_n(", "), _y(M_MOTIVACION)]
    partes.append(_n("."))
    _parrafo(doc, partes, todo_amarillo=generico)
    _vacio(doc, "FCEN Texto")

    if otros and not generico:
        titulos = ", ".join(TITULOS_TIPO[t] for t in otros)
        _parrafo(
            doc,
            [
                _y(
                    f"[El caso incluye además: {titulos}. Redacte esa parte de la decisión.]"
                )
            ],
        )
        _vacio(doc, "FCEN Texto")

    _tablas(doc, cubiertos, ctx)
    if TipoSolicitud.REEXPEDICION_RECIBO in cubiertos and ctx.aprobo:
        # Del oficio de la Secretaría para este trámite; solo si aplica.
        _parrafo(
            doc,
            [
                _n("En consecuencia, se "),
                _b("AUTORIZA"),
                _n(
                    " el desbloqueo de la Historia Académica del estudiante para "
                    "la reexpedición del recibo de pago correspondiente al período "
                    "académico mencionado por parte de la Oficina de Registro y "
                    "Matrícula."
                ),
            ],
            todo_amarillo=True,
        )
        _vacio(doc, "FCEN Texto")
    _nota_acta(doc, d)

    _p(doc, "FCEN Texto", NOTIFICACION)
    _vacio(doc, "FCEN Texto")
    _p(doc, "FCEN Texto", RECURSOS)
    _vacio(doc, "FCEN Texto")
    _firma(doc)
    registro = None if generico else any(_TRAMITES[t].registro for t in cubiertos)
    _copia(doc, ctx, registro)


# ---------------------------------------------------------------------------
# Resolución de reingreso (aprobado)
# ---------------------------------------------------------------------------

# Considerandos normativos copiados literalmente de la resolución de reingreso
# de la Secretaría. Salen en amarillo: se verifican contra la norma antes de
# firmar (el primero habla de posgrado aunque el modelo era de pregrado).
CONSIDERANDOS_REINGRESO = (
    "el Acuerdo 008 del 15 de abril de 2008 del Consejo Superior Universitario "
    "adoptó el Estatuto Estudiantil de la Universidad Nacional de Colombia en sus "
    "disposiciones académicas y en su Artículo 46. autorizó a través de los "
    "Consejos de Facultad el reingreso a estudiantes de posgrado, por una única "
    "vez, sólo a quienes presenten un Promedio Aritmético Ponderado Acumulado "
    "superior o igual a 2.7.",
    "el Parágrafo 3. – Artículo 46. del citado Acuerdo, estipula que los Consejos "
    "de Facultad no podrán otorgar reingreso a quién lo solicite después de tres "
    "años consecutivos de haber perdido la calidad de estudiante.",
    "la Resolución 012 del 11 de marzo de 2014 de la Vicerrectoría Académica, "
    "modificada en su Artículo 7. por la Resolución 037 del 10 de septiembre de "
    "2015, reglamentó el procedimiento para el estudio de solicitudes de "
    "reingreso a quienes hubieren perdido la calidad de estudiantes de pregrado o "
    "posgrado en la Universidad Nacional de Colombia.",
)

ATRIBUCIONES = (
    "En uso de sus atribuciones legales y estatutarias establecidas en el Acuerdo "
    "011 de 2005 del Consejo Superior Universitario, y"
)

M_PERIODO_REINGRESO = "[primer/segundo período académico de AAAA]"


def _centrado(doc: DocumentoDocx, partes: Sequence[Parte]) -> None:
    p = _parrafo(doc, partes, "FCEN Destinatario")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _dia_en_letras(n: int) -> str:
    return "uno" if n == 1 else en_letras(n)


def _resolucion(doc: DocumentoDocx, ctx: _Contexto) -> None:
    e, d = ctx.extraccion, ctx.decision
    fecha_acta = _fecha_iso(d.acta_fecha)
    exe = "la exestudiante" if e.tratamiento == "La estudiante" else "el exestudiante"
    cual = "la cual" if e.tratamiento == "La estudiante" else "el cual"
    al = "a la exestudiante" if e.tratamiento == "La estudiante" else "al exestudiante"
    nombre = _nombre_dni(e)
    programa = _o(ctx.programa, "[programa]")
    periodo = _y(M_PERIODO_REINGRESO)

    _centrado(doc, [_b("RESOLUCIÓN "), _y("[Res.CFCEN-___-AAAA]")])
    _centrado(
        doc,
        [
            _b("("),
            _b(f"{fecha_acta.day:02d} DE {_MESES[fecha_acta.month - 1].upper()}")
            if fecha_acta
            else _y("[DD DE MES]"),
            _b(", ACTA "),
            _ob(d.acta_numero, "[No.]"),
            _b(")"),
        ],
    )
    _centrado(doc, [_b("_" * 50)])
    _vacio(doc)
    _centrado(
        doc,
        [
            _b("“Por la cual se autoriza un Reingreso al programa curricular de "),
            _ob(ctx.programa, "[programa]"),
            _b(" de la Facultad de Ciencias Exactas y Naturales a partir del "),
            periodo,
            _b("”"),
        ],
    )
    _centrado(doc, [_b("_" * 35)])
    _vacio(doc)
    _centrado(
        doc,
        [
            _b(
                "EL CONSEJO DE LA FACULTAD DE CIENCIAS EXACTAS Y NATURALES, DE LA "
                "UNIVERSIDAD NACIONAL DE COLOMBIA"
            )
        ],
    )
    _vacio(doc)
    _centrado(doc, [_b(ATRIBUCIONES)])
    _vacio(doc, n=2)
    _centrado(doc, [_b("CONSIDERANDO")])
    _vacio(doc)
    _parrafo(
        doc,
        [
            _y(
                "[Considerandos normativos copiados del modelo de la Secretaría: "
                "verifíquelos contra la norma vigente antes de firmar.]"
            )
        ],
    )
    for texto in CONSIDERANDOS_REINGRESO:
        _parrafo(doc, [_b("QUE "), _y(texto)])
        _vacio(doc, "FCEN Texto")

    recomendo: Parte = (
        _n("recomendó aprobar")
        if ctx.caso.sentido == Sentido.APROBAR
        else _y("[recomendó aprobar]")
    )
    _parrafo(
        doc,
        [
            _b("QUE"),
            _n(" el Comité Asesor de "),
            programa,
            _n(" en sesión "),
            _y("[Ad-referéndum / presencial]"),
            _n(" del "),
            _y("[fecha de la sesión del comité]"),
            _n(", en ("),
            _b("Acta "),
            _y("[No.]"),
            _n("), "),
            recomendo,
            _n(" la solicitud de reingreso a partir del "),
            periodo,
            _n(f", realizada por {exe} "),
            *nombre,
            _n(", el "),
            _n(fecha_larga(e.fecha_solicitud))
            if _fecha_iso(e.fecha_solicitud)
            else _y("[fecha de la solicitud]"),
            _n(" a través de "),
            _y("[oficio / correo electrónico]"),
            _n(", "),
            _y(
                f"{cual} manifestó explícitamente que conoce la normatividad que "
                "tiene la Universidad Nacional de Colombia sobre los reingresos"
            ),
            _n("."),
        ],
    )
    _vacio(doc, "FCEN Texto")
    _parrafo(
        doc,
        [
            _b("QUE "),
            _n(CONSEJO[0].lower() + CONSEJO[1:]),
            *_sesion(d),
            _n(
                ", con fundamento en la recomendación presentada por el Comité "
                "precitado, aprobó la solicitud de reingreso a partir del "
            ),
            periodo,
            _n(f" {al} "),
            *nombre,
            _n(", para continuar sus estudios en el programa curricular de "),
            programa,
            _n("."),
        ],
    )
    _vacio(doc, "FCEN Texto")
    _p(doc, "FCEN Texto", "En mérito de lo anterior,")
    _vacio(doc, "FCEN Texto")
    _centrado(doc, [_b("RESUELVE")])
    _vacio(doc)
    _parrafo(
        doc,
        [
            _b("ARTÍCULO 1."),
            _n(f" Autorizar reingreso {al} "),
            *nombre,
            _n(", para que, a partir del "),
            periodo,
            _n(", continúe sus estudios en el programa curricular de "),
            programa,
            _n("."),
        ],
    )
    _vacio(doc, "FCEN Texto")
    nivel = "posgrado" if ctx.posgrado else "pregrado"
    _parrafo(
        doc,
        [
            _b("PARÁGRAFO: "),
            _n(exe[0].upper() + exe[1:] + " "),
            nombre[0],
            _n(
                f" deberá tener en cuenta el calendario académico detallado de "
                f"{nivel} para el "
            ),
            _b("pago de su matrícula"),
            _n(" en el "),
            periodo,
            _n(
                " en la Sede Manizales de la Universidad Nacional de Colombia, "
                "establecido por el Consejo de Sede, de lo contrario este acto "
                "académico quedaría sin efecto."
            ),
        ],
    )
    _vacio(doc, "FCEN Texto")
    _parrafo(
        doc,
        [
            _b("ARTÍCULO 2."),
            _n(
                " Notificar el contenido de esta disposición por medio de la "
                f"Secretaría de la Facultad de Ciencias Exactas y Naturales {al} "
            ),
            *nombre,
            _n(
                ", de conformidad con lo señalado en el artículo 67 y siguientes del "
                "Código de Procedimiento Administrativo y de lo Contencioso "
                "Administrativo y haciéndole saber que contra la misma procede el "
                "recurso de reposición, el cual podrá interponer ante el Consejo de "
                "Facultad de Ciencias Exactas y Naturales dentro de los diez (10) "
                "días siguientes a la fecha de notificación, de conformidad con lo "
                "establecido en el mismo código."
            ),
        ],
    )
    _vacio(doc, "FCEN Texto")
    _parrafo(
        doc,
        [
            _b("ARTÍCULO 3."),
            _n(" La presente Resolución rige a partir de la fecha de su expedición."),
        ],
    )
    _vacio(doc, "FCEN Texto")
    _nota_acta(doc, d)
    _centrado(doc, [_b("NOTIFÍQUESE Y CÚMPLASE")])
    _vacio(doc)
    f = ctx.fecha
    _centrado(
        doc,
        [
            _b(
                f"Dada en Manizales, a {_dia_en_letras(f.day)} ({f.day:02d}) de "
                f"{_MESES[f.month - 1]} de {f.year}"
            )
        ],
    )
    _vacio(doc, n=4)

    presidente = settings.presidente_consejo_nombre.strip()
    secretario = settings.secretaria_facultad_nombre.strip()
    p = _parrafo(
        doc,
        [
            _ob(presidente.upper() if presidente else None, M_PRESIDENTE),
            _n("\t"),
            _ob(secretario.upper() if secretario else None, M_SECRETARIO),
        ],
        "FCEN Destinatario",
    )
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9))
    p = _p(
        doc,
        "FCEN Destinatario",
        f"{settings.presidente_consejo_cargo or 'Presidente(a)'}\t"
        f"{_cargo_corto(settings.secretaria_facultad_cargo)}",
    )
    p.paragraph_format.tab_stops.add_tab_stop(Cm(9))
    _vacio(doc, n=2)
    _parrafo(
        doc,
        [
            _n(
                "Copia: \tDecanatura FCEN, Sección Biblioteca, Oficina de Registro y "
                "Matrícula, Programa de "
            ),
            programa,
        ],
        "FCEN Destinatario",
    )


def _cargo_corto(cargo: str) -> str:
    """En la resolución firma «Secretario(a)», sin «de Facultad»."""
    return (cargo or "Secretario(a)").replace(" de Facultad", "")


# ---------------------------------------------------------------------------
# Entrada
# ---------------------------------------------------------------------------


def es_resolucion(caso: Caso, decision: DecisionConsejo | None = None) -> bool:
    """El reingreso aprobado se expide por resolución, no por oficio."""
    d = decision or caso.consejo
    e = caso.extraccion
    return bool(
        d is not None
        and d.sentido == SentidoConsejo.APROBO
        and e is not None
        and e.tipos
        and e.tipos[0] == TipoSolicitud.REINGRESO
    )


def nombre_archivo(caso: Caso, decision: DecisionConsejo | None = None) -> str:
    acto = "resolucion_reingreso" if es_resolucion(caso, decision) else "oficio"
    return f"Borrador_{acto}_Consejo_caso_{caso.id}.docx"


def construir_oficio_consejo(
    caso: Caso,
    decision: DecisionConsejo | None = None,
    *,
    fecha: date | None = None,
    periodo: str | None = None,
    elaboro: str = "",
) -> bytes:
    """Borrador (.docx) de la notificación de la decisión del Consejo de
    Facultad. Lanza `SinOficio` si el caso no tiene una decisión final del
    Consejo de Facultad confirmada."""
    ctx = _contexto(caso, decision, fecha, periodo)
    resolucion = es_resolucion(caso, ctx.decision)
    doc = _documento(
        "Consejo de Facultad" if resolucion else "Secretaría de Facultad", elaboro
    )
    _aviso_borrador(doc)
    if resolucion:
        _resolucion(doc, ctx)
    else:
        _oficio(doc, ctx)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
