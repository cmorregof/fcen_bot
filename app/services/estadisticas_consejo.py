"""Estadísticas de las solicitudes que decide el Consejo de Facultad, para Calidad.

La Secretaría de Facultad entrega a Calidad un consolidado por semestre de las
solicitudes estudiantiles decididas por el Consejo de Facultad, incluidas las
movilidades. MemorIA cuenta las que pasaron por los comités asesores.

Reglas (las mismas de `suit.py`):
- Solo cifras: ni nombres ni documentos.
- Cuenta una solicitud decidida por el Consejo de Facultad en un acta
  confirmada en MemorIA (aprobó, no aprobó u otra decisión). Un aplazamiento
  no la decide; los casos descartados no entran.
- El periodo y el mes son los de la sesión del Consejo que consta en el acta.
- Un caso con varios trámites cuenta una vez en cada trámite; aparte queda el
  total de casos distintos.
- Movilidad: es un trámite de MemorIA (`TipoSolicitud.MOVILIDAD`) y va en su
  propia hoja. MemorIA no sabe si una homologación viene de una movilidad, y
  las movilidades que no pasan por el comité asesor no aparecen.
"""

from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from datetime import date

import aiosqlite
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.core.programas import PROGRAMAS
from app.schemas.caso import TITULOS_TIPO, EstadoCaso, TipoSolicitud
from app.schemas.consejo import ETIQUETA_SENTIDO_CONSEJO, Instancia, SentidoConsejo
from app.services.notificaciones import PeriodoInvalido, Semestre, semestre_de

MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
FACULTAD = "Ciencias Exactas y Naturales"
MESES = ("Ene", "Feb", "Mar", "Abr", "May", "Jun",
         "Jul", "Ago", "Sep", "Oct", "Nov", "Dic")  # fmt: skip
SENTIDOS = (SentidoConsejo.APROBO, SentidoConsejo.NO_APROBO, SentidoConsejo.OTRO)
SIN_TRAMITE = "sin_tramite"  # caso sin trámite reconocido en la ficha
SIN_PROGRAMA = "sin_programa"
NIVELES = ("pregrado", "posgrado")

# Trámites en el orden del enum, más la fila de casos sin trámite reconocido.
TRAMITES: tuple[str, ...] = tuple(t.value for t in TipoSolicitud) + (SIN_TRAMITE,)


def titulo_tramite(clave: str) -> str:
    if clave == SIN_TRAMITE:
        return "Sin trámite registrado en la ficha"
    titulo = TITULOS_TIPO[TipoSolicitud(clave)]
    return titulo[:1].upper() + titulo[1:]


def periodo_de(texto: str, hoy: date) -> Semestre:
    """«2026-1S», «2026-2S» o el año completo «2026»; vacío = semestre actual."""
    limpio = texto.strip()
    if limpio.isdigit() and len(limpio) == 4 and 2000 <= int(limpio) <= 2100:
        ano = int(limpio)
        return Semestre(limpio, date(ano, 1, 1), date(ano, 12, 31))
    try:
        return semestre_de(limpio, hoy)
    except PeriodoInvalido:
        raise PeriodoInvalido(
            "Escriba el semestre así: 2026-1S o 2026-2S, o el año completo: 2026."
        ) from None


def meses_de(periodo: Semestre) -> list[tuple[int, int]]:
    """(año, mes) de cada mes del periodo, en orden."""
    salida = []
    a, m = periodo.desde.year, periodo.desde.month
    while (a, m) <= (periodo.hasta.year, periodo.hasta.month):
        salida.append((a, m))
        a, m = (a + 1, 1) if m == 12 else (a, m + 1)
    return salida


def _tipos(extraccion_json: str | None) -> set[str]:
    try:
        datos = json.loads(extraccion_json or "")
    except ValueError:
        return set()
    salida: set[str] = set()
    for t in datos.get("tipos", []) if isinstance(datos, dict) else []:
        try:
            salida.add(TipoSolicitud(t).value)
        except ValueError:
            continue
    return salida


def _vacio() -> dict[SentidoConsejo, int]:
    return {s: 0 for s in SENTIDOS}


@dataclass
class Estadisticas:
    periodo: Semestre
    corte: date
    # trámite -> sentido -> número de solicitudes
    por_tramite: dict[str, dict[SentidoConsejo, int]] = field(default_factory=dict)
    # casos distintos por sentido
    casos: dict[SentidoConsejo, int] = field(default_factory=_vacio)
    # programa -> trámite -> número de solicitudes
    por_programa: dict[str, dict[str, int]] = field(default_factory=dict)
    # programa -> casos distintos
    casos_programa: dict[str, int] = field(default_factory=dict)
    # (año, mes) -> trámite -> número de solicitudes
    por_mes: dict[tuple[int, int], dict[str, int]] = field(default_factory=dict)
    # (año, mes) -> casos distintos
    casos_mes: dict[tuple[int, int], int] = field(default_factory=dict)
    # nivel -> sentido -> casos distintos de movilidad
    movilidad_nivel: dict[str, dict[SentidoConsejo, int]] = field(default_factory=dict)

    def total_tramite(self, tramite: str) -> int:
        return sum(self.por_tramite.get(tramite, {}).values())

    @property
    def total_casos(self) -> int:
        return sum(self.casos.values())


async def calcular(
    db: aiosqlite.Connection, periodo: Semestre, hoy: date
) -> Estadisticas:
    r = Estadisticas(periodo, hoy)
    cur = await db.execute(
        "SELECT c.programa, c.extraccion_json, d.sentido, d.acta_fecha"
        " FROM casos c JOIN decisiones_consejo d ON d.caso_id = c.id"
        " WHERE c.estado != ? AND d.instancia = ?",
        (EstadoCaso.DESCARTADO.value, Instancia.FACULTAD.value),
    )
    for programa, extraccion, sentido_txt, acta_fecha in await cur.fetchall():
        try:
            sentido = SentidoConsejo(sentido_txt)
            dia = date.fromisoformat(str(acta_fecha)[:10])
        except ValueError:
            continue
        if sentido not in SENTIDOS or not periodo.desde <= dia <= periodo.hasta:
            continue
        prog = programa if programa in PROGRAMAS else SIN_PROGRAMA
        mes = (dia.year, dia.month)
        tramites = _tipos(extraccion) or {SIN_TRAMITE}

        r.casos[sentido] += 1
        r.casos_programa[prog] = r.casos_programa.get(prog, 0) + 1
        r.casos_mes[mes] = r.casos_mes.get(mes, 0) + 1
        for t in tramites:
            r.por_tramite.setdefault(t, _vacio())[sentido] += 1
            fila = r.por_programa.setdefault(prog, {})
            fila[t] = fila.get(t, 0) + 1
            fila_mes = r.por_mes.setdefault(mes, {})
            fila_mes[t] = fila_mes.get(t, 0) + 1
        if TipoSolicitud.MOVILIDAD.value in tramites:
            nivel = PROGRAMAS[prog].nivel if prog in PROGRAMAS else SIN_PROGRAMA
            r.movilidad_nivel.setdefault(nivel, _vacio())[sentido] += 1
    return r


_GRIS = PatternFill("solid", fgColor="D9D9D9")


def _encabezado(ws: Worksheet, fila: int) -> None:
    for celda in ws[fila]:
        celda.font = Font(bold=True)
        celda.fill = _GRIS
        celda.alignment = Alignment(wrap_text=True, horizontal="center")


def _nombre_programa(clave: str) -> tuple[str, str]:
    if clave in PROGRAMAS:
        return PROGRAMAS[clave].nombre, PROGRAMAS[clave].nivel
    return "Sin programa reconocido", "—"


def _orden_programas(r: Estadisticas) -> list[str]:
    claves = sorted(
        PROGRAMAS,
        key=lambda c: (NIVELES.index(PROGRAMAS[c].nivel), PROGRAMAS[c].nombre),
    )
    if SIN_PROGRAMA in r.casos_programa:
        claves.append(SIN_PROGRAMA)
    return claves


def leeme(r: Estadisticas) -> list[str]:
    return [
        f"Estadísticas del Consejo de Facultad para Calidad — {r.periodo.texto} "
        f"({r.periodo.desde.isoformat()} a {r.periodo.hasta.isoformat()}), "
        f"generado por MemorIA el {r.corte.isoformat()}",
        "Qué cuenta: las solicitudes estudiantiles decididas por el Consejo de "
        f"Facultad de {FACULTAD}, es decir, las que tienen la decisión registrada "
        "desde un acta del Consejo confirmada en MemorIA (aprobó, no aprobó u otra "
        "decisión). El periodo y el mes son los de la sesión del Consejo.",
        "Qué no cuenta: los aplazamientos (el caso sigue pendiente), los casos "
        "descartados, las decisiones del Consejo de Sede o del Consejo Superior, los "
        "casos sin acta del Consejo confirmada y todo lo que no pasó por MemorIA.",
        "Un caso con varios trámites cuenta una vez en cada trámite; por eso la suma "
        "de los trámites puede ser mayor que el total de casos distintos, que va "
        "aparte en «Resumen», «Por programa» y «Por mes».",
        "«Sin trámite registrado en la ficha» son casos decididos cuya ficha no "
        "trae un trámite reconocido.",
        "Movilidades: la hoja «Movilidades» cuenta los casos con el trámite "
        "«movilidad» que pasaron por el comité asesor y el Consejo de Facultad en "
        "MemorIA. MemorIA no registra si una homologación o equivalencia viene de "
        "una movilidad: esas van en la fila de homologación y no se suman a "
        "movilidades. Las movilidades que no pasan por los comités en MemorIA "
        "no aparecen y hay que tomarlas de su fuente.",
        "Solo cifras: el archivo no lleva nombres ni documentos. Revise contra los "
        "registros de la Secretaría antes de entregar las cifras.",
    ]


def xlsx(r: Estadisticas) -> bytes:
    libro = Workbook()

    # Resumen: totales por trámite y sentido.
    ws = libro.active
    assert ws is not None
    ws.title = "Resumen"
    ws.append([f"Solicitudes decididas por el Consejo de Facultad — {r.periodo.texto}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"Facultad de {FACULTAD}. Corte: {r.corte.isoformat()}"])
    ws.append([])
    ws.append(["Trámite"] + [ETIQUETA_SENTIDO_CONSEJO[s] for s in SENTIDOS] + ["Total"])
    _encabezado(ws, 4)
    suma = _vacio()
    for t in TRAMITES:
        s = r.por_tramite.get(t, _vacio())
        ws.append([titulo_tramite(t)] + [s[x] for x in SENTIDOS] + [sum(s.values())])
        for x in SENTIDOS:
            suma[x] += s[x]
    ws.append(
        ["Total de trámites"] + [suma[x] for x in SENTIDOS] + [sum(suma.values())]
    )
    ws.append(["Casos distintos"] + [r.casos[x] for x in SENTIDOS] + [r.total_casos])
    for fila in ws.iter_rows(min_row=ws.max_row - 1):
        for celda in fila:
            celda.font = Font(bold=True)
    ws.column_dimensions["A"].width = 60
    ws.freeze_panes = "B5"

    # Por programa: programa × trámite.
    pp = libro.create_sheet("Por programa")
    pp.append(
        ["Programa", "Nivel"]
        + [titulo_tramite(t) for t in TRAMITES]
        + ["Total de trámites", "Casos distintos"]
    )
    _encabezado(pp, 1)

    def fila_programa(nombre: str, nivel: str, claves: list[str]) -> list[object]:
        valores = [
            sum(r.por_programa.get(c, {}).get(t, 0) for c in claves) for t in TRAMITES
        ]
        casos = sum(r.casos_programa.get(c, 0) for c in claves)
        return [nombre, nivel, *valores, sum(valores), casos]

    programas = _orden_programas(r)
    for clave in programas:
        nombre, nivel = _nombre_programa(clave)
        pp.append(fila_programa(nombre, nivel, [clave]))
    for nivel in NIVELES:
        claves = [
            c for c in programas if c in PROGRAMAS and PROGRAMAS[c].nivel == nivel
        ]
        pp.append(fila_programa(f"Total {nivel}", nivel, claves))
        for celda in pp[pp.max_row]:
            celda.font = Font(bold=True)
    pp.append(fila_programa("Total Facultad", "—", programas))
    for celda in pp[pp.max_row]:
        celda.font = Font(bold=True)
    pp.column_dimensions["A"].width = 52
    pp.freeze_panes = "C2"

    # Por mes: trámite × mes de la sesión del Consejo.
    meses = meses_de(r.periodo)
    pm = libro.create_sheet("Por mes")
    pm.append(["Trámite"] + [f"{MESES[m - 1]} {a}" for a, m in meses] + ["Total"])
    _encabezado(pm, 1)
    for t in TRAMITES:
        valores = [r.por_mes.get(am, {}).get(t, 0) for am in meses]
        pm.append([titulo_tramite(t), *valores, sum(valores)])
    casos_mes = [r.casos_mes.get(am, 0) for am in meses]
    pm.append(["Casos distintos", *casos_mes, sum(casos_mes)])
    for celda in pm[pm.max_row]:
        celda.font = Font(bold=True)
    pm.column_dimensions["A"].width = 60
    pm.freeze_panes = "B2"

    # Movilidades: el trámite «movilidad», aparte, por nivel y por programa.
    mv = libro.create_sheet("Movilidades")
    mv.append(["Movilidades decididas por el Consejo de Facultad en MemorIA"])
    mv["A1"].font = Font(bold=True, size=13)
    mv.append([])
    mv.append(["Nivel"] + [ETIQUETA_SENTIDO_CONSEJO[s] for s in SENTIDOS] + ["Total"])
    _encabezado(mv, 3)
    total_mv = _vacio()
    niveles = list(NIVELES) + (
        [SIN_PROGRAMA] if SIN_PROGRAMA in r.movilidad_nivel else []
    )
    for nivel in niveles:
        s = r.movilidad_nivel.get(nivel, _vacio())
        etiqueta = "Sin programa reconocido" if nivel == SIN_PROGRAMA else nivel
        mv.append([etiqueta] + [s[x] for x in SENTIDOS] + [sum(s.values())])
        for x in SENTIDOS:
            total_mv[x] += s[x]
    mv.append(["Total"] + [total_mv[x] for x in SENTIDOS] + [sum(total_mv.values())])
    for celda in mv[mv.max_row]:
        celda.font = Font(bold=True)
    mv.append([])
    mv.append(["Programa", "Nivel", "Movilidades"])
    _encabezado(mv, mv.max_row)
    for clave in programas:
        nombre, nivel = _nombre_programa(clave)
        mv.append([nombre, nivel, r.por_programa.get(clave, {}).get("movilidad", 0)])
    mv.append([])
    for nota in (
        "Las homologaciones y equivalencias van en «Resumen» con su propio trámite: "
        "MemorIA no sabe si vienen de una movilidad.",
        "Las movilidades que no pasan por los comités en MemorIA no aparecen: "
        "tómelas de su fuente.",
    ):
        mv.append([nota])
    mv.column_dimensions["A"].width = 52

    # Léeme.
    le = libro.create_sheet("Léeme")
    le.column_dimensions["A"].width = 120
    for linea in leeme(r):
        le.append([linea])
    le["A1"].font = Font(bold=True, size=13)
    for fila in le.iter_rows(min_row=2):
        fila[0].alignment = Alignment(wrap_text=True)

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def nombre_xlsx(r: Estadisticas) -> str:
    return f"estadisticas_consejo_fcen_{r.periodo.texto}.xlsx"
