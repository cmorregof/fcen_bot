"""Archivo oficial en Drive de actas, oficios y resoluciones firmados.

Cuándo se archiva: MemorIA solo produce borradores (Word con campos en
amarillo y «[la firma se agrega al firmar]»); la firma y el número los pone
una persona fuera de MemorIA. Por eso lo que se archiva es el PDF firmado que
esa persona sube y confirma: la Secretaría, el oficio o la resolución de un
caso con número y fecha de expedición registrados; el área, el acta firmada
del comité. Nunca se archiva un .docx ni un PDF con las marcas de borrador.
Ver docs/respaldo-drive.md.

Estructura: `<DRIVE_CARPETA_ID>/<Tipo>/<año>/<área>/<programa>/<archivo>.pdf`,
por ejemplo `Actas/2026/Ciencias Naturales/Ingeniería Física/`. En cada
carpeta de año hay un índice (xlsx) de lo archivado ese año.
"""

from __future__ import annotations

import asyncio
import io
import re
from dataclasses import dataclass
from datetime import date, datetime
from typing import Literal
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, Side

from app.core.config import settings
from app.core.programas import DIRECCIONES, PROGRAMAS
from app.services.drive import cliente as drive

TipoDocumento = Literal["acta", "oficio", "resolucion"]

CARPETA_TIPO: dict[TipoDocumento, str] = {
    "acta": "Actas",
    "oficio": "Oficios",
    "resolucion": "Resoluciones",
}
ETIQUETA_TIPO: dict[TipoDocumento, str] = {
    "acta": "Acta",
    "oficio": "Oficio",
    "resolucion": "Resolución",
}
PRODUCTORA: dict[TipoDocumento, str] = {
    "acta": "Comités Asesores de las Áreas Curriculares",
    "oficio": "Secretaría de Facultad",
    "resolucion": "Secretaría de Facultad",
}
MIME_PDF = "application/pdf"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 25 * 1024 * 1024

# Marcas que MemorIA pone en sus borradores (documento.py, oficio_consejo.py).
MARCAS_BORRADOR = (
    "[la firma se agrega al firmar]",
    "BORRADOR · pendiente de revisión",
)

COLUMNAS_INDICE = (
    "No.",
    "Fecha del documento",
    "Tipo documental",
    "Número o radicado",
    "Área curricular",
    "Programa",
    "Nombre del archivo",
    "Fecha de incorporación",
    "Formato",
    "Tamaño",
    "Archivado por",
)
_FILA_TITULOS = 10
# Una subida a la vez: el índice se lee, se completa y se reemplaza.
_candado = asyncio.Lock()


class DocumentoNoValido(ValueError):
    """No es un PDF firmado: no se archiva."""


class YaArchivado(ValueError):
    """Ya hay un archivo con ese nombre en la carpeta."""


@dataclass(frozen=True)
class Documento:
    tipo: TipoDocumento
    numero: str
    fecha: date
    programa: str  # clave de PROGRAMAS
    contenido: bytes
    archivado_por: str


def nombre_area(direccion: str) -> str:
    return DIRECCIONES[direccion].nombre.removeprefix("Área Curricular de ").strip()


def _seguro(texto: str) -> str:
    limpio = re.sub(r"[^\w.\-]+", "-", texto.strip(), flags=re.UNICODE)
    return re.sub(r"-{2,}", "-", limpio).strip("-.")


def nombre_archivo(doc: Documento) -> str:
    numero = _seguro(doc.numero) or "sin-numero"
    return f"{ETIQUETA_TIPO[doc.tipo]}_{numero}_{doc.programa}_{doc.fecha:%Y%m%d}.pdf"


def nombre_indice(tipo: TipoDocumento, ano: int) -> str:
    return f"Índice {CARPETA_TIPO[tipo]} {ano}.xlsx"


def ruta(doc: Documento) -> list[str]:
    programa = PROGRAMAS[doc.programa]
    return [
        CARPETA_TIPO[doc.tipo],
        str(doc.fecha.year),
        nombre_area(programa.direccion),
        programa.nombre,
    ]


def validar(doc: Documento) -> None:
    """Solo PDF y sin las marcas de borrador de MemorIA."""
    if doc.programa not in PROGRAMAS:
        raise DocumentoNoValido("Programa desconocido.")
    if not doc.numero.strip():
        raise DocumentoNoValido("Falta el número o radicado del documento.")
    if len(doc.contenido) > MAX_BYTES:
        raise DocumentoNoValido("El PDF pasa de 25 MB.")
    if not doc.contenido.startswith(b"%PDF-"):
        raise DocumentoNoValido(
            "Solo se archiva la versión firmada en PDF; un Word es un borrador."
        )
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        lector = PdfReader(io.BytesIO(doc.contenido))
        texto = "\n".join((p.extract_text() or "") for p in lector.pages)
    except (PdfReadError, ValueError, KeyError) as exc:
        raise DocumentoNoValido("No pude abrir el PDF.") from exc
    if any(m.casefold() in texto.casefold() for m in MARCAS_BORRADOR):
        raise DocumentoNoValido(
            "Ese PDF es un borrador de MemorIA (tiene las marcas de borrador). "
            "Suba la versión firmada."
        )


# ---------------------------------------------------------------------------
# Índice
# ---------------------------------------------------------------------------


def _filas(contenido: bytes) -> list[list[str]]:
    libro = load_workbook(io.BytesIO(contenido), read_only=True)
    ws = libro.worksheets[0]
    filas: list[list[str]] = []
    for valores in ws.iter_rows(min_row=_FILA_TITULOS + 1, values_only=True):
        celdas = [
            "" if v is None else str(v) for v in valores[1 : len(COLUMNAS_INDICE) + 1]
        ]
        if not any(celdas):
            break
        filas.append(celdas)
    libro.close()
    return filas


def indice_xlsx(tipo: TipoDocumento, ano: int, filas: list[list[str]]) -> bytes:
    """Mismo estilo que el índice electrónico de calidad: encabezado del fondo,
    títulos en negrita con borde y una fila por documento."""
    libro = Workbook()
    ws = libro.active
    assert ws is not None
    ws.title = "Índice"
    borde = Border(*(Side(style="thin"),) * 4)
    negrita = Font(bold=True)
    for col, ancho in zip(
        "ABCDEFGHIJKL", [2, 6, 13, 14, 24, 24, 26, 44, 14, 9, 10, 30], strict=True
    ):
        ws.column_dimensions[col].width = ancho
    for fila, texto in [
        (2, "FONDO UNIVERSIDAD NACIONAL DE COLOMBIA"),
        (3, "SUBFONDO SEDE MANIZALES"),
        (4, "Facultad de Ciencias Exactas y Naturales"),
    ]:
        ws.merge_cells(f"B{fila}:L{fila}")
        ws[f"B{fila}"] = texto
        ws[f"B{fila}"].font = Font(bold=True, size=13)
        ws[f"B{fila}"].alignment = Alignment(horizontal="center")
    for i, (rotulo, valor) in enumerate(
        [
            ("Oficina productora", PRODUCTORA[tipo]),
            ("Serie", CARPETA_TIPO[tipo].upper()),
            ("Año", str(ano)),
        ],
        start=6,
    ):
        ws.merge_cells(f"C{i}:D{i}")
        ws.merge_cells(f"E{i}:H{i}")
        ws[f"C{i}"], ws[f"E{i}"] = rotulo, valor
        ws[f"C{i}"].font = negrita
        for c in ("C", "E"):
            ws[f"{c}{i}"].border = borde
    for j, titulo in enumerate(COLUMNAS_INDICE, start=2):
        c = ws.cell(row=_FILA_TITULOS, column=j, value=titulo)
        c.font = negrita
        c.border = borde
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[_FILA_TITULOS].height = 32
    for k, registro in enumerate(filas):
        for j, valor in enumerate(registro, start=2):
            c = ws.cell(row=_FILA_TITULOS + 1 + k, column=j, value=valor)
            c.border = borde
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = f"B{_FILA_TITULOS + 1}"
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _tamano(n: int) -> str:
    return f"{max(1, round(n / 1024))} kB"


async def _actualizar_indice(
    cliente: drive.ClienteDrive,
    carpeta_ano: str,
    doc: Documento,
    archivo: str,
    hoy: date,
) -> None:
    nombre = nombre_indice(doc.tipo, doc.fecha.year)
    existente = next(
        (a for a in await cliente.listar(carpeta_ano) if a.nombre == nombre), None
    )
    filas = (
        await asyncio.to_thread(_filas, await cliente.descargar(existente.id))
        if existente
        else []
    )
    programa = PROGRAMAS[doc.programa]
    filas.append(
        [
            str(len(filas) + 1),
            doc.fecha.isoformat(),
            ETIQUETA_TIPO[doc.tipo],
            doc.numero.strip(),
            nombre_area(programa.direccion),
            programa.nombre,
            archivo,
            hoy.isoformat(),
            "PDF",
            _tamano(len(doc.contenido)),
            doc.archivado_por,
        ]
    )
    contenido = indice_xlsx(doc.tipo, doc.fecha.year, filas)
    if existente:
        await cliente.reemplazar(existente.id, contenido, MIME_XLSX)
    else:
        await cliente.subir(nombre, carpeta_ano, contenido, MIME_XLSX)


async def archivar(
    cliente: drive.ClienteDrive,
    raiz: str,
    doc: Documento,
    hoy: date | None = None,
) -> str:
    """Sube el PDF firmado y agrega su fila al índice del año. Devuelve la
    ruta legible en Drive."""
    validar(doc)
    hoy = hoy or datetime.now(ZoneInfo(settings.timezone)).date()
    async with _candado:
        return await _archivar(cliente, raiz, doc, hoy)


async def _archivar(
    cliente: drive.ClienteDrive, raiz: str, doc: Documento, hoy: date
) -> str:
    archivo = nombre_archivo(doc)
    partes = ruta(doc)
    carpeta_ano = await drive.carpeta(cliente, raiz, partes[:2])
    destino = await drive.carpeta(cliente, carpeta_ano, partes[2:])
    if any(a.nombre == archivo for a in await cliente.listar(destino)):
        raise YaArchivado(f"Ya está archivado: {'/'.join([*partes, archivo])}")
    await cliente.subir(archivo, destino, doc.contenido, MIME_PDF)
    await _actualizar_indice(cliente, carpeta_ano, doc, archivo, hoy)
    return "/".join([*partes, archivo])
